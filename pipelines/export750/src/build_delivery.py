"""Create the Chinese workbook, documentation and verified delivery manifest."""
from flat_common import *
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


TRANSLATIONS = {
    'category_id': '类别编号', 'category_name': '类别名称', 'documents': '记录数',
    'countable_documents': '可统计记录数（非撤稿）',
    'supported_documents': '自动规则支持数（非撤稿）',
    'needs_review_documents': '待复核记录数', 'retracted_documents': '撤稿记录数',
    'text_cleaned_documents': '文本清理记录数', 'semantic_label_changes': '研究标签调整记录数',
    'paper_documents': '论文数', 'patent_documents': '专利数', 'policy_documents': '政策数',
    'quarantine_code': '隔离状态编号', 'quarantine_reason': '隔离原因',
    'baseline_category_id': '原类别编号', 'baseline_category_name': '原类别名称',
    'status': '决定状态代码', 'description': '说明', 'quality': '质量状态代码',
    'source': '来源代码', 'review_reason': '待复核原因',
}


def main():
    validation = read(FLAT / 'VALIDATION.json')
    assert validation['passed']
    summary = read(FLAT / 'results/SUMMARY.json')
    totals = summary['totals']
    catalog = pd.read_csv(FLAT / 'results/category_dictionary.csv')
    quality = pd.read_csv(FLAT / 'results/quarantine_dictionary.csv')
    k = len(catalog)
    notes = pd.DataFrame([
        ['类别口径', f'当前共{k}类，统一编号C0001–C{k:04d}；所有类别在同一张表中平级列出。'],
        ['文献归属', '每条已归类记录只有一个category_id，只在该类别计数一次。'],
        ['质量隔离', f'{totals["quarantined"]:,}条质量隔离记录的类别为空，不计入研究类别数。'],
        ['本次修改', '将现有正式标签转换为统一平级编号；保留每条记录的类别含义和归属，不重新运行语义分类。'],
        ['标签调整统计', '研究标签调整数对照最初底稿计算，不将这次编号转换计作新的语义调整。'],
        ['待复核', f'仍有{totals["needs_review"]:,}条待复核；沿用底稿标签不表示本轮已重新验证其正确性。'],
        ['语义边界', f'{k}是实际使用的类别数，不表示所有类别的语义重合或宽泛、混杂问题都已解决。'],
        ['规则支持', '自动规则支持数沿用既有证据判断，排除撤稿记录；不是人工逐篇验收或准确率估计。'],
        ['文献数据', '全部逐条标签位于results/assignments/的259份Parquet；此工作簿展示完整类别表和汇总统计。'],
        ['索引', f'category_index为0–{k-1}，对应final_labels.npy；质量隔离为-1。CSV字典保存编号与索引。'],
        ['分析入口', '全部当前类别粗统计使用eligible_for_counts；自动规则支持分析使用eligible_for_supported_analysis。'],
        ['撤稿记录', f'{totals["retracted"]:,}条原撤稿标记保留；类别目录的记录数含撤稿，可统计记录数排除撤稿。'],
        ['历史追溯', 'audit/label_crosswalk.csv提供历史标签到当前编号的一一映射；历史候选证据按row_id追溯。'],
        ['文本与版本', '14,156条已清理文本的修订结果保留；原语料和此前版本均保留。'],
    ], columns=['事项', '说明'])
    overview = pd.DataFrame([
        ['当前类别数', k], ['总记录', totals['documents']],
        ['有研究类别的记录', summary['category_records']], ['质量隔离记录', totals['quarantined']],
        ['可作类别粗统计的非撤稿记录', totals['eligible_for_counts']],
        ['获自动规则支持的非撤稿记录', totals['eligible_for_supported_analysis']],
        ['研究主题标签调整（相对最初底稿）', totals['semantic_label_changed']],
        ['全部标签变化含质量隔离（相对最初底稿）', totals['label_changed']],
        ['本次平级编号转换引起的语义改标', 0], ['文本清理记录', totals['text_cleaned']],
        ['待复核记录', totals['needs_review']], ['其中已归类记录待复核', summary['category_records_needing_review']],
        ['其中质量隔离记录待复核', summary['quarantined_records_needing_review']],
        ['保留撤稿标记', totals['retracted']],
    ], columns=['指标', '数量'])
    csv(overview, FLAT / 'results/delivery_overview.csv')
    def chinese(frame):
        return frame.rename(columns=TRANSLATIONS)
    sheets = {
        '阅读说明': notes, '总览': overview,
        '类别目录': chinese(catalog.drop(columns='category_index')),
        '改标流向': chinese(pd.read_csv(FLAT / 'results/label_transitions.csv').fillna('')),
        '决定状态': chinese(pd.read_csv(FLAT / 'results/status_counts.csv')),
        '质量隔离': chinese(quality),
        '文本质量': chinese(pd.read_csv(FLAT / 'results/quality_counts.csv')),
        '来源统计': chinese(pd.read_csv(FLAT / 'results/source_counts.csv')),
        '待复核原因': chinese(pd.read_csv(FLAT / 'results/review_reason_counts.csv')),
    }
    workbook_path = FLAT / f'{k}类平级目录与文献标签.xlsx'
    with pd.ExcelWriter(workbook_path, engine='openpyxl') as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.book[sheet_name]
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='24486C')
            for col in sheet.iter_cols():
                values = [str(cell.value or '') for cell in col[:80]]
                width = min(65, max(14, max(sum(2 if ord(c) > 255 else 1 for c in value) for value in values) + 2))
                sheet.column_dimensions[col[0].column_letter].width = width
                for cell in col:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                        assert not any(term in cell.value for term in ['主类', '下级子类', '父类', '子类'])
                    elif isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0'
            sheet.row_dimensions[1].height = 32
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    assert workbook.sheetnames == list(sheets)
    for sheet_name, frame in sheets.items():
        sheet = workbook[sheet_name]
        assert sheet.max_row == len(frame) + 1 and sheet.max_column == len(frame.columns)
        rows = sheet.iter_rows(values_only=True)
        assert tuple(next(rows)) == tuple(frame.columns)
        for actual, expected in zip(rows, frame.itertuples(index=False, name=None)):
            assert tuple('' if v is None else v for v in actual) == tuple('' if pd.isna(v) else v for v in expected)
    assert workbook['类别目录'].max_row == k + 1
    workbook.close()
    dump(FLAT / 'audit/WORKBOOK_VALIDATION.json', dict(passed=True, created_utc=now(),
        all_cells_read_back=True, categories=k, sheets={key: len(value) for key, value in sheets.items()},
        workbook_sha256=sha(workbook_path)))
    write_documentation(summary, workbook_path.name)
    dump(FLAT / 'PROGRESS.json', dict(stage='complete', categories=k, documents=totals['documents'],
                                   data_validation_passed=True, workbook_validation_passed=True, updated_utc=now()))
    files = []
    for path in sorted(FLAT.rglob('*')):
        if path.is_file() and '__pycache__' not in path.parts and path.name != 'DELIVERY_MANIFEST.json' and path.suffix != '.log':
            files.append(dict(file=str(path.relative_to(FLAT)), bytes=path.stat().st_size, sha256=sha(path)))
    dump(FLAT / 'DELIVERY_MANIFEST.json', dict(created_utc=now(), categories=k, documents=totals['documents'],
        file_count=len(files), files=files, excludes=['DELIVERY_MANIFEST.json', '__pycache__/', '*.log']))
    for item in read(FLAT / 'DELIVERY_MANIFEST.json')['files']:
        assert sha(FLAT / item['file']) == item['sha256']
    print('FLAT_DELIVERY_COMPLETE', k, totals['documents'], 'files', len(files), flush=True)


def write_documentation(summary, workbook_name):
    totals = summary['totals']
    k = summary['category_count']
    readme = f'''# {k}类平级分类版

当前实际使用 **{k} 个类别**，统一编号 **C0001–C{k:04d}**，完整目录在一张表中平级列出。

入口：[中文工作簿]({workbook_name}) · [类别CSV](results/category_dictionary.csv) · [结果说明](REPORT.md) · [字段说明](DATA_DICTIONARY.md)。

共 **{totals['documents']:,} 条记录**：其中 **{summary['category_records']:,} 条各对应一个研究类别**，另 **{totals['quarantined']:,} 条质量隔离**，类别为空，不增加研究类别数。目录仅统计实际归入各类别的记录，每条记录只计一次。

本次统一类别结构与编号，保留已有的正式归类决定。此前相对最初底稿的 **{totals['semantic_label_changed']:,} 条研究主题标签调整**已全部继承；本次换号不计为新的语义改标。

仍有 **{totals['needs_review']:,} 条待复核**。类别数不代表语义已完全无重合，沿用标签也不表示已经重新逐篇验证。原有语义判断基于自动规则与抽样诊断，未人工读完全部记录。

- `results/assignments/part-*.parquet`：259片完整文献标签，含文献ID、题名、来源、当前类别和复核状态。
- `results/category_dictionary.csv`：{k}个实际使用的类别及直接记录数，含论文、专利、政策计数。
- `results/final_labels.npy`：按`row_id`排列的类别整数索引；0–{k-1}对应类别字典，-1表示质量隔离。
- `results/analysis_count_labels.npy`：排除撤稿后的当前类别索引，其他为-1；其中仍可能包含待复核记录。
- `results/analysis_supported_labels.npy`：自动规则支持且非撤稿的类别索引，其他为-1。
- `results/needs_review.npy`、`results/quarantined.npy`：逐条布尔标记。
- `results/pending_review/`、`results/label_changes/`、`results/quality_quarantine/`：待复核、相对底稿标签变化和质量隔离记录。
- `results/text_corrections/`：已完成的文本清理副本；原始语料保留。
- `audit/label_crosswalk.csv`：历史正式标签到当前编号的映射；无记录的候选不进入类别目录。

工作簿展示完整类别表及汇总统计；500余万条逐条标签全部保存在Parquet文件中。

读取一条文献、清理后的正文与当前类别（系统`python3`）：

```sh
python3 {FLAT}/src/read_flat.py --row-id 4006773 --output /tmp/flat_document.json
```

批量读取时把`src`加入Python搜索路径，调用`read_flat.iter_flat_parts()`；只需要题名和标签时传`include_body=False`。读取完整原文需保留原始语料目录`inputs/baseline/data/corpus/`。

结构转换脚本：`src/build_flat.py` → `src/validate_flat.py` → `src/build_delivery.py`。输入正式决定及哈希见`audit/INPUT_REFERENCES.json`。此前版本均保留；历史候选和规则证据可按`row_id`追溯。本版不重新运行语义模型。

全量一致性校验见`VALIDATION.json`，工作簿逐格回读校验见`audit/WORKBOOK_VALIDATION.json`，交付文件哈希见`DELIVERY_MANIFEST.json`。
'''
    (FLAT / 'README.md').write_text(readme)
    report = f'''# {k}类平级结果

交付时间：{now()}。

类别目录现为单一平级列表，共 **{k} 类**，编号 **C0001–C{k:04d}**。每条有研究标签的记录只有一个当前类别；质量隔离独立记录为状态。

| 指标 | 数量 |
|---|---:|
| 当前实际使用的类别 | {k:,} |
| 全部记录 | {totals['documents']:,} |
| 有研究类别的记录 | {summary['category_records']:,} |
| 质量隔离记录 | {totals['quarantined']:,} |
| 可作当前类别粗统计的非撤稿记录 | {totals['eligible_for_counts']:,} |
| 获自动规则支持的非撤稿记录 | {totals['eligible_for_supported_analysis']:,} |
| 仍需复核 | {totals['needs_review']:,} |
| 此前研究标签调整，已继承 | {totals['semantic_label_changed']:,} |
| 已完成文本清理，已继承 | {totals['text_cleaned']:,} |
| 原撤稿标记，已保留 | {totals['retracted']:,} |
| 本次编号转换导致的语义改标 | 0 |

实际类别数由正式标签中有记录的研究类别逐项去重得到。{k}个类别名称均不完全相同，每类均有记录；1个没有实际记录的候选不列入目录。数量按直接归属计算，{summary['category_records']:,}条归类记录与{totals['quarantined']:,}条质量隔离记录之和为{totals['documents']:,}，无重复汇总。

类别表中的“记录数”包含撤稿记录；“可统计记录数”排除撤稿，但可能含待复核记录；“自动规则支持数”沿用既有证据筛选并排除撤稿。三种口径不可混用，规则支持数不是人工验收数量。

本次没有新增语义决定，所有类别含义、文献归属、清理文本和待复核状态均按正式结果保留。历史状态名称只改为平级用语；历史标签映射单独存档。

尚有{totals['needs_review']:,}条待复核，其中{summary['category_records_needing_review']:,}条有当前研究类别、{summary['quarantined_records_needing_review']:,}条处于质量隔离。将目录改成平级不等于消除了宽泛、混杂或语义重合；原有自动处理和抽样诊断也不等于人工逐篇审阅。

已校验全部259片：row_id从0至5,119,003完整连续，无重复或缺失；文献ID、来源、日期和撤稿标记与原始语料逐条一致；新编号与原正式归类逐条对应，Parquet、NPY、类别数量、待复核清单及标签变化清单一致；14,156条清理文本副本哈希与既有结果一致；原正式结果和关键审计文件哈希未变。数据完整性校验不构成语义准确率评估。

使用[中文工作簿]({workbook_name})查看类别目录和统计，使用`results/assignments/`读取全部文献标签。
'''
    (FLAT / 'REPORT.md').write_text(report)
    dictionary = f'''# 文献类别与统计字段

`results/assignments/`的259个分片与原语料行序对应；全局主键是`row_id`，并非`doc_id`。不同来源可能有相同文献编号。

| 字段 | 含义 |
|---|---|
| row_id | 0至5,119,003，全局行号，也是NPY数组下标 |
| doc_id / source / date | 原文献ID、来源、日期，逐条不变 |
| title | 保留已完成清理的题名；原题名可用read_flat读取 |
| category_index | 当前类别的0至{k-1}整数索引；质量隔离为-1 |
| category_id / category_name | 单一当前类别编号和名称；质量隔离时均为空 |
| baseline_category_id / baseline_category_name | 最初底稿类别，以本版统一编号表示；原无研究类别时为空 |
| baseline_review_status | 保留底稿类别的审阅状态 |
| assignment_status | 当前决定状态，中文解释见results/status_counts.csv |
| label_origin | 自动题名规则及语义支持、沿用底稿未新验收、质量隔离三种来源 |
| label_changed | 相对最初底稿标签是否变化，包含进入质量隔离；不将本次换号算作变化 |
| semantic_label_changed | 相对最初底稿的研究主题调整，不含质量隔离 |
| quarantined / quarantine_code / quarantine_reason | 是否质量隔离、Q.*状态编号、隔离原因；不是研究类别 |
| quality_status / initial_quality_status / quality_flags | 当前文本质量、原自动决定时的质量状态和处理标记 |
| text_cleaned | 是否曾做文本清理，完整副本见results/text_corrections/ |
| retracted | 原撤稿标记，保留但排除于两个分析口径 |
| needs_review / review_reason / review_priority | 待复核标记、原因、优先级；1最优先，0不进入本轮待复核清单 |
| priority_cleanup | 是否来自此前优先清理的宽泛、混杂类别 |
| eligible_for_counts | 有当前研究类别且非撤稿，可作粗统计；不表示标签已通过语义验收 |
| eligible_for_supported_analysis | 当前研究标签获既有自动规则支持且非撤稿，不含待复核记录 |
| rule_id | 既有自动规则编号，用于审计；不是本版类别编号 |
| text_evidence_level | 3为题名提供对象与任务证据；2为题名对象加正文任务；0为无充分或其他状态；不是准确率 |
| semantic_support_cosine | 原候选中心的余弦相似度，启发式支持，不是概率 |
| semantic_support_cluster1000 / raw_cluster1000 | 1000版的支持簇和原候选簇编号，仅作证据，不是当前类别 |
| routing_body_characters | 既有判别实际使用的正文字符数，不代表逐篇全文精读 |
| clean_text_nearest1000 | 清理后重新编码的最近候选簇，仅作检索证据 |
| decision_action | 既有自动决定是否保持、撤回、统一标签或作编辑文本质量修正 |

`category_dictionary.csv`只有{k}个实际使用的类别。`documents`是直接归类记录数；`countable_documents`排除撤稿；`supported_documents`是获自动规则支持且非撤稿的记录数；`needs_review_documents`保留待复核数量；`paper_documents`、`patent_documents`、`policy_documents`之和等于该类记录数。其余列分别记录撤稿、清理文本和相对底稿研究标签调整数量。

`quarantine_dictionary.csv`独立统计4种质量状态。Q.*不是额外研究类别，相关记录在`final_labels.npy`中统一为-1，具体原因见逐条Parquet。

NPY数组都按`row_id`对齐：`final_labels.npy`保留当前类别；`analysis_count_labels.npy`只在`eligible_for_counts`为真时保留类别索引；`analysis_supported_labels.npy`只在`eligible_for_supported_analysis`为真时保留类别索引，其余为-1。读取字典时必须使用本目录的`category_dictionary.csv`。

`pending_review/`、`label_changes/`、`quality_quarantine/`与正式标签使用相同字段，分别为对应布尔条件筛选出的记录。它们可能相互交叉，不能直接把记录数相加。

历史候选可能包含无记录或已统一名称的条目，因此未将候选当作额外当前标签。原候选、初稿决定和旧编号在此前正式导出中按`row_id`保存；路径和哈希见`audit/INPUT_REFERENCES.json`。`audit/label_crosswalk.csv`明确区分实际类别、质量隔离和未使用候选。
'''
    (FLAT / 'DATA_DICTIONARY.md').write_text(dictionary)


if __name__ == '__main__':
    main()
