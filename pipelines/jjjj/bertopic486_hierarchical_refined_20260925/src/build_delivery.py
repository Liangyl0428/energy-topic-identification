from refine_common import *
from decision_guards import TITLE_GUARDS,EXCLUDE_GUARDS,ALIASES,NAME_OVERRIDES
from collections import Counter
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment

def percent(x,n):return f'{x/n:.2%}'

def main():
    s=read(BASE/'results/SUMMARY.json');t=s['totals'];v=read(BASE/'VALIDATION.json');assert v['passed']
    parents=pd.read_csv(BASE/'results/main_directory_486_refined.csv').fillna('')
    children=pd.read_csv(BASE/'results/child_directory.csv').fillna('')
    definition_notes={
        'M101.line_fault':'输配电线路的故障检测或定位；不包含电缆材料研究或雷电回击的抽象传输线模型。',
        'M137.parameters':'光伏电池或组件电学模型的参数辨识、估计或提取。',
        'M255.co2_hydrogenation':'CO2催化加氢或甲烷化；产物可以是燃料或化学品，不仅凭名称推断用途。',
        'M252.orc':'有机朗肯循环发电及其热力过程；热源不限定为余热。',
        'M114.hydrate':'二氧化碳水合物的形成、性质及基于水合物的分离；允许CO2混合气体系。',
    }
    for tid,definition in definition_notes.items():children.loc[children.topic_id.eq(tid),'definition']=definition
    children.to_csv(BASE/'results/child_directory.csv',index=False,encoding='utf-8-sig')
    rules=read(BASE/'review/final_taxonomy_rules.json')
    for r in rules['rules']:
        if r['topic_id'] in definition_notes:r['definition']=definition_notes[r['topic_id']]
    dump(BASE/'review/final_taxonomy_rules.json',rules)
    initial_csv=BASE/'review/initial_taxonomy_rules.csv'
    if not initial_csv.exists():initial_csv.write_bytes((BASE/'results/taxonomy_rules.csv').read_bytes())
    rule_rows=[]
    for r in rules['rules']:
        item={k:r.get(k,'') for k in ['rule_id','initial_topic_id','topic_id','parent_id','name','is_child','definition','seed_status','object_pattern','task_pattern','exclude_title_pattern','required_text_pattern','accepted_documents']}
        item['candidate_seed_clusters1000']='|'.join(f'T{x:04d}' for x in r['candidate_seed_clusters1000'])
        item['additional_precision_guards']='src/decision_guards.py'
        rule_rows.append(item)
    pd.DataFrame(rule_rows).to_csv(BASE/'results/taxonomy_rules.csv',index=False,encoding='utf-8-sig')
    enabled=children[children.enabled]
    final=read(BASE/'evidence/final_review_samples.json')
    audit=read(BASE/'review/FINAL_SEMANTIC_REVIEW.json');assert audit['completed']
    packets=read(BASE/'review/parent_candidate_packets.json');candidate_rows=[]
    for p in packets:
        if not p['priority']:continue
        for rank,c in enumerate(p['candidates'],1):
            candidate_rows.append(dict(parent_id=p['parent_id'],parent_name=p['name'],rank=rank,candidate_cluster1000=c['raw1000'],
                intersection_documents=c['documents'],share_of_parent=c['share_of_parent'],keywords=c['keywords'],
                center_row_id=c.get('center',{}).get('row_id'),center_title=c.get('center',{}).get('title',''),
                random_row_id=c.get('random',{}).get('row_id'),random_title=c.get('random',{}).get('title',''),
                status='仅候选；交集和样本不构成整簇认可'))
    candidates=pd.DataFrame(candidate_rows);candidates.to_csv(BASE/'results/priority_1000_candidates.csv',index=False,encoding='utf-8-sig')
    domain=parents.groupby('domain',sort=False)[['original_documents','final_inclusive_documents','accepted_child_documents','original_rows_needs_review']].sum().reset_index()
    domain.to_csv(BASE/'results/domain_counts.csv',index=False,encoding='utf-8-sig')
    quarantine=sum(count for name,count in s['counts']['topic'].items() if name.startswith('Q.'))
    priority_with_children=int((parents.priority & parents.enabled_children.gt(0)).sum())
    priority_adjusted=int((parents.priority & parents.original_rows_label_changed.gt(0)).sum())
    stats=[('总记录',t['documents']),('保留主类',486),('原宽泛或混杂主类',280),('优先处理主类的原记录',t['priority_parent']),
           ('启用子类',s['enabled_child_count']),('实际拥有子类的优先主类',priority_with_children),('有标签或质量调整的优先主类',priority_adjusted),
           ('接受的新标签',t['accepted_changed_labels']),('其中跨主类重分',t['accepted_cross_parent']),
           ('优先主类接受的新标签',t['priority_accepted_changed']),('可用于子类分析的非撤稿记录',t['eligible_for_leaf_analysis']),
           ('质量隔离记录',quarantine),('全部标签变化（包含质量隔离）',t['topic_changed']),('清理文本的记录',t['text_cleaned']),
           ('清理后接受主题标签',t['cleaned_reclassified']),('撤回不充分的自动决定',t['reverted_decisions']),
           ('需继续复核的记录',t['needs_review']),('其中来自优先主类',t['priority_needs_review']),('撤稿标记原样保留',t['retracted'])]
    overview=pd.DataFrame(stats,columns=['指标','数量']);overview.to_csv(BASE/'results/delivery_overview.csv',index=False,encoding='utf-8-sig')
    translations={
      'parent_id':'主类ID','merged_topic_id':'原主类ID','name':'名称','domain':'领域','status':'原审阅状态','priority':'优先清理',
      'original_documents':'原记录数','final_inclusive_documents':'现主类总记录_含子类','final_direct_parent_documents':'仍直接挂主类记录',
      'accepted_child_documents':'已分入子类记录','enabled_children':'启用子类数','original_rows_label_changed':'原记录标签变化数',
      'original_rows_parent_changed':'原记录迁出或隔离数','original_rows_needs_review':'原记录待复核数',
      'original_rows_quality_quarantine':'原记录质量隔离数','original_rows_to_children':'原记录分入子类数',
      'original_rows_accepted':'原记录获规则支持数','original_rows_still_at_parent':'原记录仍挂原主类数','original_rows_text_cleaned':'原记录文本清理数',
      'local_child_rules':'拟议本地子类规则数','candidate_rules':'可检索规则数','topic_id':'子类ID','parent_name':'所属主类',
      'documents':'记录数','eligible_leaf_documents':'子类分析可用数','enabled':'已启用','rule_id':'规则ID','definition':'定义与边界',
      'candidate_clusters1000':'1000版候选簇','seed_status':'候选检索依据','precision_note':'精度防护说明','review_status':'验收范围',
      'assignment_status':'决定状态','value':'状态或来源','reason':'原因','decisions':'撤回决定数','row_id':'全局行ID','doc_id':'文献ID',
      'source':'来源','title':'题名','old_parent_id':'原主类','final_topic_id':'最终标签ID','final_topic_label':'最终标签',
      'semantic_support_cosine':'向量相似度_非概率','text_evidence_level':'文本证据级别','body_excerpt':'正文节选','url':'原链接'}
    def cn(df):return df.rename(columns=translations)
    pcols=['parent_id','name','domain','status','priority','original_documents','final_inclusive_documents','final_direct_parent_documents',
           'accepted_child_documents','enabled_children','original_rows_label_changed','original_rows_needs_review','original_rows_quality_quarantine','original_rows_text_cleaned']
    ccols=['topic_id','parent_id','parent_name','name','documents','eligible_leaf_documents','definition','candidate_clusters1000','seed_status','precision_note','review_status']
    notes=pd.DataFrame([
        ['统计口径','486个主类是目录层；子类属于主类。主类含子类数和子类数不能相加。每条记录只有一个final_topic_id。'],
        ['实际改标','新旧标签已逐条写入259份Parquet；NPY标签按原row_id对齐。质量隔离与接受的研究主题调整分开统计。'],
        ['自动处理范围','对全部记录逐条做质量检查和候选判别；对象/任务题名证据和向量支持共同决定接受，非人工逐篇读完。'],
        ['待复核','宽泛/混杂类证据不足、对象任务冲突、模板清理后主题不明等，保留原主类或进入质量桶。'],
        ['分析入口','子类分析使用eligible_for_leaf_analysis；主类粗统计使用eligible_for_parent_counts。保留父类不表示已验证纯度。'],
        ['相似度','余弦相似度只作启发式语义支持，不是概率、置信度或正确率。'],
        ['1000版候选','1000簇作为候选检索，禁止整簇直接转为子类。29条规则只有名称向量近邻检索起点；仍须逐条题名证据。'],
        ['语义抽查','诊断与验收样本按主题分层；样本用于发现错误和修订规则，不代表总体准确率。'],
        ['已知局限','中英文明确题名规则优先；其他语言、笼统题名、跨任务文献、正文才说明主任务的记录大量待复核。'],
        ['文本处理','原题名正文未覆盖；精确模板移除及派生正文另存，可通过read_refined.py读取文献和新标签。'],
        ['撤稿','保留原retracted字段；已撤稿记录不进入子类/接受主题分析。'],
        ['目录统一','产沼气消化工艺归M128；沼气提纯保留M274.biogas_upgrading。餐厨/污泥与粪便子类交叉时不强分。']
    ],columns=['事项','说明'])
    ev=pd.DataFrame(final['records']);ecols=[c for c in ['row_id','doc_id','source','title','old_parent_id','sampled_topic_id','final_topic_id','final_topic_label','assignment_status','needs_review','review_reason','rule_id','semantic_support_cosine','text_evidence_level','body_excerpt','url'] if c in ev]
    ev['body_excerpt']=ev.body_excerpt.str.slice(0,600)
    sheets={
        '阅读说明':notes,'总览':overview,'486主目录':cn(parents[pcols]),'启用子类':cn(enabled[ccols]),
        '280类处理结果':cn(parents.loc[parents.priority,pcols+['original_rows_to_children','original_rows_accepted','original_rows_still_at_parent','candidate_rules']]),
        '1000版候选证据':candidates,'改标状态':cn(pd.read_csv(BASE/'results/status_counts.csv')),
        '质量统计':cn(pd.read_csv(BASE/'results/quality_counts.csv')),'文本清理标记':cn(pd.read_csv(BASE/'results/flags_counts.csv')),
        '来源统计':cn(pd.read_csv(BASE/'results/source_counts.csv')),'领域统计':cn(domain),
        '误判撤回原因':cn(pd.read_csv(BASE/'results/guard_reversions_summary.csv')),'最终抽查样本':cn(ev[ecols]),
        '初轮诊断问题':pd.DataFrame(read(BASE/'review/DIAGNOSTIC_FINDINGS.json')['detected_cases']),
        '第二轮诊断问题':pd.DataFrame(read(BASE/'review/ROUND2_DIAGNOSTIC_FINDINGS.json')['detected_cases']),
        '补充样本诊断问题':pd.DataFrame(read(BASE/'review/ROUND3_DIAGNOSTIC_FINDINGS.json')['detected_cases']),
        '目录同义统一':pd.DataFrame([dict(原拟议子类=k,统一归属=v) for k,v in ALIASES.items()]),
        '语义抽查最终决定':cn(pd.read_csv(BASE/'review/SEMANTIC_REVIEW_OUTCOMES.csv')),
    }
    path=BASE/'486主目录清理细分与文献新标签.xlsx'
    with pd.ExcelWriter(path,engine='openpyxl') as writer:
        for name,df in sheets.items():df.to_excel(writer,sheet_name=name,index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24486C')
            for col in ws.columns:
                letter=col[0].column_letter
                ws.column_dimensions[letter].width=min(65,max(15,max(len(str(c.value or '')) for c in list(col)[:60])*.8))
                for c in col:
                    c.alignment=Alignment(vertical='top',wrap_text=True)
                    if isinstance(c.value,str):c.data_type='s'
    wb=load_workbook(path,read_only=True,data_only=False)
    assert all(wb[name].max_row==len(df)+1 for name,df in sheets.items())
    assert wb['486主目录'].max_row==487 and wb['启用子类'].max_row==len(enabled)+1
    wb.close();dump(BASE/'review/WORKBOOK_VALIDATION.json',dict(passed=True,sheets={k:len(df) for k,df in sheets.items()},sha256=sha(path)))
    statuslines='\n'.join(f'| {k} | {v:,} |' for k,v in s['counts']['status'].items())
    qualitylines='\n'.join(f'| {k} | {v:,} |' for k,v in s['counts']['quality'].items())
    top=parents.loc[parents.priority].sort_values('original_rows_label_changed',ascending=False).head(12)
    toplines='\n'.join(f'| {r.parent_id} {r["name"]} | {r.original_documents:,} | {r.original_rows_label_changed:,} | {r.original_rows_needs_review:,} |' for _,r in top.iterrows())
    report=f'''# 486主目录清理、细分与逐条标签修订

完成时间：{now()}。新版本独立存放于 `{BASE.name}`，保留原486、1000、990版本。

共逐条处理 **{t['documents']:,} 条记录**（论文、专利、政策），保留 **486 个主类**，启用 **{s['enabled_child_count']} 个子类**。
接受 **{t['accepted_changed_labels']:,} 条研究主题标签调整**，其中 **{t['accepted_cross_parent']:,} 条跨主类重分**；另有 **{quarantine:,} 条进入文本质量桶**。
清理 **{t['text_cleaned']:,} 条记录的模板或题名前缀**；其中 **{t['cleaned_reclassified']:,} 条**获得题名及向量共同支持的主题标签，其余按质量和证据状态保留复核。

本版是全量自动逐条处理并经分层抽样诊断的可用版本，**不是人工读完全部文献，也不代表所有宽泛、混杂问题已消除**。
仍有 **{t['needs_review']:,} 条（{percent(t['needs_review'],t['documents'])}）需要复核**。这些记录完整保留，未用不确定的子类强行覆盖。

## 优先主类的处理情况

原207个宽泛类与73个混杂类，共280类、{t['priority_parent']:,}条记录全部进入质量检查与候选检索。
其中{priority_with_children}个主类实际有启用子类；{priority_adjusted}个主类发生标签或质量调整。优先记录中{t['priority_accepted_changed']:,}条获得新主题标签，{t['priority_needs_review']:,}条仍需复核。
未实际细分的主类保留原目录，并提供1000版交集候选、关键词和代表题名，见 `results/priority_1000_candidates.csv`。

| 主类 | 原记录 | 标签变化（含质量隔离） | 原记录待复核 |
|---|---:|---:|---:|
{toplines}

## 子类与语义边界

主类ID和名称沿用486底稿，子类用 `Mxxx.key` 表示，每个子类唯一归属一个主类。主类下面尚未明确细分的记录仍挂主类。
**486主类与子类不能相加宣称为互斥类别总数，也不能将父子数量重复相加。**

- 1000版只提供候选检索和向量佐证，不将整个簇直接作为已确认子类。
- 307条规则中有29条以名称向量近邻补充检索起点；这不等于认可对应整簇。具体使用量与依据见最终规则文件。
- 产沼气的厌氧消化工艺统一归M128；沼气提纯作为独立任务保留。餐厨垃圾/污泥与粪便/农业残余共同消化时进入复核。
- 电网建设政策、场效应晶体管、电力碳计量与调度、智能电表这4个拟议子类与已有主类同义，统一回原主类，不重复增加节点。
- 氮化碳光催化制氢、钴基析氧等材料特化归已有共同反应主题，避免复制同义标签。
- 贵金属或材料范围不明的氧还原研究不能迁入“非贵金属氧还原”主类。
- 无粘结剂、包覆与涂布、孤岛调度与动态控制、供氢重整器与燃料电池本体、制造与检测等易混任务分别设防护。

## 改标方法与证据

逐条检查质量后，使用原父类、1000版第一/第二候选及规则检索获得候选。只有题名符合规则所需的对象及任务、并达到启发式向量门槛时才接受主题调整；少数以专门技术对象定义的类没有单独的任务词要求，具体规则公开保存。
同主类余弦门槛为0.60，跨主类为0.64；**相似度不是概率，也不能据此声称正确率**。摘要背景提及、多个任务、迁移可能丢失具体对象等情况保留复核。
题名完整处理；通常正文最多取16000字符、政策正文最多4000字符用于候选判别，原文完整保留。正文不是本轮逐篇全文精读。

初轮独立诊断样本291条，发现并撤回33条不充分决定；据此及后续诊断完善精度防护。本轮全量共撤回{t['reverted_decisions']:,}个不充分自动决定，逐条理由可查。
第二轮604条分层诊断样本中有{len(read(BASE/'review/ROUND2_DIAGNOSTIC_FINDINGS.json')['detected_cases'])}条触发进一步撤回或质量隔离。修改后对新增抽样记录继续检查；这些是适应性诊断与回归检查，不是保留测试集准确率评估。
补充阅读86条新样本，再撤回8条不充分决定，并将通用防护应用于全量。最终固定的604条诊断样本同时展示最终接受与退回复核结果，完整保留修订过程。
最终分层抽查与后续防护核查范围、实际阅读记录和未解决问题见 `review/FINAL_SEMANTIC_REVIEW.json`。跨轮累计阅读{audit['unique_titles_reviewed']}条不同题名，另对{audit['unique_records_with_body_excerpt_reviewed']}条查看了正文节选。
样本按主题分层，偏向发现各类问题，不能外推全体记录准确率。

## 状态与质量数量

| 决定状态 | 记录数 |
|---|---:|
{statuslines}

| 文本质量状态 | 记录数 |
|---|---:|
{qualitylines}

已撤稿的{t['retracted']:,}条保留原标记，并从接受主题及子类分析口径中排除。原不可用记录、信息不足、封面索引勘误和未解模板分别保存在Q.*桶，未删除记录。

## 使用文件

- `486主目录清理细分与文献新标签.xlsx`：主目录、启用子类、280类结果、候选证据、诊断与最终样本。
- `results/assignments/part-*.parquet`：全部文献逐条新旧标签及审计字段，共259片。
- `results/final_labels.npy`：按原row_id排列的最终标签整数索引；对照 `results/topic_dictionary.csv`。
- `results/final_parent_labels.npy`：486主类索引，质量桶为-1；主类顺序同原486目录。
- `results/analysis_leaf_labels.npy`：可用于子类分析的索引，其他记录为-1；`needs_review.npy`是逐条复核标记。
- `results/label_changes/`、`results/pending_review/`、`results/decision_reversions/`：标签变化、待复核及撤回决定的逐条清单。
- `results/text_corrections/`：派生清理正文、题名及原文哈希；`models/cleaned_*.npz`保存重新编码向量。
- `results/initial_assignments/`：防护前结果，仅用于审计和复现，分析应使用正式assignments。
- `src/read_refined.py`：直接读取原文献、清理后正文及最终标签；无需复制原始语料。

子类分析使用 `eligible_for_leaf_analysis`；仅统计有规则支持的主题使用 `eligible_for_accepted_topic_analysis`；主类粗统计用 `eligible_for_parent_counts`，其中可能包含待复核记录。
兼容字段 `eligible_for_topic_counts`等同主类粗统计标记，**不能当作细分标签已经验收的标记**。

## 校验与边界

校验通过：5,119,004行完整且row_id无重复/缺失，文献ID、来源、日期、撤稿标记不变；新旧索引与Parquet一致；子类引用唯一有效；父类/子类/质量桶数量守恒；复核清单一致；{v['correction_records_verified']:,}条清理记录和{v['cleaned_embeddings_verified']:,}条重算向量校验通过；旧版本923个交付及语料文件哈希核对通过。Excel已回读核验。

完整性校验不能替代语义准确性。当前中英文明确题名的覆盖更好；其他语言、笼统题名、复杂交叉研究与仅正文能说明任务的记录仍大量待复核。
原486目录只是底稿，未变更的主类标签也不能据此认为已经重新验证纯度。本版保留这些限制，避免将自动处理包装成全部问题已解决。
'''
    (BASE/'REPORT.md').write_text(report)
    readme=f'''# 486主目录修订版

主入口：[结果说明](REPORT.md)、[中文工作簿](486主目录清理细分与文献新标签.xlsx)。

处理记录：{t['documents']:,}；保留486主类，启用{s['enabled_child_count']}子类；接受新研究主题标签{t['accepted_changed_labels']:,}条；仍需复核{t['needs_review']:,}条。

正式标签在`results/assignments/`。它们已经包含文献ID、题名、来源和逐条新旧标签。原文献正文以只读方式引用旧486版的`data/corpus/`。

读取一条文献及新标签（使用系统python3）：

```sh
python3 {BASE}/src/read_refined.py --row-id 4006773 --output /tmp/refined_document.json
```

批量读取：将`src`加入Python搜索路径，调用`read_refined.iter_refined_parts()`。每个批次提供原文、派生清理正文、原标签、最终标签和复核标记。

NPY标签按原始row_id逐条排列，字典见`results/topic_dictionary.csv`。不要把原cluster500或cluster1000的编号当作本版标签索引。

子类统计筛选`eligible_for_leaf_analysis`。父子计数不要相加。原始目录和防护前结果不作为正式新标签。

复现顺序：`prepare.py` → `prepare_models.py` → `complete_retrieval.py` → `run_all.py` → `finalize.py` → `sample_review.py`及语义诊断 → `validate_inputs.py` → `validate_output.py` → `build_delivery.py`。
诊断读样本及完善防护是审阅步骤，不能由脚本自动替代。复现完整首轮时应在新的输出目录执行，避免与已交付结果混写。

输入依赖与配置见`INPUT_SNAPSHOT.json`、`RUN_CONFIG.json`；完整性结果见`VALIDATION.json`，交付哈希见`DELIVERY_MANIFEST.json`。
'''
    (BASE/'README.md').write_text(readme)
    sourcehash={str(p.relative_to(BASE)):sha(p) for p in sorted((BASE/'src').glob('*.py'))}
    dump(BASE/'RUN_CONFIG.json',dict(version='486-hierarchy-20260925-v1',created_utc=now(),base_directory=str(BASE),old486=str(OLD),
        candidate1000=str(NEW),corpus_reference=str(OLD/'data/INPUT_MANIFEST.json'),total_documents=t['documents'],
        main_parents=486,rules=307,main_priority_classes=280,minimum_title_evidence_level=3,within_parent_cosine=.60,cross_parent_cosine=.64,
        routing_body_max_characters=16000,policy_body_max_characters=4000,encoder='local multilingual MiniLM, ONNX, 384 dimensions',
        source_hashes=sourcehash,canonical_aliases=ALIASES,name_overrides=NAME_OVERRIDES,
        limitations=['title-focused automated labeling plus stratified diagnostic review','similarity is not calibrated accuracy','unresolved records retain baseline parent and review flag']))
    progress('complete',documents=t['documents'],parents=486,enabled_children=s['enabled_child_count'],accepted_changed_labels=t['accepted_changed_labels'],
        text_cleaned=t['text_cleaned'],needs_review=t['needs_review'],validation_passed=True)
    files=[]
    for p in sorted(BASE.rglob('*')):
        if not p.is_file() or p.name=='DELIVERY_MANIFEST.json' or '__pycache__' in p.parts or 'logs' in p.parts:continue
        files.append(dict(path=str(p.relative_to(BASE)),bytes=p.stat().st_size,sha256=sha(p)))
    dump(BASE/'DELIVERY_MANIFEST.json',dict(created_utc=now(),version='486-hierarchy-20260925-v1',files=files,
        original_corpus_referenced_not_copied=True,official_assignments='results/assignments/',pre_guard_results_are_audit_only='results/initial_assignments/'))
    print('DELIVERY_COMPLETE',path,flush=True)

if __name__=='__main__':main()
