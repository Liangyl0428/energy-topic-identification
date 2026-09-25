"""Independently validate the delivered labels, then write the review workbook and report."""
from merge_common import *
from semantic_mapping import make_mapping, map_labels
from candidate_review import validate_candidate_decisions
import subprocess
import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from scipy import sparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BACKUP=BASE/'history/v1_998_before_synonym_merge'

def export_workbook(topics,mapping,merges,decisions,summary):
    workbook=Workbook()
    workbook.remove(workbook.active)
    notes=pd.DataFrame([
        ['版本',summary['version']],['原始类别数',summary['raw_topics']],
        ['合并后类别数',summary['final_topics']],['本轮新增合并组',summary['new_merge_groups']],
        ['完整文档记录数',summary['documents']],['已分类记录数',summary['assigned_documents']],
        ['未分类不可用文本数',summary['unusable_documents']],
        ['本轮统一名称的文档数',summary['documents_with_topic_name_changed_vs_previous']],
        ['标签分片','results/assignments/*.parquet（全部259份）'],
        ['当前主题字段','final_topic_id、final_topic_index、final_topic_label'],
        ['编号说明','S编号尾数为原类锚点；不等于连续final_topic_index。索引重排不代表语义变更。'],
        ['历史字段','new_cluster1000/new_topic_id、old_cluster500/old_merged486保留为追溯字段'],
        ['备份','history/v1_998_before_synonym_merge'],
        ['复核范围','全1000类清单筛查；本轮150类、1022条不同文献题名，其中272条新增随机样本。部分正文片段核对。'],
        ['范围说明','整类合并，不重新聚类或单篇纠错；未并入只有上下位或相关关系的类。'],
    ],columns=['项目','说明'])
    frames=[('使用说明',notes),('合并总览',merges),('990类目录',topics),('1000类映射',mapping),
        ('保留分开理由',pd.DataFrame([{'decision_id':g['decision_id'],
            'raw_topics':'|'.join(f'T{i:04d}' for i in g['members']),
            'decision':g['decision'],'reason':g['reason']} for g in decisions['keep_or_defer']]))]
    for name,frame in frames:
        sheet=workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for row in frame.fillna('').itertuples(index=False,name=None):
            sheet.append([v.item() if isinstance(v,np.generic) else v for v in row])
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24476B')
        for col in sheet.columns:
            header=str(col[0].value)
            sheet.column_dimensions[col[0].column_letter].width=65 if any(x in header for x in ['reason','说明','label','keywords','limits']) else 23
            for cell in col[1:]:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        sheet.sheet_view.showGridLines=False
    path=BASE/'同义类别合并与标签映射.xlsx'
    workbook.save(path)
    workbook.close()
    check=load_workbook(path,read_only=True,data_only=True)
    for name,frame in frames:
        assert check[name].max_row==len(frame)+1
        assert next(check[name].values)==tuple(frame.columns)
    check.close()
    return path

def main():
    checks={}
    decisions=read(BASE/'review/decisions.json')
    summary=read(BASE/'results/SUMMARY.json')
    assert summary['version']==decisions['version']=='semantic-1000-20260925-v2-synonyms'
    anchor,lookup,anchors=make_mapping(decisions['merge_groups'])
    mapping=pd.read_csv(BASE/'results/raw_to_final.csv').fillna('')
    topics=pd.read_csv(BASE/'results/final_topics.csv').fillna('')
    merges=pd.read_csv(BASE/'results/merge_summary.csv').fillna('')
    prior=pd.read_csv(BACKUP/'results/raw_to_final.csv').fillna('')
    k=len(anchors)
    assert k==summary['final_topics']==990
    assert mapping.raw_cluster1000.tolist()==list(range(1000))
    np.testing.assert_array_equal(mapping.final_topic_index,lookup)
    np.testing.assert_array_equal(mapping.final_topic_id,[f'S{i:04d}' for i in anchor])
    assert topics.final_topic_index.tolist()==list(range(k))
    assert topics.final_topic_id.is_unique and topics.final_topic_label.is_unique
    groupby={i:g for g in decisions['merge_groups'] for i in g['members']}
    for i,row in mapping.iterrows():
        expected=groupby[i]['canonical_name'] if i in groupby else row.raw_automatic_label
        assert row.final_topic_label==expected
        assert bool(row.semantic_merge_applied)==(i in groupby)
    for row in topics.itertuples():
        members=np.flatnonzero(lookup==row.final_topic_index)
        assert row.raw_topic_members=='|'.join(f'T{i:04d}' for i in members)
        assert row.raw_topic_count==len(members)
        assert row.final_topic_label==mapping.iloc[members[0]].final_topic_label
    for g in decisions['keep_or_defer']:
        assert len(set(anchor[g['members']]))==len(g['members'])
    checks['complete_disjoint_mapping_and_kept_pairs']=True
    checks['topic_names_ids_members_and_indices_match']=True
    assert validate_candidate_decisions(decisions)==summary['candidate_pairs']
    checks['all_explicit_review_pairs_including_cross_language_are_exported']=True

    # Frozen raw inputs and the entire pre-edit snapshot must be byte-identical.
    baseline=read(BASE/'review/INPUT_BASELINE.json')
    for root,files in [(RAW,baseline['files']),(BACKUP,read(BACKUP/'BACKUP_MANIFEST.json')['files'])]:
        for item in files:
            p=root/item.get('path',item.get('file'))
            assert p.stat().st_size==item['bytes'] and sha(p)==item['sha256'],str(p)
    assert sha(RAW/'DELIVERY_MANIFEST.json')==baseline['delivery_manifest_sha256']
    checks['all_frozen_raw_input_hashes_unchanged']=True
    checks['complete_previous_998_topic_snapshot_preserved']=True

    raw=np.load(RAW/'results/raw_labels.npy',mmap_mode='r')
    final=np.load(BASE/'results/final_labels.npy',mmap_mode='r')
    np.testing.assert_array_equal(final,map_labels(raw,lookup))
    np.testing.assert_array_equal(np.load(BASE/'models/raw_to_final_index.npy'),lookup)
    np.testing.assert_array_equal(np.load(BASE/'models/raw_to_final_anchor.npy'),anchor)
    manifest=read(BASE/'results/ASSIGNMENT_MANIFEST.json')
    assert manifest['version']==summary['version'] and len(manifest['parts'])==259
    counts=np.zeros((k,3),dtype=np.int64)
    cursor=0
    id_changes=0
    name_changes=0
    source_codes={'paper':0,'patent':1,'policy':2}
    for p in manifest['parts']:
        path=BASE/p['file']
        assert sha(path)==p['sha256']
        table=pq.read_table(path)
        old_table=pq.read_table(BACKUP/p['file'])
        df=table.to_pandas()
        ids=df.row_id.to_numpy()
        np.testing.assert_array_equal(ids,np.arange(cursor,cursor+len(df)))
        assert p['rows']==len(df) and p['first_row']==ids[0] and p['last_row']==ids[-1]
        y=df.new_cluster1000.to_numpy();good=y>=0
        np.testing.assert_array_equal(y,raw[ids])
        np.testing.assert_array_equal(df.final_topic_index,final[ids])
        assert set(df.source)<=set(source_codes)
        # Direct column lookup independently verifies the serialized string labels.
        for col,fallback in [('final_topic_id','UNUSABLE'),('final_topic_label','不可用文本（未分类）'),
                ('final_label_status','不可用文本；保留未分类'),('merge_decision_id','')]:
            target=np.full(len(y),fallback,dtype=object)
            target[good]=mapping[col].to_numpy()[y[good]]
            np.testing.assert_array_equal(df[col],target)
        flags=np.zeros(len(y),dtype=bool);flags[good]=mapping.semantic_merge_applied.to_numpy()[y[good]]
        np.testing.assert_array_equal(df.semantic_merge_applied,flags)
        secondary=df.secondary_candidate1000.to_numpy();valid=(secondary>=0)
        target=np.full(len(y),'UNUSABLE',dtype=object)
        target[valid]=mapping.final_topic_id.to_numpy()[secondary[valid]]
        np.testing.assert_array_equal(df.final_secondary_topic_id,target)
        np.testing.assert_array_equal(df.secondary_maps_to_primary,good & valid & (target==df.final_topic_id.to_numpy()))
        old=old_table.to_pandas()
        raw_columns=[c for c in df if not c.startswith('final_') and c not in ['semantic_merge_applied','merge_decision_id','secondary_maps_to_primary']]
        assert df[raw_columns].equals(old[raw_columns]),f'original document fields changed: {path}'
        id_changes+=int((df.final_topic_id!=old.final_topic_id).sum())
        name_changes+=int((df.final_topic_label!=old.final_topic_label).sum())
        sc=df.source.map(source_codes).to_numpy()
        counts+=np.bincount(final[ids][good].astype(np.int64)*3+sc[good],minlength=k*3).reshape(k,3)
        cursor+=len(df)
        if cursor%1000000==0:
            print('independent label verification:',cursor,flush=True)
    assert cursor==len(raw)==manifest['documents']==summary['documents']==5119004
    assert int((final==-2).sum())==summary['unusable_documents']==2391
    np.testing.assert_array_equal(counts,topics[['papers','patents','policies']])
    np.testing.assert_array_equal(counts.sum(1),topics.documents)
    np.testing.assert_array_equal(counts,np.load(BASE/'results/source_counts.npy'))
    assert id_changes==summary['documents_with_topic_id_changed_vs_previous']
    assert name_changes==summary['documents_with_topic_name_changed_vs_previous']
    checks['all_259_shards_and_5119004_rows_read_back']=True
    checks['all_document_id_source_and_original_columns_preserved']=True
    checks['all_primary_secondary_names_ids_statuses_and_flags_match_mapping']=True
    checks['unusable_records_and_all_topic_source_counts_exact']=True
    checks['reported_changes_match_previous_version_row_by_row']=True

    original_counts=sparse.load_npz(RAW/'models/class_word_counts.npz').tocsr()
    reduction=sparse.csr_matrix((np.ones(1000,dtype=np.int64),(lookup,np.arange(1000))),shape=(k,1000))
    merged_counts=sparse.load_npz(BASE/'models/class_word_counts.npz').tocsr()
    assert ((reduction@original_counts)-merged_counts).nnz==0
    ctf=sparse.load_npz(BASE/'models/ctfidf.npz')
    assert ctf.shape==merged_counts.shape and np.isfinite(ctf.data).all()
    checks['word_counts_exactly_aggregated_by_final_topic']=True
    checks['ctfidf_shape_and_values_valid']=True
    result=subprocess.run([sys.executable,'-m','unittest','discover','-s',str(BASE/'src'),'-p','test_semantic_mapping.py','-v'],capture_output=True,text=True)
    (BASE/'mapping_tests.log').write_text(result.stdout+result.stderr)
    assert result.returncode==0,result.stderr
    checks['mapping_safety_unit_tests_passed']=True

    workbook=export_workbook(topics,mapping,merges,decisions,summary)
    checks['excel_sheets_read_back_and_counts_verified']=True
    new=merges[merges.decision_id.str.startswith('SYN')]
    lines=['# 同义类别合并结果','',
        f'原始 **1000 类归并为 {k} 类**，累计确认 {len(merges)} 组同义合并。本轮在原有998类版本上新增8组合并，统一了 {name_changes:,} 条文档记录的主题名称，其中 {id_changes:,} 条的稳定主题编号发生归并。',
        '',f'已写回并逐行核对全部 **{cursor:,}** 条记录（论文、专利、政策），共259个Parquet分片；其中 {summary["assigned_documents"]:,} 条有分类，{summary["unusable_documents"]:,} 条不可用文本保持未分类。',
        '', '本轮新增合并：','', '| 原始 topic | 统一标签 | 涉及记录 |','|---|---|---:|']
    for r in new.itertuples():
        lines.append(f'| {r.raw_topics.replace("|", " + ")} | {r.final_topic_label} | {r.documents:,} |')
    lines.extend(['','原有两组合并继续保留：T0142 + T0177（并网光伏系统）、T0393 + T0498（OLED）。',
        '', '## 文件与使用', '',
        '- [同义类别合并与标签映射.xlsx](同义类别合并与标签映射.xlsx)：合并总览、990类目录、1000类映射和保留分开理由。',
        '- [results/final_topics.csv](results/final_topics.csv)：更新后的类别、关键词及来源计数。',
        '- [results/raw_to_final.csv](results/raw_to_final.csv)：原始1000类到最终990类的一对一/多对一映射。',
        '- [results/merge_summary.csv](results/merge_summary.csv)：每组合并的依据及文献证据行号。',
        '- `results/assignments/*.parquet`：所有文档的更新标签；`results/final_labels.npy`按原row_id排列。',
        '- [VALIDATION.json](VALIDATION.json)：全量校验；[review/synonym_review_v2.json](review/synonym_review_v2.json)：本轮语义决策。',
        '', '请使用 `final_topic_id`、`final_topic_index` 和 `final_topic_label`。`new_cluster1000/new_topic_id` 与旧500/486类字段保留作追溯。S编号的数字是原始类锚点，不能当成连续索引；连续索引重排不表示未合并类的含义变化。',
        '', '原始1000类输入文件经逐文件SHA-256校验保持不变；旧998类结果完整保存在 `history/v1_998_before_synonym_merge/`。现有模型通过 `src/semantic_mapping.py` 将冻结1000类预测映射到990类，词频和c-TF-IDF按最终类重新聚合；原始置信诊断量没有被当成新的合并类概率。',
        '', '## 合并依据与范围', '',
        '本轮覆盖1000类的主题词清单，对150个候选类读取了20个关键词及中心、随机、边界题名，共1022条不同记录，其中272条是排除旧样本后新抽取的随机记录；部分文献核对正文片段。实际阅读范围保存在 `review/ACTUALLY_INSPECTED_v2.json`。',
        '', '只有对象、主要任务和分类层级相同的同义类才合并。微电网调度/孤岛控制、不同电池材料体系、DSSC光阳极/对电极/染料、ORR/HER等仍分开。相似度只用于寻找候选；未按阈值批量合并。',
        '', '这是整类语义归并，不是对511万条记录逐篇重新判题；原聚类中的边界噪声仍保留，语义复核没有独立专家金标准。未调整TRL/CRL或其他成熟度结果。',
        '', '## 复现', '', '在项目根目录执行：', '', '```sh',
        'python3 jjjj/bertopic1000_semantic_merged_20260925/src/assemble_synonym_decisions_v2.py',
        'python3 jjjj/bertopic1000_semantic_merged_20260925/src/apply_merges.py',
        'python3 jjjj/bertopic1000_semantic_merged_20260925/src/finalize_synonyms_v2.py',
        '```', '', '上述命令使用已保存的语义决策和证据，不会重新训练原模型。'])
    (BASE/'REPORT.md').write_text('\n'.join(lines)+'\n')
    (BASE/'README.md').write_text('\n'.join([
        '# BERTopic同义类别合并（990类）','',
        '原始1000类累计合并10组，得到990类；本轮新增8组。全部5,119,004条文档记录的最终标签已更新并校验。','',
        '查看 [结果报告](REPORT.md) 或 [合并及标签映射工作簿](同义类别合并与标签映射.xlsx)。','',
        '当前类别表：`results/final_topics.csv`。原始类映射：`results/raw_to_final.csv`。文档标签：`results/assignments/*.parquet`，使用 `final_topic_id/final_topic_index/final_topic_label`。','',
        '原始topic字段继续保留；S编号后缀不是连续索引。原998类版本备份在 `history/v1_998_before_synonym_merge/`。完成状态与完整校验以 `VALIDATION.json` 为准。','']) )
    summary['validation_passed']=True
    summary['report']='REPORT.md'
    summary['workbook']=workbook.name
    dump(BASE/'results/SUMMARY.json',summary)
    dump(BASE/'VALIDATION.json',dict(passed=True,stage='complete',version=summary['version'],created_utc=now(),
        checks=checks,raw_topics=1000,final_topics=k,documents=cursor,assignment_parts=259,
        updated_topic_names_vs_previous=name_changes,changed_stable_topic_ids_vs_previous=id_changes,
        note='数据完整性与映射一致性校验通过；不代表逐篇语义准确率或独立专家验收。'))
    dump(BASE/'PROGRESS.json',dict(stage='complete',updated_utc=now(),raw_topics=1000,final_topics=k,documents=cursor,validation_passed=True))
    files=[]
    for p in sorted(BASE.rglob('*')):
        if not p.is_file() or 'history' in p.relative_to(BASE).parts or '__pycache__' in p.parts or p.name=='DELIVERY_MANIFEST.json' or p.suffix in ['.tmp','.log']:
            continue
        files.append(dict(path=str(p.relative_to(BASE)),bytes=p.stat().st_size,sha256=sha(p)))
    dump(BASE/'DELIVERY_MANIFEST.json',dict(created_utc=now(),version=summary['version'],files=files))
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':
    try:
        main()
    except Exception as exc:
        dump(BASE/'VALIDATION.json',dict(passed=False,stage='final_verification_failed',created_utc=now(),error=str(exc)))
        raise
