"""Verify source immutability/model mapping, then publish local review artifacts."""
from merge_common import *
from semantic_mapping import SemanticTopicModel
import numpy as np,pandas as pd,pyarrow.parquet as pq
from scipy import sparse
from threadpoolctl import threadpool_limits
import re

def excel_safe(v):
    if isinstance(v,str):
        v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',' ',v)[:32000]
        # Preserve textual excerpts as text, including formula-like prefixes.
        if v.startswith(('=','+','-','@')):v="'"+v
    return v

def main():
    threadpool_limits(4)
    s=read(BASE/'results/SUMMARY.json');dec=read(BASE/'review/decisions.json')
    validation=read(BASE/'VALIDATION.json');checks=validation['checks']
    baseline=read(BASE/'review/INPUT_BASELINE.json');unchanged=[]
    for r in baseline['files']:
        assert sha(RAW/r['path'])==r['sha256'],f'raw snapshot changed: {r["path"]}'
        unchanged.append(r['path'])
    assert sha(RAW/'DELIVERY_MANIFEST.json')==baseline['delivery_manifest_sha256']
    checks['raw_snapshot_files_unchanged']=len(unchanged)
    checks['raw_delivery_manifest_unchanged']=True
    print('raw snapshot unchanged:',len(unchanged),'files',flush=True)
    # Verify the standard transformer agrees with the exported merged topic words.
    sys.path[:0]=[str(RAW/'src'),str(CORPUS/'src')]
    from experiment_components import FiniteClassTfidf
    counts=sparse.load_npz(BASE/'models/class_word_counts.npz')
    expected=FiniteClassTfidf(reduce_frequent_words=True).fit_transform(counts)
    actual=sparse.load_npz(BASE/'models/ctfidf.npz')
    diff=(actual-expected).tocsr();error=float(np.abs(diff.data).max()) if diff.nnz else 0.
    assert error<1e-12;checks['ctfidf_max_abs_error']=error
    # Reload the actual original BERTopic model, not only its KMeans component.
    evidence=pq.read_table(RAW/'evidence/review_documents.parquet').to_pylist()
    chosen={r['row_id']:r for r in evidence[::max(1,len(evidence)//64)][:64]}
    for group in dec['merge_groups']:
        for i in group['members']:
            for r in [r for r in evidence if r['cluster']==i][:4]:chosen[r['row_id']]=r
    examples=list(chosen.values());parts={p['file']:p for p in read(CORPUS/'data/INPUT_MANIFEST.json')['parts']}
    emb=[]
    for r in examples:
        p=parts[r['source_part']]
        a=np.load(CORPUS/'models/embeddings'/f'{Path(p["file"]).stem}.npy',mmap_mode='r')
        emb.append(a[r['row_id']-p['first_row']].astype(np.float32))
    model=SemanticTopicModel()
    predicted=model.predict_labels([r['title'] or 'source record' for r in examples],np.array(emb))
    rows=np.array([r['row_id'] for r in examples])
    raw=np.load(RAW/'results/raw_labels.npy',mmap_mode='r');final=np.load(BASE/'results/final_labels.npy',mmap_mode='r')
    np.testing.assert_array_equal(predicted.raw_cluster1000.to_numpy(),raw[rows])
    np.testing.assert_array_equal(predicted.final_topic_index.to_numpy(),final[rows])
    mapping=pd.read_csv(BASE/'results/raw_to_final.csv').fillna('')
    np.testing.assert_array_equal(predicted.final_topic_id.to_numpy(),mapping.set_index('raw_cluster1000').loc[raw[rows]].final_topic_id.to_numpy())
    checks['reloaded_bertopic_and_mapping_predictions']=len(examples)
    checks['prediction_samples_include_all_merged_raw_topics']=all(i in set(raw[rows]) for g in dec['merge_groups'] for i in g['members'])
    predicted['row_id']=rows;predicted.to_csv(BASE/'evidence/prediction_verification.csv',index=False,encoding='utf-8-sig')
    tests=read(BASE/'TEST_RESULTS.json');assert tests['passed'] and tests['tests_run']==5
    checks['mapping_safety_tests_passed']=tests['tests_run']
    print('prediction and ctfidf checks passed:',len(examples),flush=True)
    topics=pd.read_csv(BASE/'results/final_topics.csv').fillna('')
    merges=pd.read_csv(BASE/'results/merge_summary.csv').fillna('')
    cand=pd.read_csv(BASE/'review/candidate_decisions.csv').fillna('')
    records=[]
    for p in sorted((BASE/'evidence').glob('packet_*.json')):records.extend(read(p)['records'])
    records.extend(read(BASE/'evidence/fresh_random_audit.json')['records'])
    review=pd.DataFrame(records).drop_duplicates('row_id').fillna('')
    review.to_csv(BASE/'evidence/reviewed_records.csv',index=False,encoding='utf-8-sig')
    keeps=pd.DataFrame(dec['keep_or_defer'])
    for c in ['members','evidence_row_ids']:keeps[c]=keeps[c].apply(lambda x:'|'.join(map(str,x)))
    keeps.to_csv(BASE/'review/keep_or_defer.csv',index=False,encoding='utf-8-sig')
    # Read back concrete examples from the delivered shards, covering merged,
    # unchanged and unusable records. This is a workbook sample, not full output.
    assignment_manifest=read(BASE/'results/ASSIGNMENT_MANIFEST.json');sample=[];want=set(rows.tolist())
    unusable_ids=np.flatnonzero(raw<0)[:5];want.update(map(int,unusable_ids))
    for p in assignment_manifest['parts']:
        ids=sorted(r for r in want if p['first_row']<=r<=p['last_row'])
        if not ids:continue
        table=pq.read_table(BASE/p['file']);sample.extend(table.take([r-p['first_row'] for r in ids]).to_pylist())
    sample=pd.DataFrame(sample)
    for r in sample.itertuples():
        if r.new_cluster1000>=0:assert r.final_topic_id==mapping.iloc[r.new_cluster1000].final_topic_id
        else:assert r.final_topic_id=='UNUSABLE' and r.final_topic_index==-2
    metrics=pd.DataFrame([{'指标':k,'值':json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v} for k,v in s.items()])
    contents={'结果概览':metrics,'确认合并':merges,'原1000到最终标签':mapping,'最终998类':topics,'保留与暂缓':keeps,
              '已读证据443条':review,'候选及审定状态':cand,'文献标签示例':sample}
    workbook=BASE/'1000类语义合并与文献标签.xlsx'
    with pd.ExcelWriter(workbook,engine='openpyxl') as writer:
        for name,f in contents.items():f.map(excel_safe).to_excel(writer,sheet_name=name,index=False)
        from openpyxl.styles import Font,PatternFill
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='214A63')
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(16,len(str(col[0].value))+4))
    from openpyxl import load_workbook
    wb=load_workbook(workbook,read_only=True,data_only=False)
    assert wb.sheetnames==list(contents)
    for name,f in contents.items():assert wb[name].max_row==len(f)+1
    wb.close();checks['workbook_sheets_verified']=len(contents)
    checks['reviewed_document_export_matches_audit']=len(review)==s['reviewed_documents']
    checks['label_sample_covers_unusable']=int((sample.new_cluster1000<0).sum())
    # Explicit classification limits accompany the successful technical checks.
    validation.update({'passed':True,'stage':'complete','created_utc':now(),
       'semantic_review':{'inventory_topics':1000,'detailed_topics':60,'unique_sample_records':443,'fresh_random_records':96,'reviewer':'Codex模型辅助审阅','independent_expert_review':False},
       'not_established':['exhaustive synonym elimination','whole-corpus semantic purity or accuracy','uniform granularity of all 998 topics','independent expert approval','TRL/CRL maturity scores']})
    dump(BASE/'VALIDATION.json',validation)
    write_readme(s,merges)
    dump(BASE/'PROGRESS.json',{'stage':'complete','updated_utc':now(),'final_topics':s['final_topics'],'documents':s['documents'],'validation_passed':True})
    # Manifest excludes itself and the currently-open log, avoiding stale hashes.
    files=[p for p in BASE.rglob('*') if p.is_file() and p.name!='DELIVERY_MANIFEST.json' and p.suffix!='.log' and '__pycache__' not in p.parts]
    dump(BASE/'DELIVERY_MANIFEST.json',{'created_utc':now(),'files':[{'path':str(p.relative_to(BASE)),'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(files)],'excluded':['*.log','__pycache__','DELIVERY_MANIFEST.json'],'raw_inputs_are_referenced':True})
    pointer=ROOT/'jjjj/bertopic1000_FINAL_LABELS.json'
    dump(pointer,{'version':s['version'],'directory':str(BASE),'primary_label_fields':['final_topic_id','final_topic_label'],'assignment_manifest':str(BASE/'results/ASSIGNMENT_MANIFEST.json'),'mapping':str(BASE/'results/raw_to_final.csv'),'topics':s['final_topics'],'validation':str(BASE/'VALIDATION.json'),'raw1000_directory':str(RAW),'trl_crl_recalculated':False})
    print(json.dumps({'status':'complete','topics':s['final_topics'],'documents_in_merged_groups':s['documents_in_merged_groups'],'workbook':str(workbook),'validation_passed':True},ensure_ascii=False),flush=True)

def write_readme(s,merges):
    table='\n'.join(f'| {r.raw_topics.replace("|", " + ")} | {r.final_topic_id} | {r.final_topic_label} | {r.documents:,} |' for r in merges.itertuples())
    text=f"""# 1000类语义合并结果

已生成 **{s['final_topics']}个非空主题类**。本轮按“对象、任务、分类层级相同，主要差异为表述”的标准确认2组同义类别，涉及 **{s['documents_in_merged_groups']:,}条记录**。未设定合并后的目标类数。

| 原始类 | 最终类 | 统一标签 | 记录数 |
| --- | --- | --- | ---: |
{table}

全部 **{s['documents']:,}条记录**已经写入259个新标签分片并逐行回读。其中 **{s['assigned_documents']:,}条可分类记录**有最终标签，**{s['unusable_documents']:,}条不可用文本**仍为 `-2 / UNUSABLE`。合并组中23,348条记录共享新的统一名称；若按组内保留最小原类编号计算，10,061条记录的类别归属锚点发生变化。其余996个原类保持独立，未拆类或按单篇重新分配。

## 文件入口

- [审阅与结果工作簿](1000类语义合并与文献标签.xlsx)：概览、2组合并、完整映射、998类、保留/暂缓理由、443条已读证据、2472候选及状态、文献标签示例。Excel不承载500万条全量文献。
- [文献最终标签目录](results/assignments/)：全量259个Parquet分片。
- [分片清单](results/ASSIGNMENT_MANIFEST.json)：行数、行号范围、输入与输出哈希。
- [1000到998完整映射](results/raw_to_final.csv)：1000行，每个原类恰好一个去向。
- [998类目录](results/final_topics.csv)：合并后文献/专利/政策计数和重算主题词。
- [明确合并及保留决定](review/decisions.json)：2组合并与38项保留/暂缓记录。
- [技术验证](VALIDATION.json)、[交付文件哈希](DELIVERY_MANIFEST.json)。

## 文献标签如何使用

读取 `results/assignments/part-*.parquet` 的 **`final_topic_id` 与 `final_topic_label`**；以 `row_id` 对接原始语料。`doc_id`、来源、日期、质量标记、旧500/486类及原1000类字段均保留。`new_topic_id` / `new_cluster1000` 明确作为原1000类追溯字段。入口指针为 `../bertopic1000_FINAL_LABELS.json`，消费程序应据此选择新版本；本轮没有修改其他历史版本或外部面板。

`final_topic_index` 为0到997的连续索引；`Sxxxx` 后缀为组内最小原类编号，允许不连续，不能把后缀当作连续索引。`semantic_merge_applied` 表示该记录属于确认的合并组。未合并类保留自动名称，`final_label_status` 区分清单初筛、重点抽查和已合并状态。

原1000类的余弦、距离间隔和第二候选均保留其原字段名；新增 `final_secondary_topic_id` 只映射候选编号，`secondary_maps_to_primary` 表明该候选是否已经落入同一最终类。没有把原距离间隔冒充合并后的置信度。

## 审阅范围与严格取舍

完成1000类主题词、主来源中心题名及最近邻清单初筛；补查相似度≥0.94的122对主题词候选。候选集合共2472对，来源于每类前3个簇心近邻、簇心余弦≥0.90或c-TF-IDF余弦≥0.80的并集；这些都是召回规则，**不是合并阈值**。没有声称逐对审定全部2472对。

重点查看60个原类的中心、随机及边界题名和已记录的正文片段，共443条唯一文献；其中96条来自排除既有样本后的新增随机抽查。空正文如实保留。证据包和实际显示范围见 `review/ACTUALLY_INSPECTED.json` 与 `evidence/reviewed_records.csv`。这是模型辅助抽样审阅，未作独立专家验收、逐篇人工标注或全量准确率测量。

新增随机抽查使若干初步候选被撤回：电池顶盖与整包结构、太阳能电池结构与制造工艺、正极改性与宽泛复合材料、TiO2污染物去除与更广的光催化反应、电价分析与不同电价核定环节、厌氧消化与污水反应器任务。只要存在实质分界或证据不足，本轮保留原类。这不证明剩余类之间完全没有同义或重叠；保留/暂缓项仍可在补充证据后复议。

合并后的类仍可能包含原聚类的边界噪声，例如OLED类混入光伏材料、并网光伏类混入组件/离网研究。本次只进行整类同义合并，**没有宣称修复这些单篇错分**。998是主题类数量，不是经专家确认的998项独立技术，也不是已有TRL/CRL结果的数量。

## 计算与验证

最终标签通过固定映射从原1000类得到；没有重新embedding、重新训练或用合并后质心再次分配文献。词频按原类精确求和，再按原配置重算c-TF-IDF。未来推理入口为 `src/semantic_mapping.py` 的 `SemanticTopicModel.predict_labels(documents, embeddings)`，必须使用与原模型相同的文本/384维向量及不可用文本筛查流程；内部执行原BERTopic1000预测后映射。不会输出伪造的合并概率。

259个分片全部回读，原列和新增列逐行核验；主题数、来源数、未分类标记和映射完整性均核对。重新加载原BERTopic并验证抽样预测映射；重算c-TF-IDF与原实现对比；5项测试覆盖负标签、越界、重复成员和重叠合并。原1000类快照的全部312个清单文件哈希保持一致。

TRL/CRL没有重新核算，不能直接把历史500/486类成熟度结果复制给新标签。

复现顺序：`python3 src/apply_merges.py` 后执行 `python3 src/finalize.py`。输入和环境路径见 `src/merge_common.py`，最终人工可复核的决定固定在 `review/decisions.json`。输出在本目录，原始输入目录只读引用。
"""
    (BASE/'README.md').write_text(text)
    (BASE/'REPORT.md').write_text(text)
if __name__=='__main__':main()
