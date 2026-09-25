"""Full document assignments, old/new flow tables, review samples and comparison."""
from common import *
import collections, heapq, math
import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import joblib
from threadpoolctl import threadpool_limits
from sklearn.metrics import adjusted_rand_score, normalized_mutual_info_score

def keep(heaps,key,score,rid,n):
    item=(float(score),int(rid));h=heaps[key]
    if len(h)<n:heapq.heappush(h,item)
    elif item>h[0]:heapq.heapreplace(h,item)

def distribution(sizes):
    a=np.asarray(sizes,dtype=float);a=np.sort(a)
    return {'clusters':len(a),'minimum':int(a.min()),'p10':float(np.quantile(a,.1)),
            'median':float(np.median(a)),'p90':float(np.quantile(a,.9)),'maximum':int(a.max()),
            'mean':float(a.mean()),'clusters_below_100':int((a<100).sum()),
            'clusters_below_500':int((a<500).sum()),'clusters_below_1000':int((a<1000).sum()),
            'gini':float(2*np.sum(np.arange(1,len(a)+1)*a)/(len(a)*a.sum())-(len(a)+1)/len(a))}

def flow_tables(table,old_names,old_statuses,prefix):
    edges=[];summaries=[]
    for old in range(len(table)):
        size=int(table[old].sum())
        if not size:continue
        order=np.argsort(table[old])[::-1];nonzero=order[table[old,order]>0]
        for new in nonzero:
            count=int(table[old,new])
            edges.append({'old_topic_id':f'{prefix}{old:03d}','old_name':old_names.get(old,''),
                          'new_topic_id':f'T{new:04d}','documents':count,'share_of_old':count/size,
                          'share_of_new':count/int(table[:,new].sum())})
        summaries.append({'old_topic_id':f'{prefix}{old:03d}','old_name':old_names.get(old,''),
                          'old_status':old_statuses.get(old,''),'documents':size,
                          'new_topics_with_any_documents':len(nonzero),
                          'significant_new_topics':int(((table[old]>=50)&(table[old]/size>=.05)).sum()),
                          'largest_new_topic':f'T{order[0]:04d}','largest_new_share':int(table[old,order[0]])/size,
                          'new_topics_covering_80pct':int(np.searchsorted(np.cumsum(table[old,order])/size,.8)+1),
                          'top5_flows':' ; '.join(f'T{j:04d}: {int(table[old,j])} ({table[old,j]/size:.1%})' for j in nonzero[:5])})
    return pd.DataFrame(edges),pd.DataFrame(summaries)

def excel_safe(value):
    if isinstance(value,str):
        value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',' ',value)
        return value[:32000]
    return value

def export():
    threadpool_limits(4)
    audit=read(BASE/'models/CLUSTER_AUDIT.json');inp=read(SOURCE/'data/INPUT_MANIFEST.json')
    parts=inp['parts'];n=inp['documents']
    labels=np.load(BASE/'results/raw_labels.npy',mmap_mode='r');old=np.load(SOURCE/'results/raw_labels.npy',mmap_mode='r')
    usable=np.load(BASE/'data/usable.npy');source=np.load(BASE/'data/source_codes.npy')
    second=np.load(BASE/'results/secondary_labels.npy',mmap_mode='r');margin=np.load(BASE/'results/distance_margin.npy',mmap_mode='r')
    distance=np.load(BASE/'results/distance_squared.npy',mmap_mode='r')
    centers=np.load(BASE/'models/topic_centroids.npy');oldcenters=np.load(SOURCE/'models/topic_centroids.npy')
    z=np.load(SOURCE/'models/reduced.npy',mmap_mode='r')
    oldkm=joblib.load(SOURCE/'models/kmeans500.joblib')
    km=joblib.load(BASE/'models/kmeans1000.joblib')
    assert np.array_equal(labels>=0,old>=0) and np.array_equal(labels>=0,usable)
    mapping=pd.read_csv(SOURCE/'results/raw_to_merged.csv')
    lookup=np.full(500,-1,dtype=int)
    for r in mapping.itertuples():lookup[r.cluster]=int(r.merged_topic_id[1:])
    oldnames={int(r.cluster):r.reviewed_name for r in mapping.itertuples()}
    oldstatuses={int(r.cluster):r.status for r in mapping.itertuples()}
    mergednames={int(r.merged_topic_id[1:]):r.merged_name for r in mapping.itertuples()}
    mergedstatuses={int(r.merged_topic_id[1:]):r.merged_status for r in mapping.itertuples()}
    cross=np.bincount(old[usable].astype(np.int64)*K+labels[usable],minlength=500*K).reshape(500,K)
    mergedcross=np.zeros_like(cross)
    np.add.at(mergedcross,lookup,cross)
    assert cross.sum()==usable.sum() and mergedcross.sum()==usable.sum()
    edges500,oldsummary=flow_tables(cross,oldnames,oldstatuses,'B')
    edges486,mergedsummary=flow_tables(mergedcross,mergednames,mergedstatuses,'M')
    for name,frame in [('old500_to_new1000',edges500),('old486_to_new1000',edges486),('old500_flow_summary',oldsummary),('old486_flow_summary',mergedsummary)]:
        frame.to_csv(BASE/f'results/{name}.csv',index=False,encoding='utf-8-sig')
    np.save(BASE/'results/contingency_500x1000.npy',cross)
    counts=np.bincount(labels[usable],minlength=K)
    sums=np.zeros(K);squares=np.zeros(K);old_cos_total=0.;new_cos_total=0.
    by_source={s:{'documents':0,'cosine500_sum':0.,'cosine1000_sum':0.,'distance500_sum':0.,'distance1000_sum':0.} for s in ['paper','patent','policy']}
    low_cos=np.zeros(K,dtype=np.int64);low_margin=np.zeros(K,dtype=np.int64)
    flags={key:np.zeros(K,dtype=np.int64) for key in ['title_only','retracted','template_record']}
    heaps=collections.defaultdict(list);assignment_manifest=[];prediction_checks=0
    source_names=['paper','patent','policy'];out=BASE/'results/assignments';out.mkdir(exist_ok=True)
    columns=['row_id','doc_id','source','date','title_only','retracted','template_record','quality_note']
    for i,p in enumerate(parts):
        path=SOURCE/p['file'];d=pq.read_table(path,columns=columns).to_pydict();ids=np.asarray(d['row_id'])
        good=usable[ids];ix=ids[good];y=labels[ix];sc=source[ix]
        x=np.load(SOURCE/'models/embeddings'/f'{path.stem}.npy').astype(np.float32)
        x/=np.maximum(np.linalg.norm(x,axis=1,keepdims=True),1e-12)
        sim=np.einsum('ij,ij->i',x[good],centers[y])
        oldsim=np.einsum('ij,ij->i',x[good],oldcenters[old[ix]])
        sums+=np.bincount(y,weights=sim,minlength=K);squares+=np.bincount(y,weights=sim**2,minlength=K)
        low_cos+=np.bincount(y[sim<.45],minlength=K);low_margin+=np.bincount(y[margin[ix]<.02],minlength=K)
        olddist=((z[ix]-oldkm.cluster_centers_[old[ix]])**2).sum(1)
        for sid,s in enumerate(source_names):
            mask=sc==sid;b=by_source[s];b['documents']+=int(mask.sum())
            for key,value in [('cosine500_sum',oldsim),('cosine1000_sum',sim),('distance500_sum',olddist),('distance1000_sum',distance[ix])]:
                b[key]+=float(np.sum(value[mask],dtype=np.float64))
        for flag in flags:
            flags[flag]+=np.bincount(labels[ids[good&np.asarray(d[flag],dtype=bool)]],minlength=K)
        group=y*3+sc;order=np.argsort(group,kind='stable')
        bounds=np.r_[0,np.flatnonzero(np.diff(group[order]))+1,len(order)]
        priority=ix.astype(np.uint64)+np.uint64(SEED)
        priority=(priority^(priority>>np.uint64(30)))*np.uint64(0xbf58476d1ce4e5b9)
        priority=(priority^(priority>>np.uint64(27)))*np.uint64(0x94d049bb133111eb)
        priority=priority^(priority>>np.uint64(31))
        for a,b in zip(bounds[:-1],bounds[1:]):
            loc=order[a:b]
            if not len(loc):continue
            cluster,sid=divmod(int(group[loc[0]]),3)
            for role,score,take in [('center',sim[loc],3),('random',priority[loc].astype(np.float64),2),('boundary',-margin[ix[loc]],2),('far',-sim[loc],1)]:
                for j in np.argsort(score)[-take:]:keep(heaps,(cluster,sid,role),score[j],ix[loc[j]],take)
        cosines=np.full(len(ids),np.nan,dtype=np.float32);cosines[good]=sim
        merged=np.full(len(ids),-2,dtype=np.int16);merged[good]=lookup[old[ids[good]]]
        table=pa.table({**d,'old_cluster500':old[ids],'old_merged486':merged,'new_cluster1000':labels[ids],
                        'new_topic_id':[f'T{j:04d}' if j>=0 else 'UNUSABLE' for j in labels[ids]],
                        'secondary_candidate1000':second[ids],'distance_margin1000':margin[ids],
                        'cosine_to_centroid1000':cosines})
        dest=out/path.name;pq.write_table(table,dest,compression='zstd')
        check=pq.read_table(dest,columns=['row_id','new_cluster1000','old_cluster500','old_merged486'])
        assert np.array_equal(check['row_id'].to_numpy(),ids)
        assert np.array_equal(check['new_cluster1000'].to_numpy(),labels[ids])
        assert np.array_equal(check['old_cluster500'].to_numpy(),old[ids])
        assert np.array_equal(check['old_merged486'].to_numpy(),merged)
        sample=ix[:8]
        if len(sample):
            assert np.array_equal(km.predict(np.array(z[sample])),labels[sample]);prediction_checks+=len(sample)
        assignment_manifest.append({'file':str(dest.relative_to(BASE)),'rows':len(ids),'sha256':sha(dest),'source_part':p['file'],'source_sha256':p['sha256']})
        if i%20==0:progress('exporting_document_assignments',parts_done=i+1,parts_total=len(parts))
    selection=collections.defaultdict(list)
    for (cluster,sid,role),h in heaps.items():
        for rank,(score,rid) in enumerate(sorted(h,reverse=True),1):
            selection[rid].append({'role':role,'rank':rank})
    evidence=[];topic_examples={i:{} for i in range(K)}
    for i,p in enumerate(parts):
        ids=np.arange(p['first_row'],p['first_row']+p['n']);positions=[j for j,rid in enumerate(ids) if int(rid) in selection]
        if not positions:continue
        path=SOURCE/p['file']
        # Only the selected rows are saved, with original full text and its hash.
        table=pq.read_table(path).take(pa.array(positions))
        for r in table.to_pylist():
            rid=r['row_id'];cluster=int(labels[rid]);r.update(cluster=cluster,topic_id=f'T{cluster:04d}',
                 old_topic_id=f'B{old[rid]:03d}',old_merged_id=f'M{lookup[old[rid]]:03d}',
                 selection_roles='|'.join(f"{x['role']}:{x['rank']}" for x in selection[rid]),
                 body_sha256=digest(r['body']),source_part=p['file'])
            evidence.append(r)
            for role in ['center:1','random:1']:
                if role in r['selection_roles'].split('|'):
                    topic_examples[cluster][r['source']+'_'+role]=r['title']
        if i%40==0:progress('saving_review_samples',parts_done=i+1,parts_total=len(parts))
    pq.write_table(pa.Table.from_pylist(evidence),BASE/'evidence/review_documents.parquet',compression='zstd')
    compact=[{k:r[k] for k in ['row_id','topic_id','old_topic_id','old_merged_id','source','title','date','selection_roles','body_sha256','source_part']}|{'body_excerpt':r['body'][:1500]} for r in evidence]
    pd.DataFrame(compact).to_csv(BASE/'evidence/review_samples.csv',index=False,encoding='utf-8-sig')
    topics=pd.read_csv(BASE/'results/raw_1000_topics.csv')
    topics['mean_cosine']=sums/counts;topics['std_cosine']=np.sqrt(np.maximum(0,squares/counts-(sums/counts)**2))
    topics['cosine_below_045']=low_cos;topics['margin_below_002']=low_margin
    for flag in flags:topics[flag+'_documents']=flags[flag]
    largest_old=cross.argmax(0);largest_merged=mergedcross.argmax(0)
    topics['dominant_old500']=[f'B{i:03d}' for i in largest_old]
    topics['dominant_old500_name']=[oldnames[i] for i in largest_old]
    topics['dominant_old500_share']=cross[largest_old,np.arange(K)]/counts
    topics['dominant_old486']=[f'M{i:03d}' for i in largest_merged]
    topics['dominant_old486_name']=[mergednames[i] for i in largest_merged]
    topics['dominant_old486_share']=mergedcross[largest_merged,np.arange(K)]/counts
    topics['old500_sources_with_any_documents']=(cross>0).sum(0)
    for key in ['paper_center:1','paper_random:1','patent_center:1','patent_random:1','policy_center:1']:
        topics[key.replace(':','_')]=[topic_examples[i].get(key,'') for i in range(K)]
    topics['review_status']='自动类群；未完成逐类语义审定'
    topics.to_csv(BASE/'results/topics1000_with_comparison.csv',index=False,encoding='utf-8-sig')
    totals={key:sum(v[key] for v in by_source.values()) for key in ['cosine500_sum','cosine1000_sum','distance500_sum','distance1000_sum']}
    weights=read(BASE/'models/CLUSTER_AUDIT.json')['per_document_training_weights']
    weighted_den=sum(weights[s]*r['documents'] for s,r in by_source.items())
    geometry={'mean_cosine500':totals['cosine500_sum']/int(usable.sum()),'mean_cosine1000':totals['cosine1000_sum']/int(usable.sum()),
              'mean_squared_distance500':totals['distance500_sum']/int(usable.sum()),'mean_squared_distance1000':totals['distance1000_sum']/int(usable.sum()),
              'weighted_mean_squared_distance500':sum(weights[s]*r['distance500_sum'] for s,r in by_source.items())/weighted_den,
              'weighted_mean_squared_distance1000':sum(weights[s]*r['distance1000_sum'] for s,r in by_source.items())/weighted_den}
    for r in by_source.values():
        for key in list(r):
            if key.endswith('_sum'):r[key.replace('_sum','_mean')]=r.pop(key)/r['documents']
    summary={'documents':n,'assigned_usable_documents':int(usable.sum()),'unusable_text_documents':int((~usable).sum()),
             'old_raw_clusters':500,'old_merged_clusters':486,'new_raw_clusters':K,'new_merges':0,
             'old500_sizes':distribution(cross.sum(1)),'new1000_sizes':distribution(counts),'geometry':geometry,'geometry_by_source':by_source,
             'flow_threshold':'A significant flow contains >=50 documents and >=5% of its old topic; diagnostic only',
             'old500_with_multiple_significant_new_topics':int((oldsummary.significant_new_topics>=2).sum()),
             'old486_with_multiple_significant_new_topics':int((mergedsummary.significant_new_topics>=2).sum()),
             'new_topics_at_least_80pct_from_one_old500':int((topics.dominant_old500_share>=.8).sum()),
             'new_topics_at_least_80pct_from_one_old486':int((topics.dominant_old486_share>=.8).sum()),
             'new_topics_contained_in_single_old500':int(((cross>0).sum(0)==1).sum()),
             'new_topics_contained_in_single_old486':int(((mergedcross>0).sum(0)==1).sum()),
             'old_status_flow_counts':mergedsummary.groupby('old_status').apply(lambda f:int((f.significant_new_topics>=2).sum()),include_groups=False).to_dict(),
             'partition_agreement_not_accuracy':{'adjusted_rand_index':float(adjusted_rand_score(old[usable],labels[usable])),
                'normalized_mutual_information':float(normalized_mutual_info_score(old[usable],labels[usable]))},
             'saved_review_documents':len(evidence),'semantic_accuracy':None,'independent_expert_review':False,
             'maturity_recalculated':False,'notes':['global reclustering, not a hierarchical split','more clusters and higher within-cluster similarity do not prove higher semantic accuracy','no automatic merges or inherited TRL/CRL grades']}
    dump(BASE/'results/COMPARISON.json',summary)
    dump(BASE/'results/ASSIGNMENT_MANIFEST.json',{'parts':assignment_manifest,'documents':n,'all_rows_read_back':True})
    dump(BASE/'evidence/EVIDENCE_AUDIT.json',{'selected_documents':len(evidence),'topics_covered':len({r['cluster'] for r in evidence}),
         'per_topic_per_source':{'center':3,'random':2,'boundary':2,'far':1},'selected_is_not_manually_reviewed':True})
    # Validate the actual saved BERTopic transform, not only the KMeans component.
    model=joblib.load(BASE/'models/bertopic1000.joblib')
    samples=evidence[::max(1,len(evidence)//150)][:150]
    sample_x=[];sample_y=[]
    for r in samples:
        p=next(p for p in parts if p['file']==r['source_part'])
        x=np.load(SOURCE/'models/embeddings'/f'{Path(p["file"]).stem}.npy',mmap_mode='r')
        sample_x.append(x[r['row_id']-p['first_row']].astype(np.float32));sample_y.append(r['cluster'])
    predicted,_=model.transform([r['title'] or 'source record' for r in samples],embeddings=np.array(sample_x))
    assert np.array_equal(predicted,sample_y)
    baseline=read(BASE/'data/SOURCE_BASELINE.json')
    unchanged=all(sha(SOURCE/r['path'])==r['sha256'] for r in baseline['files'])
    assert unchanged
    validation={'passed':True,'checks':{'exactly_1000_nonempty_clusters':len(topics)==K and (counts>0).all(),
                  'same_corpus_and_usable_mask_as500':True,'flow_counts_account_for_all_usable_documents':True,
                  'all_assignment_parts_read_back':True,'source_protected_files_unchanged':True,
                  'kmeans_reloaded_predictions_checked':prediction_checks,'bertopic_reloaded_predictions_checked':len(samples),
                  'direct_ctfidf_max_abs_error':read(BASE/'models/STREAMING_VALIDATION.json')['direct_ctfidf_max_abs_error']},
                'not_validated':['semantic accuracy','all1000 topic names','uniform granularity','independent expert acceptance'], 'created_utc':now()}
    validation['checks']['exactly_1000_nonempty_clusters']=bool(validation['checks']['exactly_1000_nonempty_clusters'])
    dump(BASE/'VALIDATION.json',validation)
    make_workbook(summary,topics,oldsummary,mergedsummary,edges500,compact)
    draw_comparison(summary,cross.sum(1),counts)
    write_report(summary,topics,mergedsummary)
    progress('delivery_ready',clusters=K,documents=n,review_samples=len(evidence))

def make_workbook(summary,topics,oldsummary,mergedsummary,edges,samples):
    from openpyxl import load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    path=BASE/'1000类试验与500类对比.xlsx'
    metrics=[{'指标':k,'值':json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v} for k,v in summary.items()]
    renames={'topic_id':'新类编号','documents':'文档数','papers':'文献数','patents_2026':'2026专利数','policies':'政策数','keywords':'自动主题词',
             'automatic_label':'自动标签','dominant_old500':'主要原500类','dominant_old500_name':'原500类名称参考','dominant_old500_share':'来自主要原500类的比例',
             'dominant_old486':'主要原486类','dominant_old486_name':'原486类名称参考','dominant_old486_share':'来自主要原486类的比例',
             'mean_cosine':'平均簇心余弦相似度','review_status':'审阅状态','old_topic_id':'旧类编号','old_name':'旧类名称','old_status':'旧类状态',
             'significant_new_topics':'显著分流新类数','new_topics_with_any_documents':'涉及新类数','largest_new_topic':'最大去向类','largest_new_share':'最大去向比例',
             'top5_flows':'前五个去向','new_topic_id':'新类编号','share_of_old':'占旧类比例','share_of_new':'占新类比例',
             'row_id':'原文行号','title':'标题','body_excerpt':'原文节选','selection_roles':'抽样类型'}
    with pd.ExcelWriter(path,engine='openpyxl') as writer:
        for name,frame in [('阅读说明',pd.DataFrame(metrics)),('1000类目录',topics),('原500类去向',oldsummary),('原486类去向',mergedsummary),('500到1000交叉表',edges),('中心随机边界样本',pd.DataFrame(samples))]:
            frame=frame.map(excel_safe).rename(columns=renames)
            frame.to_excel(writer,sheet_name=name,index=False)
    wb=load_workbook(path)
    for ws in wb:
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B54')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=26
            for c in col:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
        for row in ws.iter_rows(min_row=2):ws.row_dimensions[row[0].row].height=44
    wb.save(path)

def draw_comparison(summary,old_sizes,new_sizes):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig,axes=plt.subplots(1,2,figsize=(11,4))
    for values,name,color in [(old_sizes,'K = 500','#527692'),(new_sizes,'K = 1000','#ce8d44')]:
        axes[0].plot(np.arange(1,len(values)+1)/len(values),np.sort(values),label=name,color=color)
    axes[0].set_yscale('log');axes[0].set_xlabel('Cluster-size percentile');axes[0].set_ylabel('Documents per cluster (log)');axes[0].legend()
    g=summary['geometry'];axes[1].bar(['K = 500','K = 1000'],[g['mean_cosine500'],g['mean_cosine1000']],color=['#527692','#ce8d44'])
    axes[1].set_ylim(0,1);axes[1].set_ylabel('Mean cosine to assigned centroid')
    fig.suptitle('Same corpus and embeddings; global reclustering')
    fig.text(.5,.01,'Geometric diagnostics are not semantic accuracy.',ha='center',fontsize=9)
    fig.tight_layout(rect=[0,.04,1,.93]);fig.savefig(BASE/'results/comparison.png',dpi=160);plt.close(fig)

def write_report(s,topics,mergedsummary):
    a,b=s['old500_sizes'],s['new1000_sizes'];g=s['geometry']
    lines=['**1000类试验已完成。这是相同文档和向量上的全局重新聚类，保留现有500/486类结果；本轮未自动合并。**','',
      f"输入 {s['documents']:,} 个文档单位，其中 {s['assigned_usable_documents']:,} 条可用文本全部获得1000类之一的标签，{s['unusable_text_documents']:,} 条无可用文本保留UNUSABLE。分配率不是准确率。",'',
      '方法：固定原多语言MiniLM向量、原PCA64与归一化、来源训练权重（文献70%/专利25%/政策5%）、随机种子500及三遍全量训练，只把MiniBatchKMeans的类数改为1000。通过BERTopic预计算词频适配器重算全部文档的c-TF-IDF；复用的是原文全量词频，不是旧主题词或旧类别标签。','',
      '| 指标 | 原500类 | 新1000类 |','|---|---:|---:|',
      f"| 最小类文档数 | {a['minimum']:,} | {b['minimum']:,} |",f"| 类大小中位数 | {a['median']:,.1f} | {b['median']:,.1f} |",
      f"| 最大类文档数 | {a['maximum']:,} | {b['maximum']:,} |",f"| 少于100篇的类 | {a['clusters_below_100']} | {b['clusters_below_100']} |",
      f"| 平均簇心余弦相似度 | {g['mean_cosine500']:.4f} | {g['mean_cosine1000']:.4f} |",
      f"| 按训练来源权重计算的平均平方距离 | {g['weighted_mean_squared_distance500']:.4f} | {g['weighted_mean_squared_distance1000']:.4f} |",'',
      '类数增加通常会使类更小、类内距离下降，这些几何变化不能直接证明语义准确率提高。当前没有独立金标准，不能给出准确率提升百分比。','',
      f"以‘至少50篇且占旧类至少5%’作为分流诊断条件，原500类中有 {s['old500_with_multiple_significant_new_topics']} 类流向至少两个新类；原486类中有 {s['old486_with_multiple_significant_new_topics']} 类出现这种分流。该条件只是统计口径，不代表拆分已获语义认可。",'',
      f"1000个新类中，{s['new_topics_at_least_80pct_from_one_old500']} 个至少80%的文章来自同一个原500类；完全包含在单一原500类中的新类只有 {s['new_topics_contained_in_single_old500']} 个。旧类与新类是交叉对应关系，不能解释为每个旧类都被一分为二。",'',
      f"已保存 {s['saved_review_documents']:,} 篇中心、随机、边界和远点样本。自动关键词与旧主题名称参考用于浏览，不等于1000个新类都已命名审定。抽样保存不等于已经逐篇人工核读。",'',
      '[Excel工作簿](1000类试验与500类对比.xlsx) · [1000类目录](results/topics1000_with_comparison.csv) · [原486类去向](results/old486_flow_summary.csv) · [逐文档分配清单](results/ASSIGNMENT_MANIFEST.json) · [完整比较指标](results/COMPARISON.json) · [验证记录](VALIDATION.json)','',
      '![类大小及聚合度对比](results/comparison.png)','',
      '原500/486类模型及结果保持不变。1000类没有自动继承热点排名或TRL/CRL等级；已有成熟度评价仍对应原来限定的技术对象。后续应以具体子类内容与随机样本决定是否采用新类、局部拆分或合并。']
    (BASE/'REPORT.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

if __name__=='__main__':export()
