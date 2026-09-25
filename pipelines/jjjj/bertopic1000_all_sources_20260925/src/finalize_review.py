"""Persist the assistant's bounded spot review and verify delivered artifacts."""
from common import *
import csv
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

COMMENTS={
    0:('电池与充电设施散热仍交叠','中心包含电池冷却和储能散热模组，随机专利为充电桩散热；仍需区分电池本体与充电设施对象。'),
    616:('存在明确跨对象混入','中心文献是电池热管理，但所抽专利含热电池供热及TOPCon太阳能电池退火。尤其row_id=5102665不属于动力/储能电池热管理；该类仅28篇专利，占总数很小，不能据这两个专利估计全类错误率。'),
    315:('器件与电路层级仍宽泛','中心文献讨论MOSFET/IGBT/SiC功率器件，随机文献为CMOS高速I/O收发器，专利为DCDC变换器和LLC保护电路；未形成单一同层技术对象。'),
    549:('变流器与变压器等仍交叠','中心文献是可再生能源变流器，随机文献为站用备用变压器改造；专利包括UPS功率转换和移动ESS装置。该类只有8篇专利，不将来源分层样本比例外推总体。'),
    816:('电价与负荷需求任务尚未完全分开','主题词偏电价预测，但中心文献为电力需求预测，随机文献为低压馈线场景分析；两个所抽专利则为电价预测。不能把B400向T0816/T0738分流直接解释成电价/负荷两个纯类。'),
    738:('负荷预测集中但有越界线索','中心文献为短期负荷预测，另有需求响应潜力评价；随机专利为可再生能源场址气象降尺度及布局，需逐文档核分。专利仅7篇，样本不是总体比例样本。'),
    646:('隔膜方向有集中迹象，完整边界待核','所抽中心与随机专利都围绕电池隔膜，两篇文献为多孔聚合物及电化学膜/隔膜。样本支持隔膜方向，但自动主题词仍覆盖电解质与粘结剂，不能据四篇样本认定三者已完全分离。'),
    444:('风电预测与状态监测仍混杂','中心文献为风机预测性维护/异常检测，随机文献为风速估计，两篇专利为风电发电量/功率预测；按研究任务划分仍有混杂。'),
}

def main():
    assert read(BASE/'PROGRESS.json')['stage']=='delivery_ready'
    df=pq.read_table(BASE/'evidence/review_documents.parquet',filters=[('cluster','in',list(COMMENTS))]).to_pandas()
    topics=pd.read_csv(BASE/'results/topics1000_with_comparison.csv').set_index('cluster')
    viewed={};findings=[]
    for cid,(judgement,reason) in COMMENTS.items():
        sub=df[df.cluster==cid];ids=[]
        for source in ['paper','patent','policy']:
            for role in ['center:1','random:1']:
                hit=sub[(sub.source==source)&sub.selection_roles.str.split('|').apply(lambda roles:role in roles)]
                if not len(hit):continue
                r=hit.iloc[0];rid=int(r.row_id);ids.append(rid)
                viewed[rid]={'row_id':rid,'topic_id':f'T{cid:04d}','source':source,
                             'selection_roles':r.selection_roles,'title_shown':r.title[:180],
                             'body_excerpt_shown':r.body[:360],'body_chars_shown':min(len(r.body),360),
                             'body_sha256':r.body_sha256,'source_part':r.source_part,
                             'url':str(r.get('url','')),'reviewer':'assistant','review_type':'agent_assisted_spot_check'}
        t=topics.loc[cid]
        findings.append({'topic_id':f'T{cid:04d}','documents':int(t.documents),'papers':int(t.papers),
                         'patents':int(t.patents_2026),'policies':int(t.policies),'judgement':judgement,'reason':reason,
                         'viewed_row_ids':';'.join(map(str,dict.fromkeys(ids))),'review_scope':'中心/随机题名及前360字符；非全文或总体准确率评估'})
    dump(BASE/'evidence/ACTUALLY_INSPECTED.json',{'topics':len(COMMENTS),'unique_documents':len(viewed),
         'documents_with_body_excerpt':sum(v['body_chars_shown']>0 for v in viewed.values()),'records':list(viewed.values()),
         'limitations':'有目的选择8个与旧宽泛/混杂类相关的新类，按来源选样；非全1000类代表性抽检，不能估计总体准确率'})
    pd.DataFrame(findings).to_csv(BASE/'results/spot_review.csv',index=False,encoding='utf-8-sig')
    dump(BASE/'results/SPOT_REVIEW.json',{'findings':findings,'all_1000_reviewed':False,'independent_expert_review':False})
    lines=['**对8个新类进行了有目的的辅助抽查；以下观察不能用于估计1000类的总体准确率。**','',
           f'实际查看{len(viewed)}篇唯一文档的题名与可用正文前360字符。中心与随机样本按来源选取，稀少来源会被过度代表，因此正文中的反例只证明存在问题，不代表其总体发生比例。完整样本行号和显示内容见 [核读记录](evidence/ACTUALLY_INSPECTED.json)。','',
           '| 新类 | 文档数 | 抽查结论 | 具体依据与限制 |','|---|---:|---|---|']
    for r in findings:lines.append(f"| {r['topic_id']} | {r['documents']} | {r['judgement']} | {r['reason']} |")
    lines+=['','T0000包含14篇文献和1340篇专利；T0616包含4242篇文献和28篇专利。两类都涉及电池热管理，而来源比例差异很大，说明后续需要核查分簇是否部分受来源或表达方式影响。不能把这种分流直接称为新增了两个独立技术方向。','',
            '这次增加类数使平均类规模下降，并提供更多局部细分线索；仍需要对关键技术方向开展对象/任务边界审阅，再决定哪些类采用、局部拆分或合并。当前1000类适合作为待复核底稿，尚不能当作已验收的技术目录。']
    (BASE/'SPOT_REVIEW.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    report=BASE/'REPORT.md';text=report.read_text()
    note=f'\n\n已完成{len(COMMENTS)}个新类、{len(viewed)}篇唯一文档的辅助抽查，观察与反例见 [抽查记录](SPOT_REVIEW.md)。电池热管理中的TOPCon工艺混入、风电功率预测与状态监测混杂等问题仍存在；该抽查不是总体准确率评估。\n'
    if '[抽查记录](SPOT_REVIEW.md)' not in text:report.write_text(text+note,encoding='utf-8')
    from openpyxl.styles import Font,PatternFill,Alignment
    path=BASE/'1000类试验与500类对比.xlsx'
    book=load_workbook(path)
    for title,records in [('8类抽查结论',findings),('实际核读记录',list(viewed.values()))]:
        if title in book:del book[title]
        ws=book.create_sheet(title);fields=list(records[0]);ws.append(fields)
        for r in records:ws.append([r[k] for k in fields])
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B54')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=35
            for c in col:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
        for row in ws.iter_rows(min_row=2):ws.row_dimensions[row[0].row].height=65
    book.save(path);book.close()
    wb=load_workbook(path,read_only=True,data_only=False)
    expected={'1000类目录':1000,'原500类去向':500,'原486类去向':486,'中心随机边界样本':read(BASE/'results/COMPARISON.json')['saved_review_documents']}
    expected.update({'8类抽查结论':len(COMMENTS),'实际核读记录':len(viewed)})
    for name,count in expected.items():assert wb[name].max_row==count+1
    formula_count=sum(c.data_type=='f' for ws in wb for row in ws for c in row)
    assert formula_count==0
    links=[]
    for p in BASE.glob('*.md'):
        for url in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',p.read_text()):
            if '://' in url:continue
            exists=(p.parent/url.split('#')[0]).resolve().exists();links.append({'document':p.name,'target':url,'exists':exists})
    assert all(r['exists'] for r in links)
    tests=(BASE/'logs/component_tests.txt').read_text()
    assert 'Ran 2 tests' in tests and tests.strip().endswith('OK')
    audit=read(BASE/'VALIDATION.json');audit['artifact_checks']={'workbook_sheets':wb.sheetnames,'checked_data_rows':expected,
               'no_excel_formulas':True,'document_links_checked':len(links),'links_valid':True,'component_tests_passed':2,
               'spot_review_topics':len(COMMENTS),'spot_review_unique_documents':len(viewed)}
    dump(BASE/'VALIDATION.json',audit)
    manifest=BASE/'DELIVERY_MANIFEST.json'
    files=[p for p in sorted(BASE.rglob('*')) if p.is_file() and '__pycache__' not in p.parts and p.suffix!='.pyc' and p!=manifest]
    dump(manifest,{'created_utc':now(),'input_cache_is_referenced_not_copied':True,
                  'files':[{'path':str(p.relative_to(BASE)),'bytes':p.stat().st_size,'sha256':sha(p)} for p in files]})
    print(json.dumps({'validated':True,'delivery_files':len(files),'spot_review_documents':len(viewed),'workbook_sheets':len(wb.sheetnames)},ensure_ascii=False))

if __name__=='__main__':main()
