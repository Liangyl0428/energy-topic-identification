"""Build a reviewable audit package, preserving the 990-class input release."""
from audit_metrics import *
from semantic_judgments import FINDINGS
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl import load_workbook
from collections import Counter

def clean_cell(x):
    if isinstance(x,(list,dict)):
        x=json.dumps(x,ensure_ascii=False)
    if isinstance(x,str):
        x=ILLEGAL_CHARACTERS_RE.sub('',x)
        if x.startswith(('=','+','-','@')):x="'"+x
    return x

def main():
    metrics=read(AUDIT/'results/METRICS.json')
    ev=read(AUDIT/'evidence/audit_samples.json')
    quality=read(AUDIT/'results/QUALITY_CENSUS.json')
    topics=pd.read_csv(BASE/'results/final_topics.csv').fillna('')
    t=topics.set_index('final_topic_id')
    evidence={r['row_id']:r for r in ev['records']}
    cases={r['case_id']:r for r in ev['cases']}
    y=np.load(BASE/'results/final_labels.npy',mmap_mode='r')
    raw=np.load(RAW/'results/raw_labels.npy',mmap_mode='r')
    read_chars={}
    for p in sorted((AUDIT/'evidence').glob('read_log*.json')):
        for r in read(p)['records']:
            read_chars[r['row_id']]=max(read_chars.get(r['row_id'],0),r.get('body_characters_shown',0))
    assert len(evidence)==len(ev['records'])==572
    assert set(evidence)==set(read_chars), 'Every evidence title must have a read log'
    for rid,r in evidence.items():
        assert y[rid]==t.loc[r['topic_id']].final_topic_index,(rid,'wrong current label')
        assert raw[rid]==r['raw_cluster'],(rid,'wrong raw label')
    sem=np.load(AUDIT/'results/semantic_cosines.npy')
    lex=np.load(AUDIT/'results/lexical_cosines.npy')
    flow=np.load(AUDIT/'results/secondary_flow_counts.npy')
    pair_rows=[]; case_rows=[]; cited=set()
    for f in FINDINGS:
        c=cases[f['case_id']]
        f['topics']=c['topics']; f['title']=c['title']
        assert all(rid in evidence for rid in f['evidence_row_ids']),f['case_id']
        assert all(evidence[rid]['topic_id'] in c['topics'] for rid in f['evidence_row_ids']),f['case_id']
        cited.update(f['evidence_row_ids'])
        for p in f['pair_judgments']:
            a,b=p['a'],p['b'];i,j=int(t.loc[a].final_topic_index),int(t.loc[b].final_topic_index)
            assert a in c['topics'] and b in c['topics']
            assert {evidence[rid]['topic_id'] for rid in p['evidence_row_ids']}=={a,b},(f['case_id'],a,b)
            pair_rows.append(dict(case_id=f['case_id'],title=c['title'],a=a,b=b,relationship=p['relationship'],
                confidence=p['confidence_in_overlap'],semantic_cosine=float(sem[i,j]),lexical_cosine=float(lex[i,j]),
                a_original_runner_up_b_fraction=float(flow[i,j]/t.loc[a].documents),
                b_original_runner_up_a_fraction=float(flow[j,i]/t.loc[b].documents),
                evidence_row_ids='|'.join(map(str,p['evidence_row_ids'])),
                evidence='\n'.join(f"{evidence[rid]['topic_id']} / row_id={rid} / {evidence[rid]['title']}" for rid in p['evidence_row_ids']),
                whole_class_synonym_equivalence_proven=False))
        case_rows.append(dict(case_id=f['case_id'],theme=c['title'],topics='|'.join(c['topics']),
            relationship=f['relationship'],confirmed_overlap=f['confirmed_overlap'],
            confirmed_pair_count=len(f['pair_judgments']),finding=f['summary'],limits=f['limits'],
            recommended_action=f['recommended_action'],evidence_row_ids='|'.join(map(str,f['evidence_row_ids']))))
    pairs=pd.DataFrame(pair_rows)
    assert not pairs[['a','b']].duplicated().any()
    findings=pd.DataFrame(case_rows)
    positive_topics=set(pairs.a)|set(pairs.b)
    near=pairs[pairs.relationship=='主主题近重复']
    summary=dict(created_utc=now(),input_directory=str(BASE),topic_count=990,all_pairs=metrics['all_pairs'],
        candidate_pairs=metrics['candidate_pairs'],case_count=len(findings),sampled_topics=ev['selected_topics'],
        sampled_documents=ev['total_records'],fresh_random_documents=ev['fresh_records'],
        existing_center_documents=ev['total_records']-ev['fresh_records'],titles_read=len(read_chars),
        nonempty_body_excerpts_read=sum(v>0 for v in read_chars.values()),cited_documents=len(cited),
        confirmed_overlap_cases=int(findings.confirmed_overlap.sum()),confirmed_overlap_pairs=len(pairs),
        topics_in_confirmed_overlap_pairs=len(positive_topics),near_duplicate_core_pairs=len(near),
        topics_in_near_duplicate_core_pairs=len(set(near.a)|set(near.b)),
        full_quality_marker_scan_topics=7,full_quality_marker_scan_documents=sum(c['documents'] for c in quality['topics']),
        labels_modified=False,whole_class_synonym_equivalence_proven=False,
        overall_overlap_rate_estimated=False,final_nonoverlapping_topic_count_estimated=False,
        review_type='Codex模型辅助逐项语义判读；不是独立人工专家审定',
        sample_design='风险定向选类；每当前类6条新增随机文献，另取原簇中心文献。不能推断全体990类的重合率。',
        input_delivery_manifest_sha256=metrics['input_delivery_manifest_sha256'])
    assert summary['confirmed_overlap_cases']==22
    assert summary['topics_in_confirmed_overlap_pairs']==57
    assert summary['near_duplicate_core_pairs']==4
    dump(AUDIT/'results/AUDIT_SUMMARY.json',summary)
    dump(AUDIT/'results/semantic_judgments.json',dict(review_type=summary['review_type'],findings=FINDINGS))
    findings.to_csv(AUDIT/'results/semantic_findings.csv',index=False,encoding='utf-8-sig')
    pairs.to_csv(AUDIT/'results/pair_findings.csv',index=False,encoding='utf-8-sig')

    inventory=pd.read_csv(AUDIT/'results/topic_audit_inventory.csv').fillna('')
    sampled={r['topic_id'] for r in ev['records']}
    inventory['this_audit_content_sampled']=inventory.final_topic_id.isin(sampled)
    inventory['this_audit_cases']=inventory.final_topic_id.map(lambda tid:'|'.join(c['case_id'] for c in ev['cases'] if tid in c['topics']))
    inventory['in_confirmed_overlap_pair']=inventory.final_topic_id.isin(positive_topics)
    inventory['audit_status']=inventory.final_topic_id.map(lambda tid:
        '已抽样并发现具体重合关系；不代表整类同义' if tid in positive_topics else
        '已抽样；详见案例中的混杂、任务差异或质量结论' if tid in sampled else
        '本轮仅全量数值筛查；未作内容裁决，不能认定无重合')
    inventory.to_csv(AUDIT/'results/topic_audit_inventory.csv',index=False,encoding='utf-8-sig')
    evidence_rows=[]
    for r in ev['records']:
        evidence_rows.append(dict(topic_id=r['topic_id'],row_id=r['row_id'],source=r['source'],title=r['title'],
            selection=r['selection'],cited_in_findings=r['row_id'] in cited,
            body_characters_read=read_chars[r['row_id']],body_excerpt_read=r['body_excerpt'][:read_chars[r['row_id']]],
            body_excerpt_saved=r['body_excerpt'],url=r.get('url',''),source_part=r['source_part'],body_sha256=r['body_sha256']))
    evidence_df=pd.DataFrame(evidence_rows)
    evidence_df.to_csv(AUDIT/'results/evidence_read_index.csv',index=False,encoding='utf-8-sig')

    methods=[
        ('审计结论','990类仍有主主题近重复、父子包含、范围交叉及输入污染，不能验收为互斥语义分类体系。'),
        ('主主题近重复','两类代表性文献共享对象和核心任务，未形成稳定边界；不是每条记录都同义，也不是直接执行整类合并。'),
        ('范围交叉','共享具体子主题、父子范围或使用不同分类维度；可通过层级、多标签或重新分配解决，不能一概合并。'),
        ('区分依据','不同研究对象、反应、部件或任务可支持分开；必须检查文献分配是否真正符合定义。'),
        ('语义相似度','全量文献原始缓存向量、按原训练来源权重重建990个中心。余弦是候选筛选指标，不是重复率或概率。'),
        ('词汇相似度','当前合并版c-TF-IDF向量的余弦。公共背景词、文体和模板均可能拉高相似度。'),
        ('第二候选','原1000类PCA64 KMeans的第二候选映射到990类；未重算990类近邻，不是分类概率。'),
        ('抽样限制',summary['sample_design']),
        ('阅读范围',f"{summary['titles_read']}条题名及{summary['nonempty_body_excerpts_read']}条非空正文节选；未阅读全文或专利权利要求。"),
        ('模板检查','七类35,413条做已知字符串/正则模式全量匹配；模式命中不等于所有记录均无研究价值。'),
        ('统计口径','22组为审计案例，不是22次合并；57类为已判定关系端点去重，不是57个应删除类；不对关系做传递闭包。'),
        ('未审类','未抽样的919类仅经过数值筛查，不默认通过语义验收。'),
        ('审阅身份',summary['review_type']),
        ('标签变更','本轮仅审计，不修改990类目录、主题模型或文献标签。'),
    ]
    sheets={
        '审计摘要':pd.DataFrame([{'item':k,'value':v} for k,v in summary.items()]),
        '逐组结论':findings,
        '具体重合关系':pairs,
        '文献证据':evidence_df,
        '990类目录':inventory,
        '5135候选对':pd.read_csv(AUDIT/'results/candidate_pairs.csv'),
        '7类文本质量':pd.read_csv(AUDIT/'results/quality_marker_census.csv'),
        '数值筛查统计':pd.DataFrame(metrics['threshold_screens']),
        '方法与限制':pd.DataFrame(methods,columns=['item','description']),
    }
    xlsx=AUDIT/'严格重合审计.xlsx'
    with pd.ExcelWriter(xlsx,engine='openpyxl') as writer:
        for name,df in sheets.items():
            safe=df.copy()
            for col in safe.columns:safe[col]=safe[col].map(clean_cell)
            safe.to_excel(writer,sheet_name=name,index=False)
            ws=writer.book[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='29445F')
            for cells in ws.iter_cols():
                values=[str(c.value or '') for c in list(cells)[:100]]
                ws.column_dimensions[cells[0].column_letter].width=min(72,max(15,max(map(len,values),default=15)+2))
            if name in ['逐组结论','具体重合关系','方法与限制']:
                for row in ws.iter_rows(min_row=2):
                    for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb=load_workbook(xlsx,read_only=True)
    for name,df in sheets.items():assert wb[name].max_row==len(df)+1,(name,'rows')
    wb.close()

    # Verify every file in the pre-existing release manifest, not merely its filename.
    manifest=read(BASE/'DELIVERY_MANIFEST.json')
    assert sha(BASE/'DELIVERY_MANIFEST.json')==metrics['input_delivery_manifest_sha256']
    altered=[]
    for f in manifest['files']:
        p=BASE/f['path']
        if not p.is_file() or p.stat().st_size!=f['bytes'] or sha(p)!=f['sha256']:altered.append(f['path'])
    assert not altered,altered
    validation=dict(created_utc=now(),evidence_current_label_checks=len(evidence),evidence_raw_label_checks=len(evidence),
        all_evidence_titles_have_read_logs=True,all_cited_rows_present=True,all_pairs_have_both_side_evidence=True,
        workbook_sheets_checked=len(sheets),input_manifest_files_checked=len(manifest['files']),
        input_release_all_hashes_unchanged=True,labels_modified=False)
    dump(AUDIT/'results/VALIDATION.json',validation)

    lines=[
        '# 990类严格重合审计（2026-09-25）','',
        '**结论：有，而且不止零星重合。当前990类不能验收为语义已去重、边界清晰、互斥的分类体系。**',
        '上一轮1000→990只是完成了10组归并及标签映射，没有消除全部语义重合。部分保留判断过于保守：把局部侧重点、父子范围差异或异类噪声当作独立类别的依据，却没有证明其边界能解释实际文献分配。','',
        f"本轮对全部990类计算了**{metrics['all_pairs']:,}对**指标；在风险定向选取的**71类、25个案例**中审读572条题名、155条非空正文节选，确认**22个问题组、{len(pairs)}条具体类别关系，涉及57类**存在主主题近重复或范围交叉。其中4对属于主主题近重复。另对7个疑似质量问题类的35,413条记录做了精确文本模式检查。",
        '**这些是已发现问题，不是全库重合率：22组不等于应该减少22类，57类不等于57个冗余类；未抽样的919类不默认通过。** 本审计不证明任一整类的每条文献完全同义，也不对关系作传递合并。','',
        '输入：[990类现有结果](../bertopic1000_semantic_merged_20260925/results/final_topics.csv)。本轮保持全部现有主题与文献标签不变；完整交付清单逐文件哈希复核通过。','',
        '## 最清楚的具体问题','',
        '|主题|类别ID|严格判定|', '|---|---|---|',
        '|TiO₂污染物光降解|S0308 / S0866|主主题近重复，中心文献均为改性TiO₂降解有机污染物|',
        '|OLED|S0393 / S0966|主主题近重复，两类均有exciplex、传输层和器件效率研究|',
        '|厌氧消化|S0772 / S0785|主主题近重复，餐厨垃圾、污泥、反应器及产甲烷交叉|',
        '|电价政策|S0619 / S0651|主主题近重复，销售与输配电价均跨两类分布|',
        '|钙钛矿电池|S0066 / S0220 / S0384 / S0787|界面、传输层、叠层与器件模拟交叉，原侧重点不能形成互斥边界|',
        '|聚合物电解质|S0049 / S0126 / S0940|固态聚合物锂电解质交叉，S0940更宽，不能整类等同|',
        '|有机光伏|S0140 / S0158|聚合物光伏属于OPV，宽类实际也收录聚合物设计|',
        '|电动汽车充电|S0253 / S0879|充电站选址与充放电调度在两类双向交叉|','',
        '类别ID为稳定ID，数字不是990类中的连续行号。完整逐组判定和具体row_id见下文及工作簿。','',
        '## 模板与占位文本的全类检查','',
        '|类别|精确命中|比例|含义|','|---|---:|---:|---|',
    ]
    q={r['topic_id']:r for r in quality['topics']}
    for tid,key,meaning in [('S0576','sigma_publisher_template','同一句SIGMA-NOT出版社宣传正文，类别主要按模板聚集'),
                            ('S0952','cheminform_template','ChemInform服务说明，研究题名可能有效但正文不是研究摘要'),
                            ('S0568','doi_only_title','完整标题匹配DOI格式，无法仅据该标题解释研究主题')]:
        r=q[tid];lines.append(f"|{tid}|{r[key]:,} / {r['documents']:,}|{r[key]/r['documents']:.2%}|{meaning}|")
    lines += ['',
        '上述三类的现有 `template_record` 标记全部为0，说明原有模板检测没有捕获这些具体模式。S0576的5,729条命中正文完全相同。S0205、S0316、S0700、S0888则存在主题分散、期刊封面/会议元数据等抽样证据；不能据此判定四类全部无效。',
        '需要先清除公共模板、补全题名/摘要，再重新分类；将这些簇简单合并为“其他”会继续掩盖问题。','',
        '## 全量筛查指标及其限制','',
        '|指标阈值|候选对数|涉及类数|','|---|---:|---:|']
    for s in metrics['threshold_screens']:
        label='主题中心语义余弦' if s['metric']=='semantic_cosine' else 'c-TF-IDF词汇余弦'
        lines.append(f"|{label} ≥ {s['threshold']:.2f}|{s['pairs']:,}|{s['topics']}|")
    lines += ['',
        '以上均为筛查数量，**不是已确认重复对数、重复概率或重合率**。936对语义余弦≥0.90只能说明有待核查的高相似关系，不能说有936对重复。',
        '以全部原始缓存文献向量和原训练来源权重（论文1、专利6.166948569140668、政策52.10973741558246）重建当前990类中心。复算原1000中心的最大绝对误差为1.05×10⁻⁸。词汇指标来自当前c-TF-IDF。5,135对候选包含每类语义/词汇前5近邻及语义≥0.90或词汇≥0.80的全部对。',
        '第二候选流向仅为原PCA64 KMeans第二候选映射，不是重新计算的990类后验概率。例如S0534/S0608双向第二候选比例为42.6%/51.7%，但对电极与染料仍有真实概念差别；需要解决的是混入彼此的文献，不能用阈值自动整类合并。',
        '原模型固定MiniBatchKMeans的K=1000；990是合并后的簇数，不是数据自然证明的独立语义类数。980类仍用数字+关键词自动名；上一版823类仅清单初筛、167类有抽样复核。单标签分配与14项工程一致性检查均不能替代语义互斥验收。',
        '759类同一来源占比≥95%也不能直接当作来源伪主题的证据，因为全库论文占比本就很高。原模型距离margin<0.02有1,137,516条，也仅是边界诊断，不能当作错分率。','',
        '## 判定标准与覆盖范围','',
        '主主题近重复：两类的研究对象和核心任务反复相同，未找到可以稳定解释样本的区分规则。包含/交叉：父子范围或材料、部件、方法、任务维度相交。可区分：对象、反应、部件或主任务存在实质差异；即使保留分开，仍需处理跨任务和错分记录。',
        '语义判读由Codex模型辅助完成，不是独立领域专家审定。71类为风险定向选取：每个当前类取6条未出现在先前证据集中的随机记录（种子2026092503及2026092504），另取各原簇两个中心样本，共426条新增随机记录及146条原中心记录。逐条题名和实际显示的正文字符数有阅读日志；未声称读完文献全文或专利权利要求。',
        '中心样本、随机样本、相似度、类名与已知任务定义共同用于定性判断。定向抽样不支持估计全库重合率或自动确定最终应剩多少类。','',
        '## 25个案例的逐组证据','']
    for f in FINDINGS:
        lines += [f"### {f['case_id']} {f['title']} — {f['relationship']}",'',
            '**类别：** '+', '.join(f['topics']), '',f['summary'],'',
            '**边界与反证：** '+f['limits'],'',
            '**建议：** '+f['recommended_action'],'',
            '|具体关系|语义余弦|词汇余弦|'] if f['pair_judgments'] else [
                f"### {f['case_id']} {f['title']} — {f['relationship']}",'',
                '**类别：** '+', '.join(f['topics']),'',f['summary'],'',
                '**边界与反证：** '+f['limits'],'','**建议：** '+f['recommended_action'],'']
        if f['pair_judgments']:
            lines.append('|---|---:|---:|')
            for p in pair_rows:
                if p['case_id']==f['case_id']:
                    lines.append(f"|{p['a']} / {p['b']}：{p['relationship']}|{p['semantic_cosine']:.4f}|{p['lexical_cosine']:.4f}|")
            lines.append('')
        lines += ['具体文献证据（题名保留原文；row_id可追溯原数据）：','']
        for rid in f['evidence_row_ids']:
            r=evidence[rid]
            lines.append(f"- {r['topic_id']} · row_id={rid} · {r['title']}")
        lines.append('')
    lines += ['## 处理优先级','',
        '1. 先修复模板/DOI占位与范围外记录，避免继续用污染文本定义主题。',
        '2. 优先复核并统一四对近重复核心：S0308/S0866、S0393/S0966、S0772/S0785、S0619/S0651；合并前清理异类记录并决定是否保留一致的细分规则。',
        '3. 对父子范围与多维交叉建立层级或多标签体系；需要单标签时必须给出叶子类定义和冲突裁决规则。',
        '4. 对919个本轮未抽样类别继续审查候选关系。当前证据不足以给出最终互斥类别数，不能以某个余弦阈值直接压缩到任意类数。','',
        '## 交付与校验','',
        '- [审计工作簿](严格重合审计.xlsx)：逐组结论、具体关系、文献证据、990类目录、5,135对候选、模板统计及方法限制。',
        '- [逐组结论CSV](results/semantic_findings.csv)；[具体类别关系CSV](results/pair_findings.csv)。',
        '- [全部489,555对指标](results/all_489555_pairs.parquet)；[990类审计目录](results/topic_audit_inventory.csv)。',
        '- [证据与抽样信息](evidence/audit_samples.json)；[阅读范围索引](results/evidence_read_index.csv)。',
        '- [机器可读摘要](results/AUDIT_SUMMARY.json)；[校验结果](results/VALIDATION.json)。','',
        f"已核对572条证据的当前/原始类别、所有具体关系的双边证据及{len(sheets)}个工作表的行数；输入交付清单{len(manifest['files'])}个文件的大小与SHA-256全部保持原值。**本轮未修改任何文献标签。**",'']
    (AUDIT/'REPORT.md').write_text('\n'.join(lines))
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    print('Input release verified unchanged:',len(manifest['files']),'files')

if __name__=='__main__':main()
