"""Publish only explicit, evidence-backed review decisions; never auto-merge pairs."""
from common import *
import collections, numpy as np, pandas as pd
import pyarrow as pa, pyarrow.parquet as pq
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

def markdown_table(frame):
    def cell(v): return str(v).replace('|','\\|').replace('\n',' ')
    rows=['| '+' | '.join(map(cell,frame.columns))+' |','| '+' | '.join('---' for _ in frame.columns)+' |']
    rows += ['| '+' | '.join(cell(v) for v in row)+' |' for row in frame.itertuples(index=False,name=None)]
    return '\n'.join(rows)

inp = json.loads((BASE / 'data/INPUT_MANIFEST.json').read_text())
cluster_audit = json.loads((BASE / 'models/CLUSTER_AUDIT.json').read_text())
review_method = json.loads((BASE / 'results/REVIEW_METHOD.json').read_text())
review = pd.read_csv(BASE / 'results/manual_reviews.tsv', sep='\t', keep_default_na=False)
assert set(review.cluster) == set(range(500)) and len(review) == 500
decisions = json.loads((BASE / 'results/merge_decisions.json').read_text())
raw = pd.read_csv(BASE / 'results/raw_500_topics.csv')
evidence = pd.read_parquet(BASE / 'evidence/review_documents.parquet')
known_evidence = set(evidence.row_id)
evidence_clusters = dict(zip(evidence.row_id,evidence.cluster))
mapping = {i: i for i in range(500)}; merged_names = {}; rationales = {}; taken = set(); merge_metadata = {}
for decision in decisions:
    members = decision['members']
    assert len(members) >= 2 and len(set(members)) == len(members)
    assert set(members) <= set(range(500)) and not taken.intersection(members)
    assert decision['reason'].strip() and decision['name'].strip()
    assert set(decision['evidence_row_ids']) <= known_evidence
    assert set(members) <= {evidence_clusters[rid] for rid in decision['evidence_row_ids']}
    taken.update(members); target = min(members)
    for i in members: mapping[i] = target
    merged_names[target] = decision['name']; rationales[target] = decision['reason']
    merge_metadata[target] = decision

review = raw.merge(review, on='cluster', validate='one_to_one')
review['merged_topic_id'] = [f'M{mapping[i]:03d}' for i in review.cluster]
review['merged_name'] = [merged_names.get(mapping[i], review.loc[review.cluster.eq(mapping[i]), 'reviewed_name'].iloc[0]) for i in review.cluster]
review['merge_reason'] = [rationales.get(mapping[i], '保留独立类群') for i in review.cluster]
review.to_csv(BASE / 'results/cluster_reviews.csv', index=False)
review['merged_domain'] = [merge_metadata.get(mapping[i],{}).get('domain',review.loc[review.cluster.eq(mapping[i]),'domain'].iloc[0]) for i in review.cluster]
review['merged_status'] = [merge_metadata.get(mapping[i],{}).get('status',review.loc[review.cluster.eq(mapping[i]),'status'].iloc[0]) for i in review.cluster]
review.to_csv(BASE / 'results/cluster_reviews.csv', index=False)
review[['raw_topic_id','cluster','reviewed_name','merged_topic_id','merged_name','domain','status','merged_domain','merged_status','merge_reason']].to_csv(BASE / 'results/raw_to_merged.csv', index=False)

final_rows = []
for target in sorted(set(mapping.values())):
    group = review[review.cluster.map(mapping).eq(target)]
    final_rows.append({'merged_topic_id':f'M{target:03d}', 'name':group.merged_name.iloc[0], 'domain':group.merged_domain.iloc[0], 'status':group.merged_status.iloc[0], 'raw_clusters':','.join(group.raw_topic_id), 'raw_cluster_count':len(group), 'documents':int(group.documents.sum()), 'papers':int(group.papers.sum()), 'patents_2026':int(group.patents_2026.sum()), 'policies':int(group.policies.sum()), 'review_note':'；'.join(dict.fromkeys(group.review_note)), 'merge_reason':group.merge_reason.iloc[0]})
final = pd.DataFrame(final_rows)
final.to_csv(BASE / 'results/merged_topics.csv', index=False)

destination = BASE / 'results/merged_assignments'; destination.mkdir(exist_ok=True)
label_lookup = np.array([mapping[i] for i in range(500)], dtype=np.int16)
assignment_counts = collections.Counter(); final_counts = collections.Counter()
for part in sorted((BASE / 'results/assignments').glob('part-*.parquet')):
    table = pq.read_table(part); labels = table.column('raw_cluster').to_numpy(); assigned = labels >= 0
    merged = np.full(len(labels), -2, dtype=np.int16); merged[assigned] = label_lookup[labels[assigned]]
    table = table.append_column('merged_cluster',pa.array(merged)).append_column('merged_topic_id',pa.array([f'M{x:03d}' if x>=0 else 'UNUSABLE' for x in merged]))
    pq.write_table(table, destination / part.name, compression='zstd')
    for source, label in zip(table.column('source').to_pylist(), merged): assignment_counts[(source,'assigned' if label>=0 else 'unusable')] += 1
    final_counts.update(int(x) for x in merged if x >= 0)
    if len(list(destination.glob('part-*.parquet'))) % 50 == 0:
        print('MAPPED',part.name,flush=True)
assert all(final_counts[int(r.merged_topic_id[1:])] == r.documents for r in final.itertuples())
coverage = []
for source,total in inp['source_counts'].items():
    assigned = assignment_counts[(source,'assigned')]
    coverage.append({'source':source,'unique_documents':total,'assigned':assigned,'unusable':assignment_counts[(source,'unusable')],'assignment_rate':assigned/total,'meaning':'分配率；不是人工校验准确率'})
coverage = pd.DataFrame(coverage)
coverage.to_csv(BASE/'results/coverage.csv',index=False)

evidence = evidence.merge(review[['cluster','reviewed_name','merged_topic_id','merged_name']],on='cluster',validate='many_to_one')
evidence.to_parquet(BASE/'evidence/review_documents_with_topics.parquet',index=False)
excel_evidence = evidence.copy(); excel_evidence['body']=excel_evidence.body.str.slice(0,2000)
decision_frame = pd.DataFrame([{'merged_topic_id':f'M{min(x["members"]):03d}','name':x['name'],'status':x['status'],'raw_topics':','.join(f'B{i:03d}' for i in x['members']),'reason':x['reason'],'evidence_row_ids':','.join(map(str,x['evidence_row_ids']))} for x in decisions])
retained_frame = pd.read_csv(BASE/'results/retained_overlap_cases.csv',keep_default_na=False)
raw_status_counts=review.groupby('status').agg(raw_clusters=('cluster','size'),documents=('documents','sum')).reset_index()
status_counts=final.groupby('status').agg(merged_groups=('merged_topic_id','size'),documents=('documents','sum')).reset_index()
status_counts.to_csv(BASE/'results/status_summary.csv',index=False)
domain_counts=final.groupby('domain').agg(merged_groups=('merged_topic_id','size'),documents=('documents','sum'),papers=('papers','sum'),patents_2026=('patents_2026','sum'),policies=('policies','sum')).reset_index()
domain_counts.to_csv(BASE/'results/domain_summary.csv',index=False)
inspected=json.loads((BASE/'evidence/excerpts_inspected.json').read_text())
detail_frame=pd.DataFrame(inspected).sort_values('body_chars_shown').drop_duplicates('row_id',keep='last').sort_values(['cluster','row_id'])
detail_frame.to_csv(BASE/'evidence/inspected_excerpts.csv',index=False)
method_rows = [
('数据范围','全部已有文献和政策；公开年份2026的全部已获取专利；不是完整2026全年。'),
('去重','文献沿用作品/版本归并；专利按国家地区+申请号；政策按URL/文号等。保留原始映射账本。'),
('缺失专利文本恢复','2941件无摘要专利从现有权利要求/说明书恢复正文，其中210件同时缺标题；原始输入和恢复账本均保留。'),
('时间范围',json.dumps(inp['date_ranges'],ensure_ascii=False)),
('语义模型','本地多语言MiniLM；所有已有文本tokens分段参与；每段<=128tokens；保存384维文档向量。文献使用已有标题摘要，并非PDF全文。'),
('政策向量','为防超长附件清单淹没政策叙事，采用前4096 tokens正文向量60% + 全文向量40%；没有把附件完全丢弃。'),
('BERTopic配置','PCA64+MiniBatchKMeans500+全量c-TF-IDF；非默认UMAP/HDBSCAN。'),
('训练','全量PCA统计；分来源初始化后3遍全部可用文档训练；文献/专利/政策目标权重70%/25%/5%。'),
('分配率与准确率','固定500类的标签覆盖率不能当作准确率；未声称达到招标要求90%。'),
('审阅范围',f'500簇初审为主题词和{review_method["initial_unique_documents"]}篇抽样题名；另查{review_method["detailed_clusters"]}簇的{review_method["detailed_unique_documents"]}篇证据，其中{review_method["detailed_documents_with_body"]}篇有可读正文片段。非全部文档逐篇人工复核。'),
('审阅者','本会话助手根据记录的题名和正文片段审阅；不是独立行业专家验收。'),
('未决重叠','3016个相似候选对未逐对全部裁决；未合并不代表已证明互斥。保留案例与其理由另表列出。'),
('单标签与粒度','每个文档只分到一个簇；复合政策和跨学科文献没有多标签。合并后粒度不统一，混杂和关联基础类群仍保留。'),
('整簇合并边界','只改变类群映射，不悄悄逐篇重分；宽泛、混杂和外围类群保留状态提示。'),
('热点使用限制','本轮交付主题结构，不是核心/新兴/潜在热点排名；政策全文跨度多年，后续潜在热点需另筛2026。'),
('期刊视角','期刊专题、高被引等需要期刊/引用元数据形成附加视角，不由语义聚类自动等价实现。'),
('原文可追溯','evidence/review_documents_with_topics.parquet保存完整已有文本；Excel正文仅展示前2000字符。'),
('向量落盘','models/embeddings逐批保存；results/merged_assignments保存全部文档标签。')]
book=BASE/'500类BERTopic聚类与审阅合并.xlsx'
def excel_safe(value):
    return ILLEGAL_CHARACTERS_RE.sub('',value)[:32767] if isinstance(value,str) else value
with pd.ExcelWriter(book,engine='openpyxl') as writer:
    for name,frame in [('合并后类群',final),('原始500簇审阅',review),('整簇合并依据',decision_frame),('保留与待重分类群',retained_frame),('覆盖率',coverage),('类群状态统计',status_counts),('领域分布',domain_counts),('实际补查片段',detail_frame),('抽样原文证据',excel_evidence),('方法与口径',pd.DataFrame(method_rows,columns=['项目','说明']))]:
        frame.apply(lambda col:col.map(excel_safe)).to_excel(writer,sheet_name=name,index=False)
        ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for column in ws.columns:
            key=column[0].column_letter
            width=min(65,max(12,max(len(str(c.value or '')) for c in list(column)[:100])+2))
            ws.column_dimensions[key].width=width
        for cell in ws[1]:
            from openpyxl.styles import Font,PatternFill
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245B78')
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type=='f':cell.data_type='s'

review_record={'raw_clusters':500,'merged_groups':len(final),'merge_groups':len(decisions),'removed_duplicate_cluster_count':500-len(final),'all_500_reviewed':True,'review_scope':review_method,'decisions_sha256':sha(BASE/'results/merge_decisions.json'),'manual_reviews_sha256':sha(BASE/'results/manual_reviews.tsv'),'completed_utc':now()}
dump(BASE/'results/MERGE_REVIEW.json',review_record)
total_assigned=int(coverage.assigned.sum());total=int(coverage.unique_documents.sum())
lines=['# 全来源 BERTopic 500 类聚类与整簇审阅合并','',f'本轮处理 {total:,} 个去重后的文档单位，得到 500 个原始簇；审阅后以 {len(decisions)} 组合并决定形成 {len(final)} 个类群。所有原始簇均保留可逆映射。','',
'## 数据及覆盖口径','',markdown_table(coverage),'',f'分配成功 {total_assigned:,}/{total:,}（{total_assigned/total:.4%}）。这表示文本获得了固定簇数模型的标签，**不表示分类准确率**。本轮未用独立金标准证明90%准确率。','',
f'文献日期：{inp["date_ranges"]["paper"]}；专利日期：{inp["date_ranges"]["patent"]}；政策日期：{inp["date_ranges"]["policy"]}。2026年专利为库中已抓取部分，截止8月31日；政策截止9月23日。',
'','原始公开专利记录与归并后的申请单位分别保存在专利账本中。没有把多份公开文本或政策编码分段重复当作独立文档。撤稿、模板化记录、只有标题的记录保留质量标记。','',
'本次采用固定簇数模型，每条可用文本获得一个标签。不能把99%以上分配率与旧版弱监督接受率直接比较并宣称准确率已改善，也不能据此断言所有符合招标范围的文献均已正确归类。','',
'## 审阅结果','',
f'合并后共{len(final)}类，包含方向候选、宽泛/混杂类群、关联基础研究及外围文本，并非{len(final)}个同样细、互不重叠的正式能源方向。下表统计合并后的类群状态；原始500簇状态另存`results/cluster_reviews.csv`。','',markdown_table(status_counts),'',
'审阅区分技术方向、政策治理、关联基础研究、范围外噪声及混杂类群。500类是探索性类群；若一个原始簇已混合多种研究对象，整簇合并无法让它变成统一细粒度主题，因此保留明确状态。没有用改名或隐蔽逐文档重分掩盖这一限制。','',
'例如热电材料、钙钛矿光伏、废旧锂电池回收、染料敏化太阳能电池等有明确的重复类群合并依据；欧洲一般政治混入、通用机械装配混杂及电力预测/计算混杂保留待重分。不同技术的父子关系和跨应用材料交叠仍然存在。','',
'合并清单：','',markdown_table(decision_frame[['merged_topic_id','name','raw_topics','status']]),'',
'合并仅针对审阅证据支持的同一研究对象与任务，逐项理由在工作簿“整簇合并依据”及 `results/merge_decisions.json`。语义相似度、词频相似度和共同边界仅用于候选筛查，不直接决定合并。','',
f'实际审阅范围：先查看全部500簇的主题词及{review_method["initial_unique_documents"]}篇中心/随机等样本题名，再对{review_method["detailed_clusters"]}簇补查{review_method["detailed_unique_documents"]}篇证据，其中{review_method["detailed_documents_with_body"]}篇具有已展示的正文片段。初审不等于核读500簇的全部摘要，12,508篇留存证据也不等于全部被逐篇阅读。`evidence/inspected_excerpts.csv`记录实际补查内容及长度。审阅由本会话助手完成，未调用外部模型推理API，非独立行业专家验收。','',
'对3016个相似候选对尚未逐对全部裁决；无合并决定的类群不能解释为已通过去重验收。`results/retained_overlap_cases.csv`列出重点保留案例及审阅深度。需要互斥、统一细粒度主题时，后续应仅对明确的宽泛/混杂簇再细分、逐文档核分；本轮未执行该额外步骤。','',
'## 可复现方法','',
'使用 BERTopic 0.17.3，本地多语言MiniLM语义向量、全量加权PCA64、MiniBatchKMeans500和全量c-TF-IDF。固定500类使用了BERTopic可替换组件，不是默认HDBSCAN自动发现500类。全部可用文档参与PCA统计和3轮聚类训练；词表样本只用于控制词汇规模，主题词频累计使用全量文本。','',
'文献使用库中已有标题和摘要，不代表PDF全文。已有长文本全部分段进入模型，按新覆盖内容长度聚合；政策采用前4096 tokens叙事向量60% + 全文向量40%，避免超长项目附件主导语义，但全部附件仍参与。模型权重从公共仓库下载，推理在本机完成，未调用外部大模型接口。原401主题没有用作训练标签且文件校验和保持不变。每文档采用单一类群标签，跨主题政策未自动变为多标签。','',
'流式词频与直接求和、固定类别映射、模型重载预测已做针对性校验。中英技术短句匹配只是功能检查，不是实际文献的分类准确率评测。原始证据可在JSONL/Parquet中追溯，Excel正文为便于阅读截取前2000字符。','',
'## 交付文件','',
'- `500类BERTopic聚类与审阅合并.xlsx`：原始500簇、合并后类群、合并理由、覆盖率与原文证据。',
'- `results/raw_to_merged.csv`：500簇到合并类群的一一映射。',
'- `results/retained_overlap_cases.csv`：高相似或相近类群仍予保留的具体原因及证据。',
'- `results/status_summary.csv`、`results/domain_summary.csv`：类群状态及17个展示领域统计；领域归组不是行业标准认证。',
'- `results/merged_assignments/*.parquet`：全部文档的原始与合并标签；原始无语义文本标记为UNUSABLE。',
'- `models/bertopic500.joblib`、`models/embeddings/`：可重载模型和全部落盘向量。',
'- `data/INPUT_MANIFEST.json`：快照、日期、去重和全部输入分片校验和。','',
'本轮完成的是类群结构及审阅合并，不是新一轮热点榜单。后续核心/新兴热点应只用文献计量，潜在热点应重新筛选2026专利与2026政策；行政公告、撤稿、模板和范围外噪声不能直接作为热点证据。期刊聚焦视角仍需合并期刊、引用等元数据。','']
(BASE/'REPORT.md').write_text('\n'.join(lines))
checks={'raw_clusters_exactly500':len(raw)==500,'review_covers500_once':len(review)==500,'merge_groups_disjoint':True,'merge_evidence_ids_exist':True,'all_document_assignments_accounted':sum(assignment_counts.values())==total,'assigned_counts_equal_topic_counts':int(final.documents.sum())==total_assigned,'whole_cluster_mapping_only':True,'source_assignment_counts_match_manifest':all(r.assigned+r.unusable==r.unique_documents for r in coverage.itertuples()),'all_source_parts_written':len(list(destination.glob('part-*.parquet')))==len(list((BASE/'results/assignments').glob('part-*.parquet'))),'old_401_unchanged':sha(ROOT/'hhhh/revised_v2/results/taxonomy.csv')=='94bcaaab8722bab67c6e4278bee94f0cb98acc8d903a6046061bdeb3e97dac66'}
dump(BASE/'VALIDATION.json',{'pass':all(checks.values()),'checks':checks,'not_validated':['document-level semantic accuracy','equal granularity','absence of all taxonomy overlaps','tender acceptance'],'created_utc':now()})
assert all(json.loads((BASE/'VALIDATION.json').read_text())['checks'].values())
dump(BASE/'PROGRESS.json',{'stage':'complete','documents':total,'raw_clusters':500,'merged_groups':len(final),'updated_utc':now()})
print('PACKAGED',total,len(final),len(decisions),flush=True)
