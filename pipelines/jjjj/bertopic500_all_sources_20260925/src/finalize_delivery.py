"""Read back the published assignments and workbook, then finalize readable artifacts."""
from common import *
import collections
import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

review = json.loads((BASE/'results/MERGE_REVIEW.json').read_text())
manifest = json.loads((BASE/'data/INPUT_MANIFEST.json').read_text())
mapping = pd.read_csv(BASE/'results/raw_to_merged.csv')
lookup = np.array([int(mapping.loc[mapping.cluster.eq(i),'merged_topic_id'].iloc[0][1:]) for i in range(500)],dtype=np.int16)
seen = 0; counts = collections.Counter(); hashes = []
for original in sorted((BASE/'results/assignments').glob('part-*.parquet')):
    published = BASE/'results/merged_assignments'/original.name
    a = pq.read_table(original)
    b = pq.read_table(published)
    assert a.schema.equals(b.select(a.column_names).schema), f'Original schema changed: {original.name}'
    for col in a.column_names:
        left=a.column(col).to_numpy(); right=b.column(col).to_numpy()
        # Arrow equality treats NaN as unequal; missing-text distances intentionally contain NaN.
        equal=np.array_equal(left,right,equal_nan=True) if left.dtype.kind=='f' else a.column(col).equals(b.column(col))
        assert equal, f'Original column changed: {original.name}/{col}'
    rid = b.column('row_id').to_numpy()
    assert np.array_equal(rid,np.arange(seen,seen+len(b))), original.name
    raw = b.column('raw_cluster').to_numpy()
    expected = np.full(len(b),-2,dtype=np.int16)
    valid = raw>=0; expected[valid]=lookup[raw[valid]]
    assert np.array_equal(expected,b.column('merged_cluster').to_numpy()), original.name
    assert b.column('merged_topic_id').to_pylist()==[f'M{i:03d}' if i>=0 else 'UNUSABLE' for i in expected],original.name
    counts.update(int(x) for x in expected)
    seen += len(b)
    hashes.append(dict(part=original.name,documents=len(b),sha256=sha(published)))
assert seen==manifest['documents']
final = pd.read_csv(BASE/'results/merged_topics.csv')
assert all(counts[int(r.merged_topic_id[1:])]==r.documents for r in final.itertuples())
assert counts[-2]==2391
print('READBACK',seen,'documents; original columns unchanged; complete whole-cluster mapping',flush=True)

# Chinese headers improve the reading copy; machine-readable CSV/Parquet schemas stay stable.
headers = {
 'merged_topic_id':'合并类群编号','name':'类群名称','domain':'领域归组','status':'审阅状态',
 'raw_clusters':'原始类群编号','raw_cluster_count':'合入原簇数','documents':'文档数',
 'papers':'文献数','patents_2026':'2026专利申请数','policies':'政策数','review_note':'审阅说明','merge_reason':'合并理由',
 'raw_topic_id':'原始类群编号','cluster':'原簇数值编号','reviewed_name':'审阅命名',
 'merged_name':'合并后名称','merged_domain':'合并后领域','merged_status':'合并后状态',
 'keywords':'原始主题词','training_label':'训练标签','mean_cosine':'平均簇心相似度','std_cosine':'簇心相似度标准差',
 'cosine_below_045':'簇心相似度低于0.45记录数','margin_below_002':'前两簇距离差小于0.02记录数',
 'title_only_documents':'仅题名记录数','retracted_documents':'撤稿标记记录数','template_documents':'模板标记记录数',
 'reviewed_row_ids':'实际审阅样本行号','detailed_evidence_count':'补查证据篇数',
 'raw_topics':'原始类群','reason':'具体理由','evidence_row_ids':'证据行号',
 'members':'相关原簇','decision':'审阅决定','reason_type':'保留原因类型','review_depth':'审阅深度',
 'source':'数据来源','unique_documents':'去重后文档数','assigned':'已分配数','unusable':'无可用文本数',
 'assignment_rate':'标签分配率','meaning':'统计口径','merged_groups':'合并后类群数',
 'row_id':'文档行号','doc_id':'文档编号','title':'题名','body':'已有正文或摘要',
 'date':'日期','language':'语言','url':'原文链接','source_identifier':'来源记录标识',
 'title_only':'仅题名','retracted':'撤稿标记','template_record':'模板标记','usable':'文本可用',
 'publisher':'出版方','quality_note':'文本质量说明','selection_roles':'抽样角色',
 'secondary_candidate':'次近候选簇','distance_margin':'前两簇距离差',
 'body_chars_shown':'展示正文字数','body_chars_total':'已有正文总字数','excerpt':'实际补查片段','inspected_utc':'补查记录时间',
}
book = BASE/'500类BERTopic聚类与审阅合并.xlsx'
wb=load_workbook(book)
expected_rows={'合并后类群':len(final),'原始500簇审阅':500,'整簇合并依据':review['merge_groups'],'覆盖率':3,'抽样原文证据':12508,'实际补查片段':249}
for name,n in expected_rows.items():
    assert wb[name].max_row==n+1, (name,wb[name].max_row,n)
for ws in wb.worksheets:
    for c in ws[1]:
        c.value=headers.get(c.value,c.value)
    for row in ws.iter_rows(min_row=2):
        assert all(c.data_type!='f' for c in row), f'Unexpected spreadsheet formula in {ws.title}'
    if ws.title=='覆盖率':
        for row in ws.iter_rows(min_row=2):
            row[0].value={'paper':'文献','patent':'2026年专利申请','policy':'政策文本'}.get(row[0].value,row[0].value)
            row[4].number_format='0.0000%'
wb.save(book)

def table(frame):
    def fmt(x): return str(x).replace('|','\\|').replace('\n',' ')
    lines=['| '+' | '.join(map(fmt,frame.columns))+' |','| '+' | '.join(['---']*len(frame.columns))+' |']
    lines+=['| '+' | '.join(map(fmt,row))+' |' for row in frame.itertuples(index=False,name=None)]
    return '\n'.join(lines)

report_path=BASE/'REPORT.md';report=report_path.read_text()
cov=pd.read_csv(BASE/'results/coverage.csv')
friendly=cov.drop(columns=['meaning']).rename(columns=headers)
friendly['数据来源']=friendly['数据来源'].map({'paper':'文献','patent':'2026年专利申请','policy':'政策文本'})
for c in ['去重后文档数','已分配数','无可用文本数']: friendly[c]=friendly[c].map(lambda x:f'{x:,}')
friendly['标签分配率']=friendly['标签分配率'].map(lambda x:f'{x:.4%}')
if '| source |' in report:
    start=report.index('| source |');end=report.index('\n\n',start)
    report=report[:start]+table(friendly)+report[end:]
st=pd.read_csv(BASE/'results/status_summary.csv').rename(columns=headers)
st['文档数']=st['文档数'].map(lambda x:f'{x:,}')
if '| status |' in report:
    start=report.index('| status |');end=report.index('\n\n',start)
    report=report[:start]+table(st)+report[end:]
report=report.replace('| merged_topic_id | name | raw_topics | status |','| 合并类群编号 | 类群名称 | 原始类群 | 审阅状态 |')
for source,label in [('paper','文献'),('patent','专利'),('policy','政策')]:
    dates=manifest['date_ranges'][source]
    report=report.replace(str(dates),f'{dates["min"]} 至 {dates["max"]}')
report=report.replace('500簇到合并类群的一一映射','500簇到合并类群的逐簇映射（允许多对一）')
if '类群状态按内容标记' not in report:
    report=report.replace('下表统计合并后的类群状态；','类群状态按内容标记，不是数据来源分类；政策类群中也可能有政策研究文献。下表统计合并后的类群状态；')
report_path.write_text(report)

validation=json.loads((BASE/'VALIDATION.json').read_text())
validation['checks'].update(dict(all_written_original_columns_match=True,all_written_row_ids_contiguous=True,all_written_merge_labels_match_lookup=True,workbook_row_counts_match=True,workbook_has_no_formulas=True))
validation['pass']=all(validation['checks'].values());validation['readback_completed_utc']=now()
dump(BASE/'VALIDATION.json',validation)
dump(BASE/'results/MERGED_ASSIGNMENT_MANIFEST.json',dict(documents=seen,parts=hashes,completed_utc=now()))
artifacts=['500类BERTopic聚类与审阅合并.xlsx','REPORT.md','VALIDATION.json','results/raw_to_merged.csv','results/merged_topics.csv','results/merge_decisions.json','results/MERGE_REVIEW.json','results/REVIEW_METHOD.json','results/manual_reviews.tsv','results/retained_overlap_cases.csv','results/coverage.csv','results/status_summary.csv','results/domain_summary.csv','results/MERGED_ASSIGNMENT_MANIFEST.json','evidence/inspected_excerpts.csv']
dump(BASE/'DELIVERY_MANIFEST.json',dict(completed_utc=now(),documents=seen,raw_clusters=500,merged_groups=len(final),merge_decision_groups=review['merge_groups'],artifacts=[dict(path=f,bytes=(BASE/f).stat().st_size,sha256=sha(BASE/f)) for f in artifacts]))
print('FINALIZED',book,flush=True)
