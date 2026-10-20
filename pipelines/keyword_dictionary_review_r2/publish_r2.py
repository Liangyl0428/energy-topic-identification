"""Publish the frozen round-two diagnostic review without changing v1."""
from common import *
import shutil, importlib.util, html
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OLD = ROOT/'work/keyword_dictionary_multilabel_20261019'
V1_PUBLISH = HERE.parent/'keyword_dictionary/publish.py'
spec=importlib.util.spec_from_file_location('v1_presentation',V1_PUBLISH)
base=importlib.util.module_from_spec(spec);spec.loader.exec_module(base)

def readjson(name, root=WORK):
    return json.loads((root/name).read_text())

def records(frame):
    return [{k:(v.tolist() if hasattr(v,'tolist') else v) for k,v in r.items()}
            for r in frame.fillna('').to_dict('records')]

def workbook(sheets):
    wb=Workbook(write_only=True)
    for name,frame in sheets:
        ws=wb.create_sheet(name);ws.freeze_panes='A2';header=[]
        for label in frame.columns:
            cell=WriteOnlyCell(ws,value=label)
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='184E58')
            header.append(cell)
        ws.append(header)
        for row in frame.itertuples(index=False,name=None):ws.append([base.safe(x) for x in row])
        ws.auto_filter.ref=f'A1:{get_column_letter(len(frame.columns))}{len(frame)+1}'
    wb.save(ASSETS/'第二轮语义审核.xlsx')

def make_manual(stats, ds, review, comparison):
    n=stats['counts'];diff=n['relations']-readjson('CLASSIFIED.json',OLD)['counts']['relations']
    rows=''.join(f'<tr><td>{html.escape(str(r.name))}</td><td>{int(r.documents_before):,}</td><td>{int(r.documents):,}</td><td>{int(r.document_delta):+,}</td></tr>'
                 for r in comparison.assign(abs_delta=comparison.document_delta.abs()).sort_values('abs_delta',ascending=False).head(12).itertuples())
    return f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>第二轮语义审核：结果与接入说明</title><style>body{{max-width:1050px;margin:35px auto;padding:0 22px;font:16px/1.85 system-ui;color:#23434c}}h1,h2{{color:#124b57}}a{{color:#087d8b}}.note{{background:#fff2d7;padding:17px;border-left:4px solid #cc8b22}}table{{border-collapse:collapse;width:100%}}th,td{{padding:9px;border:1px solid #cedcde;text-align:left}}pre{{background:#eef4f4;padding:16px;white-space:pre-wrap;overflow-wrap:anywhere}}code{{overflow-wrap:anywhere}}.scroll{{overflow:auto}}</style>
<h1>第二轮语义审核：结果与接入说明</h1>
<p>版本：{stats['release_id']}。在第一版词典的基础上继续审读真实词汇，并用题名、可用摘要检查具体误分。由助手整理，尚未专家审定。</p>
<p><a href="关键词主题浏览.html">主题、词典和样本浏览</a> · <a href="第二轮语义审核.xlsx">审核工作簿</a> · <a href="REVIEW.json">审核统计</a> · <a href="platform_data.json">接口摘要</a></p>
<h2>本轮完成了什么</h2>
<p>逐项审读 {review['term_name_reviewed']:,} 个词名，其中 {review['newly_read_terms']:,} 个原先未读、{review['old_terms_revisited']:,} 个为既有词复核。审读队列包括高频未读词、含能源候选字样的词、照明设备候选词及旧辅助词。<strong>读过词名不等于已解决每篇论文中的词义。</strong>多数通用词保留为不触发分类的辅助词，review_scope 明确标记需语境复核。</p>
<p>逐项检查 {review['fresh_topic_samples']} 条新主题关系样本（每主题2条）和 {review['new_direct_term_samples']} 条新增直接词实例，共181条。对疑点阅读可用摘要；无摘要的记录保留“证据不足”。样本避开了第一轮抽样，但之后用于修正规则，故属于诊断样本，不能据此估计总体准确率。</p>
<ul><li>High voltage、Insulator (electricity) 改为需要相关电气设备锚点；Algorithm、Nonlinear system、Automatic frequency control 等泛化词停用或转为辅助词。</li>
<li>Power station 从火电移入通用供能系统，避免水电站被归入火电；Computer network 停止触发能源通信主题。</li>
<li>收紧电极/电解质和功率半导体的相关方向锚点；增加母线、开关设备、电机控制、油气开采等具体词。</li>
<li>新增 KWD-T0072“电光源、照明与照明节能”；LED/OLED 等泛词须有明确照明锚点。实验用灯只说明电光源关联，不能当作照明节能研究成果。</li>
<li>对7个具体论文关键词撤销触发，包括种马论文的 Wind power、光伏退役的 Nuclear decommissioning、轻轨语境的 Electric light、稀土离子回收的 Supercapacitor 等。保留原始来源词和撤销理由。</li></ul>
<p>上一轮9条被标为疑似误分或过度推断的关系已有8条不再出现，D137仍待核。另补正上一轮D028的漏判：该文摘要是光伏退役，原“题名支持”评语不正确，本轮撤销其核电退役触发。</p>
<h2>全量结果</h2>
<p>仍为 {n['papers']:,} 篇冻结论文、{ds['raw_keyword_ids']:,} 个来源关键词ID、{ds['canonical_keywords']:,} 个规范词条。当前 {ds['topics']} 个主题，{ds['decisions']['direct']} 个直接词、{ds['decisions']['conditional']} 个条件词；{n['assigned_papers']:,} 篇命中，{n['multilabel_papers']:,} 篇属于多个主题，共 {n['relations']:,} 条论文—主题关系（较第一版净变化 {diff:+,} 条）。这个变化同时含新增和撤销，不能视为纠错数。</p>
<p class="note">仍有 {ds['decisions']['unreviewed']:,} 个词尚未逐项审读。未读词的论文出现次数合计从 {review['unread_keyword_occurrence_mass_before']:,} 降至 {review['unread_keyword_occurrence_mass_after']:,}，该值按词累加，同一论文会重复计数。全量规则和来源完整性校验通过，<strong>不等于全库语义准确率已验收</strong>。</p>
<h2>篇数变化较大的主题</h2><div class="scroll"><table><tr><th>主题（本轮名称）</th><th>第一版</th><th>本轮</th><th>净变化</th></tr>{rows}</table></div>
<p>完整对照见 <a href="topic_count_comparison.csv">topic_count_comparison.csv</a>。主题之间有重叠，不可将各主题篇数相加作为独立论文总数。主题代码沿用 KWD-T，含义和词表调整按本轮定义解释。</p>
<h2>规则和证据如何使用</h2><ol>
<li>使用冻结版本的完整来源关键词；OpenAlex标签不保证是作者关键词。通过关键词ID和规范映射精确匹配，不使用模糊子串匹配。</li>
<li>明确直接词触发主题；具体技术条件词须同篇存在规定相关方向的直接词。通用支撑主题允许任一直接能源锚点，具体列表见 topic_definitions.json。</li>
<li>仅明确标为 withhold_trigger 的逐篇词不参与触发或锚点计算；needs_context 保留待核。来源 observations 没有删除或补造关键词。</li>
<li>取所有匹配主题并集，不设贡献阈值或主题上限。通用词、学科大类、未读词不触发；未命中不等于能源无关。</li>
<li>浏览器中的样本命中词来自抽样时快照；“当前归属保留/复核后已撤销”和“当前命中词”反映最终版本。新增词实例还显示该词本身是否撤销，其他词仍可能维持粗主题归属。</li></ol>
<h2>开发文件</h2><div class="scroll"><table><tr><th>文件</th><th>用途</th></tr>
<tr><td>keyword_dictionary.csv / keyword_aliases.csv</td><td>完整词典及来源ID映射；含决策、依据和本轮审核范围。</td></tr>
<tr><td>term_semantic_review.csv</td><td>本轮2472个词的审读结论、原状态、审核范围。</td></tr>
<tr><td>document_keyword_review.csv</td><td>7条撤销触发和3条显式待核，含论文标识与文本依据。</td></tr>
<tr><td>topic_definitions.json / 有效关键词主题映射.csv</td><td>72个主题定义及1221个有效词；当前一个词对应一个主题。</td></tr>
<tr><td>paper_assignments/*.parquet</td><td>全量论文的doc_id、主题列表、归类状态、能源锚点与质量标记。</td></tr>
<tr><td>paper_topic_relations/*.parquet</td><td>最终论文—主题关系与实际命中词、来源关键词ID。</td></tr>
<tr><td>paper_semantic_review.csv / reviewed_examples.parquet</td><td>181条逐项评语及审核前后归属；后者附带题名、摘要摘录和命中证据。</td></tr>
<tr><td>previous_semantic_issues.csv</td><td>旧版问题回归及D028评语补正。</td></tr>
<tr><td>missing_keyword_suggestions.csv</td><td>7条补词建议，均未写入来源或直接施加归属。</td></tr>
<tr><td>待审核关键词.csv</td><td>未读词按论文频次排序；辅助词的语境待核请另看完整词典。</td></tr>
<tr><td>platform_data.json</td><td>主题计数、181条样本及审核摘要；全量关系需读取Parquet。</td></tr>
<tr><td>VALIDATION.json / SEMANTIC_REVIEW_CHECK.json / RELEASE.json</td><td>全量规则检查、审核回归检查及交付文件SHA256。</td></tr></table></div>
<p>classification_status=assigned 表示命中；no_keywords 表示无关键词；conditional_only_no_energy_anchor 表示只有条件词但缺少合适锚点；no_active_dictionary_keyword 表示未命中有效词。保留撤稿和模板记录，可用 quality_eligible=true 过滤。原始关键词分数用于追溯，不是本方案的归属概率。</p>
<pre>import duckdb
db = duckdb.connect()
# 在本交付目录运行；跨主题汇总按 doc_id 去重。
db.sql("SELECT topic_code, count(*) FROM read_parquet('paper_topic_relations/*.parquet') WHERE quality_eligible GROUP BY topic_code").show()
db.sql("SELECT * FROM read_parquet('paper_assignments/*.parquet') WHERE doc_id = 'paper:W3140535007'").show()</pre>
<h2>复现与继续审核</h2>
<p>实际沿用第一版已验证的全量提取快照，PREPARED.json 记录的是该次提取。第二轮数据库的 observations、papers 指向冻结文件视图。完整运行顺序和复用的输入见配置与 INPUT_REUSE.json；词典复审决策在 round2_review.py。修改交付CSV不会自动改变程序规则。</p>
<pre># 项目根目录；仅在规则修改后重跑分类。第二轮数据库已存在。
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary_review_r2/build_dictionary.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary_review_r2/classify.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary_review_r2/validate.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary_review_r2/audit_r2.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary_review_r2/publish_r2.py</pre>
<p>fresh_topic_samples.parquet 和 new_direct_term_samples.parquet 已固定哈希。它们是本次已审诊断样本，勿重抽后复用旧评语；新增盲测应另建批次。第一版交付全部514个文件已校验未变。</p>
<h2>剩余问题与下一批重点</h2><ul>
<li>来源错标、漏词和题录错误仍存在；只凭词名无法解决。R125题名像期刊名，但摘要是真实电动车论文内容，未据题名删除。</li>
<li>主题相关可能只是背景或实验条件，下一批优先区分“主要研究对象/辅助方法/背景关联”，特别是电光源、材料和通用能源主题。</li>
<li>收紧泛化词会降低召回。R134车队IVR可能属于通信应用，因缺明确关键词暂未归入通信主题，已记录补词候选。</li>
<li>医学供能、航天核能及油气等属于广义能源关联，是否纳入招标业务范围需单独审定，不能与语义误分混为一谈。</li>
<li>该路线是助手整理的受控词典多主题分类；本轮仅处理论文，没有重跑NMF，也没有迁移专利、政策或重算趋势、热点及TRL/CRL。</li></ul></html>'''

def main():
    assert readjson('VALIDATION.json')['passed']
    assert readjson('SEMANTIC_REVIEW_CHECK.json')['passed']
    ASSETS.mkdir(parents=True,exist_ok=True)
    names=['all_raw_keywords.csv','keyword_dictionary.csv','keyword_aliases.csv','topic_catalog.csv',
      'topic_definitions.json','topic_counts.csv','topics_per_paper.csv','keywords_per_paper.csv',
      'assignment_status_counts.csv','unobserved_candidate_terms.csv','PREPARED.json',
      'DICTIONARY_COMPLETE.json','CLASSIFIED.json','VALIDATION.json','REVIEW.json','SEMANTIC_REVIEW_CHECK.json',
      'term_semantic_review.csv','document_keyword_review.csv','fresh_topic_samples.parquet',
      'new_direct_term_samples.parquet','review_metadata.parquet','sample_source_keywords.csv',
      'previous_semantic_review_regression.parquet','previous_semantic_issues.csv',
      'paper_semantic_review.csv','missing_keyword_suggestions.csv','SAMPLES.json']
    for name in names:shutil.copy2(WORK/name,ASSETS/name)
    shutil.copy2(HERE/'review_sample_hashes.json',ASSETS/'review_sample_hashes.json')
    for folder in ['paper_assignments','paper_topic_relations']:
        (ASSETS/folder).mkdir(exist_ok=True)
        for f in (WORK/folder).glob('*.parquet'):shutil.copy2(f,ASSETS/folder/f.name)
    c=connect()
    views=c.execute("SELECT view_name,sql FROM duckdb_views() WHERE view_name IN ('papers','observations')").df()
    write(ASSETS/'INPUT_REUSE.json',{'source_release':'keyword_dictionary_multilabel_20261019',
      'prepared_manifest_is_reused':True,'source_extraction_repeated_in_r2':False,
      'view_definitions':views.to_dict('records'),
      'reused_work_files':{n:sha(WORK/n) for n in ['all_raw_keywords.csv','raw_name_aliases.parquet','PREPARED.json']},
      'old_release_manifest_sha256':sha(ROOT/'final/assets/keyword_dictionary_multilabel/RELEASE.json')})
    d=pd.read_csv(WORK/'keyword_dictionary.csv').fillna('')
    catalog=pd.read_csv(WORK/'topic_catalog.csv').fillna('')
    topics=catalog.merge(pd.read_csv(WORK/'topic_counts.csv'),on='topic_code',validate='one_to_one')
    oldtopics=pd.read_csv(OLD/'topic_catalog.csv')[['topic_code','name']].merge(pd.read_csv(OLD/'topic_counts.csv'),on='topic_code')
    comparison=topics[['topic_code','name','documents','quality_eligible_documents']].merge(
      oldtopics[['topic_code','name','documents','quality_eligible_documents']].rename(columns={
        'name':'name_before','documents':'documents_before','quality_eligible_documents':'quality_eligible_documents_before'}),
      on='topic_code',how='left').fillna({'name_before':'新增主题','documents_before':0,'quality_eligible_documents_before':0})
    comparison['document_delta']=comparison.documents-comparison.documents_before
    for col in ['documents_before','quality_eligible_documents_before','document_delta']:comparison[col]=comparison[col].astype(int)
    comparison.to_csv(ASSETS/'topic_count_comparison.csv',index=False,encoding='utf-8-sig')
    stats=readjson('CLASSIFIED.json');ds=readjson('DICTIONARY_COMPLETE.json');review=readjson('REVIEW.json')
    checked=pd.read_csv(WORK/'paper_semantic_review.csv').fillna('')
    fresh=pd.read_parquet(WORK/'fresh_topic_samples.parquet')
    new=pd.read_parquet(WORK/'new_direct_term_samples.parquet')
    new['matched_keywords']=new.display_name.map(lambda s:[s]);new['has_direct_hit']=True
    new['matched_keyword_codes']=new.keyword_code.map(lambda s:[s])
    anchors=c.execute('SELECT row_id,anchor_keywords,anchor_topic_codes FROM assignments WHERE row_id IN (SELECT row_id FROM read_parquet(?))',[str(WORK/'new_direct_term_samples.parquet')]).df()
    new=new.merge(anchors,on='row_id',validate='many_to_one')
    examples=pd.concat([fresh,new],ignore_index=True)
    annotations=checked.drop(columns=['row_id','doc_id','topic_code','title'])
    examples=examples.merge(annotations,on='evidence_id',validate='one_to_one')
    current=c.execute('''SELECT row_id,topic_code,matched_keywords current_matched_keywords FROM topic_relations
      WHERE row_id IN (SELECT row_id FROM read_parquet(?) UNION SELECT row_id FROM read_parquet(?))''',
      [str(WORK/'fresh_topic_samples.parquet'),str(WORK/'new_direct_term_samples.parquet')]).df()
    examples=examples.merge(current,on=['row_id','topic_code'],how='left',validate='many_to_one')
    for name in ['current_matched_keywords','matched_source_keyword_ids']:
        examples[name]=examples[name].map(lambda x:x if isinstance(x,(list,tuple)) or hasattr(x,'tolist') else [])
    examples.to_parquet(ASSETS/'reviewed_examples.parquet',index=False)
    statuses=pd.read_csv(WORK/'assignment_status_counts.csv');distrib=pd.read_csv(WORK/'topics_per_paper.csv')
    unreviewed=d[d.decision.eq('unreviewed')].sort_values('document_frequency',ascending=False)
    unreviewed.to_csv(ASSETS/'待审核关键词.csv',index=False,encoding='utf-8-sig')
    active=d[d.decision.isin(['direct','conditional'])].merge(catalog[['topic_code','name']],on='topic_code')
    active.to_csv(ASSETS/'有效关键词主题映射.csv',index=False,encoding='utf-8-sig')
    summary=[('版本',stats['release_id']),('路线','真实关键词 → 助手整理词典 → 有条件精确匹配 → 多主题'),
      ('词名审读',review['term_name_reviewed']),('本轮首次阅读词名',review['newly_read_terms']),
      ('尚未逐项审读',review['unread_terms_after']),('论文诊断实例',len(examples)),
      ('样本用途','用于诊断及规则迭代，非独立盲测；题名支持不代表已读全文或主旨确认。'),
      ('审核前后','命中词为抽样时快照；retained_after_review表示最终主题归属，current_matched_keywords为最终命中词。'),
      ('逐篇撤销触发',review['document_trigger_withholdings']),
      ('旧疑似关系不再出现',f"{review['old_flagged_relations_no_longer_present']}/{review['old_flagged_semantic_relations']}"),
      ('来源','沿用4,831,088篇冻结论文；OpenAlex来源标签不保证是作者关键词。'),
      ('已归类论文',stats['counts']['assigned_papers']),('多主题论文',stats['counts']['multilabel_papers']),
      ('关系数',stats['counts']['relations']),('主题数',len(topics)),
      ('限制','程序校验通过不等于语义验收通过；背景关联、缺词、错标、业务边界仍待核。'),
      ('原数据','未删除论文或原始关键词，未给来源补造标签；原版本保留。')]
    workbook([('阅读说明',pd.DataFrame(summary,columns=['项目','说明或数值'])),('主题与篇数',topics),
      ('主题篇数前后',comparison),('本轮词名审核',pd.read_csv(WORK/'term_semantic_review.csv')),
      ('有效词主题映射',active),('完整关键词词典',d),('待逐项审核词',unreviewed),
      ('论文语义审核',examples),('逐篇词语义处理',pd.read_csv(WORK/'document_keyword_review.csv')),
      ('旧问题回归',pd.read_csv(WORK/'previous_semantic_issues.csv')),
      ('补词建议未施加',pd.read_csv(WORK/'missing_keyword_suggestions.csv')),
      ('论文归类状态',statuses),('每篇主题数量',distrib)])
    payload={'release_id':stats['release_id'],'method':'assistant_curated_keyword_multilabel',
      'counts':stats['counts'],'dictionary':ds,'status_names':base.STATUS_NAMES,'decision_names':base.STATE_NAMES,
      'statuses':records(statuses),'distribution':records(distrib),'topics':records(topics),
      'examples':records(examples),'review':review,'topic_count_comparison':records(comparison),
      'semantic_accuracy_validated':False,'independent_blind_test':False,
      'sample_evidence_timing':'sampling-time keyword snapshots; retained_after_review and current_matched_keywords are final',
      'full_paper_assignments':'paper_assignments/*.parquet','full_membership_evidence':'paper_topic_relations/*.parquet',
      'dictionary_path':'keyword_dictionary.csv','active_mapping_path':'有效关键词主题映射.csv'}
    write(ASSETS/'platform_data.json',payload)
    preview=dict(payload);preview['words']=d[['keyword_code','display_name','decision','topic_code','document_frequency','decision_reason']].values.tolist()
    page=base.PAGE.replace('关键词词典与多主题审阅.xlsx','第二轮语义审核.xlsx').replace('sample_review.csv','paper_semantic_review.csv')
    page=page.replace('2026-10-19 · 初审版','2026-10-19 · 第二轮语义审核')
    page=page.replace('能源关键词词典 · 多主题归类','能源关键词词典 · 第二轮语义审核')
    page=page.replace('能源关键词词典 · 多主题浏览','能源关键词词典 · 第二轮语义审核')
    page=page.replace('个长尾词仍待逐项审核','个词仍待逐项审核')
    page=page.replace('确定性抽样证据','本轮诊断样本与复核结果')
    page=page.replace('<span class="tag">${e.evidence_id}', '<span class="tag">${e.retained_after_review?"当前归属保留":"复核后已撤销"}</span> <span class="tag">${e.evidence_id}')
    page=page.replace('<p>命中词：${E(e.matched_keywords.join(\' / \'))}</p>',
      '<p>抽样时命中词：${E(e.matched_keywords.join(\' / \'))}</p><p>当前命中词：${E((e.current_matched_keywords||[]).join(\' / \'))||"无"}</p>${e.sample_kind==="new_direct_term"?`<p>该新增词触发：${e.sample_keyword_retained_after_review?"保留":"已撤销；粗主题仍可能由其他词支持"}</p>`:""}')
    page=page.replace('初审：${E(e.review_note)}','复核：${E(e.review_note)}')
    page=page.replace('校验与语义初审','校验与第二轮语义审核').replace('下载逐主题样本初审','下载181条逐项评语')
    page=page.replace('<h3>本轮诊断样本与复核结果</h3>',
      '<h3>本轮诊断样本与复核结果</h3><p class="muted">这些样本用于修正规则，不是独立准确率测试。抽样时命中词与最终归属分别展示；关联不等于论文主旨。</p>')
    page=page.replace('__PAYLOAD__',json.dumps(preview,ensure_ascii=False,separators=(',',':'),default=str,allow_nan=False).replace('</','<\\/'))
    (ASSETS/'关键词主题浏览.html').write_text(page)
    (ASSETS/'使用与接入说明.html').write_text(make_manual(stats,ds,review,comparison))
    config={'release_id':stats['release_id'],'assets':str(ASSETS.relative_to(ROOT)),
      'dictionary':str((ASSETS/'keyword_dictionary.csv').relative_to(ROOT)),
      'assignments':str((ASSETS/'paper_assignments/*.parquet').relative_to(ROOT)),
      'relations':str((ASSETS/'paper_topic_relations/*.parquet').relative_to(ROOT)),
      'document_keyword_review':str((ASSETS/'document_keyword_review.csv').relative_to(ROOT)),
      'pipeline':str(HERE.relative_to(ROOT)),'topic_id_namespace':'KWD-T0001..KWD-T0072',
      'conditional_anchor_rule':'per-topic compatible direct anchors; explicit global-support exceptions',
      'document_withholding_before_anchor_construction':True,'quality_filter':'quality_eligible = true',
      'paper_status_field':'classification_status','nmf_threshold':None,'maximum_topics_per_paper':None,
      'supports_patents':False,'supports_policies':False,'trends_recomputed':False,'maturity_recomputed':False,
      'semantic_accuracy_validated':False}
    write(ROOT/'final/config/keyword_dictionary_semantic_review_r2.json',config)
    manifest={'release_id':stats['release_id'],'method':'curated_dictionary_exact_match_multilabel_semantic_review_r2',
      'scope':'Same frozen 4,831,088 papers; raw source keywords retained; 7 explicit per-paper trigger withholdings',
      'semantic_accuracy_validated':False,'independent_blind_test':False,'counts':stats['counts'],
      'files':{str(p.relative_to(ASSETS)):sha(p) for p in sorted(ASSETS.rglob('*')) if p.is_file() and p.name!='RELEASE.json'},
      'pipeline_files':{str(p.relative_to(ROOT)):sha(p) for p in sorted(HERE.glob('*')) if p.suffix in ['.py','.json']},
      'presentation_dependency':{str(V1_PUBLISH.relative_to(ROOT)):sha(V1_PUBLISH)}}
    write(ASSETS/'RELEASE.json',manifest)
    print(json.dumps({'assets':str(ASSETS),'files':len(manifest['files']),'counts':stats['counts']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
