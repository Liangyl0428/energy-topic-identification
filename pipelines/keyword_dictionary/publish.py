"""Package the keyword route with complete evidence, an offline browser and workbook."""
from common import *
import shutil, html
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

STATE_NAMES={'direct':'直接触发','conditional':'条件触发','general_auxiliary':'通用辅助词',
    'broad_auxiliary':'学科大类','excluded':'逐项排除词','excluded_rule':'规则筛出待复核','unreviewed':'尚未逐项审核'}
STATUS_NAMES={'assigned':'已归类','no_keywords':'无关键词','conditional_only_no_energy_anchor':'仅有条件词，缺明确能源词',
    'no_active_dictionary_keyword':'未命中当前触发词典'}

def readjson(name): return json.loads((WORK/name).read_text())

def safe(v):
    if hasattr(v,'tolist'):v=v.tolist()
    if isinstance(v,(list,dict)):v=json.dumps(v,ensure_ascii=False)
    if v is None or (isinstance(v,float) and pd.isna(v)):return ''
    if isinstance(v,str):
        v=re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]','',v)[:32000]
        if v.startswith(('=','+','-','@')):v="'"+v
    return v

def main():
    assert readjson('VALIDATION.json')['passed'],'Cannot release invalid data'
    ASSETS.mkdir(parents=True,exist_ok=True)
    names=['all_raw_keywords.csv','keyword_dictionary.csv','keyword_aliases.csv','topic_catalog.csv',
      'topic_definitions.json','topic_counts.csv','topics_per_paper.csv','keywords_per_paper.csv',
      'assignment_status_counts.csv','unobserved_candidate_terms.csv','PREPARED.json',
      'DICTIONARY_COMPLETE.json','CLASSIFIED.json','VALIDATION.json','new_topic_samples.parquet',
      'previous_audit_regression.parquet','review_metadata.parquet','sample_all_keywords.csv','sample_review.csv','rule_refinement_regression.csv','REVIEW.json']
    for name in names:shutil.copy2(WORK/name,ASSETS/name)
    for folder in ['paper_assignments','paper_topic_relations']:
        (ASSETS/folder).mkdir(exist_ok=True)
        for f in (WORK/folder).glob('*.parquet'):shutil.copy2(f,ASSETS/folder/f.name)
    d=pd.read_csv(WORK/'keyword_dictionary.csv').fillna('')
    catalog=pd.read_csv(WORK/'topic_catalog.csv').fillna('')
    counts=pd.read_csv(WORK/'topic_counts.csv')
    topics=catalog.merge(counts,on='topic_code',how='left').fillna(0)
    stats=readjson('CLASSIFIED.json');ds=readjson('DICTIONARY_COMPLETE.json');review=readjson('REVIEW.json')
    samples=pd.read_parquet(WORK/'new_topic_samples.parquet')
    checked=pd.read_csv(WORK/'sample_review.csv').fillna('')
    samples=samples.merge(checked,on='evidence_id',validate='one_to_one')
    statuses=pd.read_csv(WORK/'assignment_status_counts.csv')
    distrib=pd.read_csv(WORK/'topics_per_paper.csv')
    unreviewed=d[d.decision.eq('unreviewed')].sort_values('document_frequency',ascending=False)
    unreviewed.to_csv(ASSETS/'待审核关键词.csv',index=False,encoding='utf-8-sig')
    active=d[d.decision.isin(['direct','conditional'])].merge(catalog[['topic_code','name']],on='topic_code')
    active.to_csv(ASSETS/'有效关键词主题映射.csv',index=False,encoding='utf-8-sig')
    summary_rows=[('数据版本',stats['release_id']),('路线','实际关键词 → 助手逐项整理触发词 → 精确匹配 → 多标签'),
       ('说明','主题由助手整理，尚未专家审定；不是自动主题发现。'),
       ('论文总数',stats['counts']['papers']),('已归类论文',stats['counts']['assigned_papers']),
       ('多主题论文',stats['counts']['multilabel_papers']),('论文主题关系',stats['counts']['relations']),
       ('主题数',len(topics)),('原始词条ID',ds['raw_keyword_ids']),('规范化词典词条',len(d)),
       ('待逐项审核词',len(unreviewed)),('规则说明','每篇论文可命中全部主题；具体技术条件词须同篇相关方向能源锚点；通用支撑词须任一能源锚点，不设主题上限。'),
       ('抽查说明','每主题2条确定性调试抽样，部分样本用于修正规则；非独立盲测，不能估计全库准确率。'),
       ('来源限制','OpenAlex来源标签不等同于作者关键词；错误/缺失标签仍会误分/漏分。'),
       ('计数说明','各主题篇数相加会重复计数；按主题内部去重；跨主题合计按doc_id去重。'),
       ('统计过滤','全库保留撤稿/模板记录；使用quality_eligible字段可过滤。'),
       ('业务范围','广义能源与电力，不等同于招标所有模块验收；尚未重算热点、TRL/CRL。')]
    wb=Workbook(write_only=True)
    sheets=[('阅读说明',pd.DataFrame(summary_rows,columns=['项目','说明或数值'])),('主题与篇数',topics),
      ('有效词主题映射',active),('完整关键词词典',d),('待逐项审核词',unreviewed),
      ('论文归类状态',statuses),('每篇主题数量',distrib),('逐主题样本初审',samples),
      ('旧误分回归',pd.read_parquet(WORK/'previous_audit_regression.parquet'))]
    for name,frame in sheets:
        ws=wb.create_sheet(name);ws.freeze_panes='A2'
        header=[]
        for label in frame.columns:
            cell=WriteOnlyCell(ws,value=label);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='184E58');header.append(cell)
        ws.append(header)
        for row in frame.itertuples(index=False,name=None):ws.append([safe(x) for x in row])
        from openpyxl.utils import get_column_letter
        ws.auto_filter.ref=f'A1:{get_column_letter(len(frame.columns))}{len(frame)+1}'
    wb.save(ASSETS/'关键词词典与多主题审阅.xlsx')
    sample_records=[]
    for row in samples.to_dict('records'):
        sample_records.append({k:(v.tolist() if hasattr(v,'tolist') else v) for k,v in row.items()})
    payload={'release_id':stats['release_id'],'method':'assistant_curated_keyword_multilabel',
       'counts':stats['counts'],'dictionary':ds,'status_names':STATUS_NAMES,'decision_names':STATE_NAMES,
       'statuses':statuses.to_dict('records'),'distribution':distrib.to_dict('records'),
       'topics':topics.to_dict('records'),'examples':sample_records,'review':review,
       'semantic_accuracy_validated':False,'full_paper_assignments':'paper_assignments/*.parquet',
       'full_membership_evidence':'paper_topic_relations/*.parquet',
       'dictionary_path':'keyword_dictionary.csv','active_mapping_path':'有效关键词主题映射.csv'}
    write(ASSETS/'platform_data.json',payload)
    preview=dict(payload)
    preview['words']=d[['keyword_code','display_name','decision','topic_code','document_frequency','decision_reason']].values.tolist()
    page=PAGE.replace('__PAYLOAD__',json.dumps(preview,ensure_ascii=False,separators=(',',':'),default=str).replace('</','<\\/'))
    (ASSETS/'关键词主题浏览.html').write_text(page)
    (ASSETS/'使用与接入说明.html').write_text(make_manual(stats,ds,review))
    manifest={'release_id':stats['release_id'],'created_utc':'2026-10-19','method':'curated_dictionary_exact_match_multilabel',
        'scope':'Frozen 4,831,088 paper records only; patents/policies not transferred',
        'source_keywords':'All observed unfiltered IDs from canonical source versions, including low-frequency and zero-score labels',
        'semantic_accuracy_validated':False,'curation_status':'initial assistant curation; explicit unresolved tail',
        'counts':stats['counts'],'files':{}}
    for path in sorted(ASSETS.rglob('*')):
        if path.is_file() and path.name!='RELEASE.json':manifest['files'][str(path.relative_to(ASSETS))]=sha(path)
    manifest['pipeline_files']={str(p.relative_to(ROOT)):sha(p) for p in sorted(HERE.glob('*.py'))}
    write(ASSETS/'RELEASE.json',manifest)
    config={'release_id':stats['release_id'],'assets':'final/assets/keyword_dictionary_multilabel',
       'dictionary':'final/assets/keyword_dictionary_multilabel/keyword_dictionary.csv',
       'assignments':'final/assets/keyword_dictionary_multilabel/paper_assignments/*.parquet',
       'relations':'final/assets/keyword_dictionary_multilabel/paper_topic_relations/*.parquet',
       'topic_id_namespace':'KWD-T0001..KWD-T0071','conditional_anchor_rule':'per-topic compatible direct anchors; explicit global-support exceptions','quality_filter':'quality_eligible = true',
       'paper_status_field':'classification_status','nmf_threshold':None,'maximum_topics_per_paper':None,
       'supports_patents':False,'supports_policies':False,'trends_recomputed':False,'maturity_recomputed':False}
    write(ROOT/'final/config/keyword_dictionary_multilabel.json',config)
    print(json.dumps({'assets':str(ASSETS),'files':len(manifest['files']),'counts':stats['counts']},ensure_ascii=False,indent=2))

def make_manual(stats,ds,review):
    n=stats['counts'];coverage=n['assigned_papers']/n['papers']
    return '''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>关键词词典多主题：使用与接入说明</title>
    <style>body{max-width:1000px;margin:45px auto;padding:0 24px;font:17px/1.8 system-ui;color:#243b43}h1,h2{color:#124653}pre{background:#eff5f5;padding:18px;white-space:pre-wrap}table{border-collapse:collapse}td,th{padding:9px;border:1px solid #bbc}a{color:#167b86}.note{background:#fff4dc;padding:18px}</style>
    <h1>关键词词典与论文多主题归类</h1><p>版本：keyword_dictionary_multilabel_20261019。根据本库实际词汇，由助手整理主题和触发词，再按明确规则匹配论文。主题词可编辑、每条归属可追溯；这是一版受控词典分类结果，尚未经专家审定。</p>
    <p><a href="关键词主题浏览.html">打开主题及词典浏览器</a> · <a href="关键词词典与多主题审阅.xlsx">审阅工作簿</a> · <a href="platform_data.json">平台接口摘要</a></p>
    <h2>当前结果</h2>'''+f'''<p>处理 {n['papers']:,} 篇论文，{ds['raw_keyword_ids']:,} 个原始关键词ID，合并明确同义词后 {ds['canonical_keywords']:,} 条。共 {ds['topics']} 个主题；{n['assigned_papers']:,} 篇命中（{coverage:.1%}），其中 {n['multilabel_papers']:,} 篇属于多个主题，形成 {n['relations']:,} 条论文—主题关系。</p>
    <p class="note">当前仍有 {ds['decisions']['unreviewed']:,} 个词未逐项审核。完整词典已经提取，人工词表整理尚未覆盖全部长尾词。未命中只说明当前词典没有识别到，不能解释为论文与能源无关。覆盖率不等于召回率。</p>'''+'''
    <h2>归类规则</h2><ol><li>使用全部冻结论文的完整来源关键词，保留低频词及原始分数为0的词。关键词来自来源数据/OpenAlex标签，不保证是作者提供的关键词。</li>
    <li>按Unicode、大小写、连字符与空白规范化；只合并明确列出的同义词，保留括号内词义。通过原始关键词ID匹配，不做模糊子串匹配。</li>
    <li>命中直接触发词，归入对应主题。具体技术的条件触发词须同篇有规定相关方向的直接能源词；通用方法、材料、经济等支撑主题接受任一直接能源词。allowed families见主题定义中的conditional_anchor_topics，空列表表示任一能源锚点。每篇取全部主题的并集；不设贡献阈值、不设主题上限。</li>
    <li>学科大类和通用辅助词不单独触发。明确范围外词、规则筛出词、未审核词各自保留状态，不删除原始词或整篇论文。</li>
    <li>原始分数保留作追溯，不作为归属置信度。共现满足规则不代表通用材料/方法一定是该论文的主要研究方向。</li></ol>
    <p>例如“Lithium-ion battery + Machine learning + Graphene”可同时归入电池、能源人工智能、能源材料结构三个主题；仅“Machine learning + Graphene”没有明确能源词，会留为待判断。</p>
    <h2>开发接入</h2><table><tr><th>文件</th><th>用途</th></tr>
    <tr><td>keyword_dictionary.csv / keyword_aliases.csv</td><td>完整规范词典、处理依据、文献频次及来源ID映射。</td></tr>
    <tr><td>topic_definitions.json / 有效关键词主题映射.csv</td><td>主题定义、直接/条件触发词，可作为编辑与审核入口。</td></tr>
    <tr><td>paper_assignments/*.parquet</td><td>所有论文的doc_id、主题列表、归类状态、质量标记与能源锚点。</td></tr>
    <tr><td>paper_topic_relations/*.parquet</td><td>逐条论文主题关系，含命中词及来源关键词ID、直接/条件标记。</td></tr>
    <tr><td>platform_data.json</td><td>71个主题、计数和抽样示例，供页面初次展示；不是全库关系文件。</td></tr>
    <tr><td>待审核关键词.csv</td><td>按文献频次排序的词典待审队列。</td></tr>
    <tr><td>sample_review.csv / REVIEW.json</td><td>逐主题样本初审、语义风险与旧问题回归证据。</td></tr>
    <tr><td>VALIDATION.json / RELEASE.json</td><td>全库规则完整性校验与交付文件SHA256。</td></tr></table>
    <p>主题编号使用KWD-T命名空间，与F、RBK500编号不能直接互换。主题数71来自当前词典整理，不是聚类K。一个词当前对应一个主题，同一论文可有多个主题。</p>
    <pre>import duckdb
db = duckdb.connect()
# 在本交付目录运行；跨主题去重请使用 doc_id。
db.sql("SELECT topic_code, count(*) FROM read_parquet('paper_topic_relations/*.parquet') WHERE quality_eligible GROUP BY topic_code").show()
db.sql("SELECT * FROM read_parquet('paper_assignments/*.parquet') WHERE doc_id = 'paper:W3140535007'").show()</pre>
    <p>status=assigned表示已归类；no_keywords表示没有关键词；conditional_only_no_energy_anchor表示有条件词但无能源锚点；no_active_dictionary_keyword表示未命中当前触发词。quality_eligible=false表示撤稿或模板记录，仍保留字面归属，统计时可排除。主题计数分全量与quality_eligible两种；各主题篇数相加不能作为独立论文总量。</p>
    <h2>如何审阅和修改</h2><p>先在浏览器选择主题查看触发词和论文证据，再在工作簿逐项审阅。建议优先处理高频待审词、直接词的多义性、条件词造成的额外主题。词表配置位于 final/energy-topic-identification/pipelines/keyword_dictionary/curation.py、curation_additions.py 和 refine_rules.py，分类程序为 classify.py。CSV是本次配置的产物，直接修改CSV不会自动更新规则；更改配置后按以下流程重新生成并重新验收。</p>
    <pre># 在项目根目录运行；初次提取用prepare.py，已提取快照可复用。
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary/build_dictionary.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary/classify.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary/validate.py
.venv-hotspots/bin/python final/energy-topic-identification/pipelines/keyword_dictionary/samples.py
# 逐项审阅样本，更新sample_review.csv及REVIEW.json后运行publish.py。</pre>
    <h2>已知边界</h2><ul><li>源关键词错标或缺失时，这条路线仍会漏分或误分。题名/摘要未用于补词或强制回填。</li>
    <li>“任一能源词 + 通用方法词”的共现规则会产生额外方法/材料标签。审阅结果提供具体例子，不能把程序校验通过当作语义准确率通过。</li>
    <li>语义抽查每主题2条，共142条关系；部分抽样用于迭代词典，非独立盲测、非专家验收，不足以推断总体准确率。样本及评语全部可查。</li>
    <li>当前为广义能源电力词典，含油气、供热、生物能源等；若业务仅保留电网特定范围，应审阅主题开关，而非物理删除论文。</li>
    <li>本次处理论文，不自动迁移专利/政策，没有重算趋势、热点和TRL/CRL。旧NMF文件保留并校验哈希一致。</li></ul></html>'''

PAGE='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>能源关键词词典 · 多主题浏览</title>
<style>:root{--ink:#173e48;--muted:#536c72;--line:#dce6e5;--accent:#087e8b}*{box-sizing:border-box}body{margin:0;background:#f3f6f5;color:var(--ink);font:15px/1.6 system-ui,sans-serif}header{background:#153f48;color:white;padding:30px max(24px,calc((100vw - 1440px)/2));}h1{margin:0;font-size:29px}header p{margin:8px 0 0;color:#cce2e3}main{max-width:1488px;margin:auto;padding:24px}a{color:var(--accent)}header a{color:#b6e5eb}.stats{display:grid;grid-template-columns:repeat(5,1fr);gap:14px}.stat,section{background:white;border:1px solid var(--line);border-radius:10px;padding:19px}.stat strong{font-size:27px;display:block}.muted{color:var(--muted)}.note{padding:16px 20px;background:#fff3dc;border-left:4px solid #c78922;border-radius:5px;margin:18px 0}nav{display:flex;gap:8px;margin:20px 0}button,select,input{font:inherit;border:1px solid #b8cece;padding:9px 12px;border-radius:6px;background:white;color:var(--ink)}button{cursor:pointer}button.active{background:var(--ink);color:white}.layout{display:grid;grid-template-columns:370px 1fr;gap:18px}.topiclist{max-height:830px;overflow:auto}.topic{display:block;width:100%;text-align:left;margin-bottom:8px}.topic small{display:block;opacity:.7}.topic.selected{background:#e0f1f0;border-color:var(--accent)}h2{margin-top:0}h3{margin:23px 0 8px}.chips{display:flex;flex-wrap:wrap;gap:7px}.chip{background:#e9f3f3;padding:4px 9px;border-radius:5px}.conditional{background:#f8f1e5}.example{padding:14px 0;border-top:1px solid var(--line)}.example p{margin:6px 0}.tag{font-size:12px;padding:3px 6px;border:1px solid #b6ccce;border-radius:4px}.tools{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:12px}input{flex:1;min-width:220px}table{width:100%;border-collapse:collapse}th,td{text-align:left;padding:10px;border-bottom:1px solid var(--line);vertical-align:top}th{background:#edf4f4}details{margin-top:8px}summary{cursor:pointer}.scroll{overflow-x:auto}[hidden]{display:none!important}footer{margin-top:20px;font-size:13px;color:var(--muted)}@media(max-width:900px){.stats{grid-template-columns:repeat(2,minmax(0,1fr))}.stat strong{font-size:23px}.chip{max-width:100%;overflow-wrap:anywhere}.layout{grid-template-columns:1fr}.topiclist{max-height:350px}main{padding:14px}h1{font-size:24px}}</style>
<header><h1>能源关键词词典 · 多主题归类</h1><p>从实际关键词建立受控词典，每条归属都能追溯到命中词。2026-10-19 · 初审版</p><p><a href="关键词词典与多主题审阅.xlsx">下载审阅工作簿</a>　<a href="使用与接入说明.html">使用与接入说明</a>　<a href="platform_data.json">接口数据</a></p></header>
<main><div class="stats" id="stats"></div><div class="note" id="note"></div>
<nav><button id="bt" class="active">主题与证据</button><button id="bd">完整关键词词典</button><button id="bs">归类状态与数量</button></nav>
<div id="topics" class="layout"><section><input id="tq" placeholder="搜索主题或触发词" style="width:100%;margin-bottom:14px"><div class="topiclist" id="topiclist"></div></section><section id="detail"></section></div>
<section id="dictionary" hidden><h2>完整关键词词典</h2><div class="tools"><input id="dq" placeholder="搜索关键词、主题编号或词条编号"><select id="state"><option value="">全部处理状态</option></select></div><p id="dn" class="muted"></p><div class="scroll"><table><thead><tr><th>词条</th><th>状态</th><th>主题</th><th>论文频次</th><th>处理依据</th></tr></thead><tbody id="drows"></tbody></table></div><p><button id="prev">上一页</button> <button id="next">下一页</button></p></section>
<section id="status" hidden><h2>论文归类状态</h2><p class="muted">未归类不等于范围外；各主题篇数相加会重复计算论文。</p><div id="stable"></div><h3>每篇论文的主题数量</h3><div id="hist"></div><h3>校验与语义初审</h3><p id="review"></p><p><a href="VALIDATION.json">查看全量规则校验</a> · <a href="sample_review.csv">下载逐主题样本初审</a> · <a href="REVIEW.json">查看已知问题与回归结果</a></p></section>
<footer>主题由助手整理，尚未专家审定；这版是词典规则分类。来源关键词可能缺失或标错，规则匹配不能消除这些问题。专利、政策、趋势和成熟度未在此次重算。</footer></main>
<script id="payload" type="application/json">__PAYLOAD__</script><script>
const D=JSON.parse(document.getElementById('payload').textContent),$=id=>document.getElementById(id),E=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c])),N=n=>Number(n).toLocaleString('zh-CN');
const C=D.counts,TM=Object.fromEntries(D.topics.map(t=>[t.topic_code,t]));let chosen=D.topics[0].topic_code,page=0,filtered=[];
$('stats').innerHTML=[['论文总数',N(C.papers)],['规范关键词',N(D.dictionary.canonical_keywords)],['主题数量',D.topics.length],['已归类论文',N(C.assigned_papers)],['多主题论文',N(C.multilabel_papers)]].map(([k,v])=>`<div class="stat"><span class="muted">${k}</span><strong>${v}</strong></div>`).join('');
$('note').textContent=`完整词典已提取；${N(D.dictionary.decisions.unreviewed)} 个长尾词仍待逐项审核。当前命中率 ${(100*C.assigned_papers/C.papers).toFixed(1)}%，不等于召回率。学科大类不触发主题；具体技术条件词须相关方向的能源锚点；通用支撑词须任一能源锚点。`;
function list(){let q=$('tq').value.toLowerCase();let ts=D.topics.filter(t=>[t.name,t.topic_code,t.direct_keywords,t.conditional_keywords].join(' ').toLowerCase().includes(q));$('topiclist').innerHTML=ts.map(t=>`<button class="topic ${chosen===t.topic_code?'selected':''}" data-topic="${t.topic_code}">${E(t.name)}<small>${t.topic_code} · ${N(t.documents)} 篇</small></button>`).join('')||'<p>没有匹配主题</p>';}
function chips(s,cls=''){return String(s||'').split(' | ').filter(Boolean).map(w=>`<span class="chip ${cls}">${E(w)}</span>`).join('')||'<span class="muted">无</span>'}
function detail(){const t=TM[chosen];$('detail').innerHTML=`<span class="muted">${t.topic_code}</span><h2>${E(t.name)}</h2><p>${E(t.definition)}</p><p><b>${N(t.documents)}</b> 篇 · 直接命中 ${N(t.direct_documents)} · 仅条件命中 ${N(t.conditional_documents)} · 质量过滤后 ${N(t.quality_eligible_documents)}</p><h3>直接触发词 <small>(${t.direct_keyword_count})</small></h3><div class="chips">${chips(t.direct_keywords)}</div><h3>条件触发词 <small>(${t.conditional_keyword_count})</small></h3><p class="muted">同篇允许的直接主题锚点：${E(t.conditional_anchor_topics || "任一能源主题")}。共现提供分类线索，不能证明该方法/材料是论文主旨。</p><div class="chips">${chips(t.conditional_keywords,'conditional')}</div><h3>确定性抽样证据</h3>`+D.examples.filter(e=>e.topic_code===chosen).map(e=>`<article class="example"><span class="tag">${e.evidence_id} · ${E(e.review_status)}</span><h3>${E(e.title)}</h3><p class="muted">${E(e.doc_id)} · ${e.has_direct_hit?'直接命中':'条件命中'}</p><p>命中词：${E(e.matched_keywords.join(' / '))}</p><p class="muted">同篇能源锚点：${E((e.anchor_keywords||[]).join(' / '))}</p><p>初审：${E(e.review_note)}</p><details><summary>题录内容摘录</summary><p>${E(e.body_excerpt)}</p></details></article>`).join('');}
$('topiclist').onclick=e=>{const b=e.target.closest('[data-topic]');if(b){chosen=b.dataset.topic;list();detail()}};$('tq').oninput=list;
for(const [k,v] of Object.entries(D.decision_names)){const o=document.createElement('option');o.value=k;o.textContent=v;$('state').append(o)}
function dict(reset=true){if(reset)page=0;const q=$('dq').value.toLowerCase(),s=$('state').value;filtered=D.words.filter(w=>(!s||w[2]===s)&&[w[0],w[1],w[3],TM[w[3]]?.name||''].join(' ').toLowerCase().includes(q));$('dn').textContent=`匹配 ${N(filtered.length)} 个词；第 ${page+1} 页，每页100条。频次为包含该词的论文数。`;$('drows').innerHTML=filtered.slice(page*100,(page+1)*100).map(w=>`<tr><td>${E(w[1])}<br><small>${w[0]}</small></td><td>${E(D.decision_names[w[2]])}</td><td>${E(TM[w[3]]?.name||'—')}<br><small>${w[3]}</small></td><td>${N(w[4])}</td><td>${E(w[5])}</td></tr>`).join('');$('prev').disabled=!page;$('next').disabled=(page+1)*100>=filtered.length;}
$('dq').oninput=()=>dict();$('state').onchange=()=>dict();$('prev').onclick=()=>{page--;dict(false)};$('next').onclick=()=>{page++;dict(false)};
function tab(id){for(const [b,p] of [['bt','topics'],['bd','dictionary'],['bs','status']]){$(p).hidden=p!==id;$(b).classList.toggle('active',p===id)}};$('bt').onclick=()=>tab('topics');$('bd').onclick=()=>{tab('dictionary');dict()};$('bs').onclick=()=>tab('status');
$('stable').innerHTML='<table><tr><th>状态</th><th>论文</th><th>占全部</th></tr>'+D.statuses.map(s=>`<tr><td>${E(D.status_names[s.classification_status])}</td><td>${N(s.papers)}</td><td>${(100*s.papers/C.papers).toFixed(2)}%</td></tr>`).join('')+'</table>';
$('hist').innerHTML='<table><tr><th>主题个数</th><th>论文</th></tr>'+D.distribution.map(x=>`<tr><td>${x.topic_count}</td><td>${N(x.papers)}</td></tr>`).join('')+'</table>';
$('review').textContent=D.review.summary;list();detail();
</script></html>'''

if __name__=='__main__':main()
